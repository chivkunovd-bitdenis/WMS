"""Review-only synthetic exports; never reads a database or external service."""
from __future__ import annotations

import asyncio
import hashlib
import io
import json
import sys
from datetime import UTC, datetime
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / "backend"))
import fitz
from openpyxl import load_workbook
from app.services import inbound_acceptance_act_service as svc

OUT = Path(__file__).resolve().parent


def request(count: int):
    lines = []
    for i in range(count):
        name = "=1+1 худи" if i == 0 else f"Товар-{i:03d} длинное кириллическое название"
        vendor = f"000АРТ-{i:03d}"
        if count > 50:
            name = (name + " зимняя коллекция утеплённая хлопковая одежда " * 3).strip()
            vendor += "-зимняя-коллекция-длинный-артикул-селлера"
        product = SimpleNamespace(name=name, wb_vendor_code=vendor, sku_code=f"00{i:04d}", wb_barcode=f"046000000{i:04d}")
        plan, fact = ((i % 11, i % 13) if count > 50 else [(10, 8), (5, 7), (4, 4), (0, 3)][i])
        lines.append(SimpleNamespace(product=product, expected_qty=plan, actual_qty=fact))
    return SimpleNamespace(id="review-only", display_number="№000586", document_number="INB-000586",
                           operation_type="inbound", created_at=datetime(2026, 9, 28, 21, 30, tzinfo=UTC),
                           status="sorting", lines=lines)


async def main():
    results = []
    for count, multiline in ((4, False), (321, False), (4, True)):
        req = request(count)
        if multiline:
            req.lines[0].product.name = "Худи\nРазмер XL\nЦвет чёрный\nЗимняя коллекция\nУтеплённый хлопок"
        with patch.object(svc, "get_request", AsyncMock(return_value=req)):
            _, pdf = await svc.build_acceptance_act_pdf(None, None, None)
            _, xlsx = await svc.build_acceptance_act_workbook(None, None, None)
        label = "act-multiline" if multiline else f"act-{count}"
        (OUT / f"{label}.pdf").write_bytes(pdf)
        (OUT / f"{label}.xlsx").write_bytes(xlsx)
        doc = fitz.open(stream=pdf, filetype="pdf")
        sheet = load_workbook(io.BytesIO(xlsx)).active
        normalized = "".join("".join(page.get_text().split()) for page in doc)
        losses = [i + 1 for i, line in enumerate(req.lines)
                  if "".join(line.product.name.split()) not in normalized]
        outside = []
        for n, page in enumerate(doc, 1):
            for block in page.get_text("dict")["blocks"]:
                for line in block.get("lines", []):
                    for span in line["spans"]:
                        if not page.rect.contains(fitz.Rect(span["bbox"])):
                            outside.append({"page": n, "text": span["text"], "bbox": span["bbox"]})
        for index in sorted({0, len(doc) // 2, len(doc) - 1}):
            doc[index].get_pixmap(dpi=110).save(OUT / f"{label}-page-{index+1}.png")
        if count == 321:
            # All pages in bounded contact sheets, for visual layout inspection.
            for start in range(0, len(doc), 8):
                contact = fitz.open()
                target = contact.new_page(width=1200, height=1760)
                for slot, index in enumerate(range(start, min(start + 8, len(doc)))):
                    x, y = (slot % 2) * 600, (slot // 2) * 440
                    target.insert_text((x + 10, y + 12), f"Page {index+1}", fontsize=10)
                    target.insert_image(fitz.Rect(x, y + 18, x + 590, y + 435),
                                        stream=doc[index].get_pixmap(matrix=fitz.Matrix(0.7, 0.7)).tobytes("png"))
                target.get_pixmap().save(OUT / f"act-321-contact-{start+1}-{min(start+8,len(doc))}.png")
        results.append({"case": label, "rows": count, "pages": len(doc), "page_sizes": sorted({tuple(p.rect) for p in doc}),
                        "empty_pages": [n for n, p in enumerate(doc, 1) if not p.get_text().strip()],
                        "lost_names": losses, "out_of_page_spans": outside,
                        "total_excel": [c.value for c in sheet[sheet.max_row]][5:],
                        "source_total": [sum(getattr(line, field) for line in req.lines) for field in ("expected_qty", "actual_qty")],
                        "first_name_excel": sheet["B4"].value,
                        "fonts": doc[0].get_fonts(), "pdf_sha256": hashlib.sha256(pdf).hexdigest(),
                        "xlsx_sha256": hashlib.sha256(xlsx).hexdigest()})
    (OUT / "pdf-probe.json").write_text(json.dumps(results, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(results, ensure_ascii=False))


if __name__ == "__main__":
    asyncio.run(main())
