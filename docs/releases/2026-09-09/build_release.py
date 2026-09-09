"""Rebuild the WMS-411 customer release document with bundled python-docx."""
from pathlib import Path
import json

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BASE = Path(__file__).resolve().parent
data = json.loads((BASE / 'release_content.json').read_text())
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.68), Inches(.66)
sec.left_margin, sec.right_margin = Inches(.82), Inches(.82)
sec.footer_distance = Inches(.30)

for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2'):
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Arial')
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
normal = doc.styles['Normal']
normal.font.size = Pt(11.5)
normal.paragraph_format.line_spacing = 1.13
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.widow_control = True
title = doc.styles['Title']
title.font.size, title.font.bold = Pt(28), True
title.paragraph_format.space_after = Pt(7)
title.paragraph_format.line_spacing = 1.05
heading = doc.styles['Heading 1']
heading.font.size, heading.font.bold = Pt(15), True
heading.paragraph_format.space_before = Pt(17)
heading.paragraph_format.space_after = Pt(6)
heading.paragraph_format.keep_with_next = True
subtitle = doc.styles['Subtitle']
subtitle.font.size = Pt(11)
subtitle.paragraph_format.space_after = Pt(17)
subtitle.font.italic = False

doc.core_properties.title = data['title']
doc.core_properties.subject = 'Новые возможности для селлеров и фулфилментов'
doc.core_properties.author = 'Короб ВМС'
doc.core_properties.keywords = 'WMS-411, обновления, сентябрь 2026'

for n, page in enumerate(data['pages']):
    if n:
        doc.add_page_break()
    doc.add_paragraph(page['title'], 'Title')
    doc.add_paragraph(data['period'] if n == 0 else 'Короб ВМС  •  ' + data['period'], 'Subtitle')
    if page.get('intro'):
        doc.add_paragraph(page['intro'])
    for block in page['sections']:
        doc.add_paragraph(block['title'], 'Heading 1')
        for text in block['paragraphs']:
            doc.add_paragraph(text)
        if block.get('note'):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            run = p.add_run(block['note'])
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor.from_string('505060')

footer = sec.footer.paragraphs[0]
footer.paragraph_format.space_after = Pt(0)
footer.paragraph_format.tab_stops.clear_all()
footer.paragraph_format.tab_stops.add_tab_stop(Inches(6.6), 2)
run = footer.add_run('Короб ВМС   /   Обновления за неделю\t')
run.font.name, run.font.size = 'Arial', Pt(9)
run.font.color.rgb = RGBColor.from_string('666672')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

for border in list(doc.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)

out = BASE / 'Короб ВМС — обновления 3–9 сентября 2026.docx'
doc.save(out)
print(out)
