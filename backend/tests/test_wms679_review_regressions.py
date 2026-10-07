"""WMS-679: API regressions found by independent review 2026-10-07.

These tests use the actual WMS route and its visibility checks.  They never
call Wildberries, create credentials, or write any external state.
"""

from __future__ import annotations

import uuid
from collections.abc import AsyncIterator
from contextlib import asynccontextmanager
from io import BytesIO

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook  # type: ignore[import-untyped]
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.deps import get_current_user, get_effective_seller_id
from app.db.session import get_db
from app.main import create_app
from app.models.marketplace_unload import (
    MarketplaceUnloadBox,
    MarketplaceUnloadBoxLine,
    MarketplaceUnloadRequest,
)
from app.models.product import Product
from app.models.seller import Seller
from app.models.tenant import Tenant
from app.models.user import User
from app.models.warehouse import Warehouse
from app.models.warehouse_box import WarehouseBox

_EXPORT_PATH = "/operations/marketplace-unload-requests/{request_id}/wb-fbw-packaging.xlsx"


def _xlsx_rows(content: bytes) -> list[tuple[object, ...]]:
    sheet = load_workbook(BytesIO(content), data_only=False)["Sheet1"]
    return [tuple(cell.value for cell in row) for row in sheet.iter_rows(min_row=2, max_col=5)]


class Ctx:
    tenant: Tenant
    other_tenant: Tenant
    warehouse: Warehouse
    other_warehouse: Warehouse
    foreign_warehouse: Warehouse
    seller: Seller
    user: User
    request: MarketplaceUnloadRequest
    box: MarketplaceUnloadBox
    product: Product
    foreign_box: WarehouseBox
    other_warehouse_box: WarehouseBox


async def _fixture(session: AsyncSession) -> Ctx:
    suffix = uuid.uuid4().hex[:8]
    ctx = Ctx()
    ctx.tenant = Tenant(name="WMS-679 review", slug=f"wms679-review-{suffix}")
    ctx.other_tenant = Tenant(name="Other tenant", slug=f"wms679-foreign-{suffix}")
    session.add_all((ctx.tenant, ctx.other_tenant))
    await session.flush()
    ctx.warehouse = Warehouse(tenant_id=ctx.tenant.id, name="Main", code=f"main-{suffix}")
    ctx.other_warehouse = Warehouse(
        tenant_id=ctx.tenant.id, name="Other", code=f"other-{suffix}"
    )
    ctx.foreign_warehouse = Warehouse(
        tenant_id=ctx.other_tenant.id, name="Foreign", code=f"foreign-{suffix}"
    )
    ctx.seller = Seller(tenant_id=ctx.tenant.id, name="Seller")
    ctx.user = User(
        tenant_id=ctx.tenant.id,
        email=f"wms679-review-{suffix}@example.test",
        password_hash="unused",
        role="fulfillment_admin",
    )
    session.add_all(
        (
            ctx.warehouse,
            ctx.other_warehouse,
            ctx.foreign_warehouse,
            ctx.seller,
            ctx.user,
        )
    )
    await session.flush()
    ctx.product = Product(
        tenant_id=ctx.tenant.id,
        seller_id=ctx.seller.id,
        name="WB product",
        sku_code=f"wb-product-{suffix}",
        wb_barcode="001234",
    )
    ctx.request = MarketplaceUnloadRequest(
        tenant_id=ctx.tenant.id,
        warehouse_id=ctx.warehouse.id,
        seller_id=ctx.seller.id,
        marketplace="wb",
        status="confirmed",
        document_number="WB-679-REVIEW",
    )
    own_box = WarehouseBox(
        tenant_id=ctx.tenant.id,
        warehouse_id=ctx.warehouse.id,
        internal_barcode="WHB-OWN-REVIEW",
    )
    ctx.foreign_box = WarehouseBox(
        tenant_id=ctx.other_tenant.id,
        warehouse_id=ctx.foreign_warehouse.id,
        internal_barcode="WHB-FOREIGN-REVIEW",
    )
    ctx.other_warehouse_box = WarehouseBox(
        tenant_id=ctx.tenant.id,
        warehouse_id=ctx.other_warehouse.id,
        internal_barcode="WHB-OTHER-WAREHOUSE",
    )
    session.add_all((ctx.product, ctx.request, own_box, ctx.foreign_box, ctx.other_warehouse_box))
    await session.flush()
    ctx.box = MarketplaceUnloadBox(
        request_id=ctx.request.id,
        warehouse_box_id=own_box.id,
        box_preset="60_40_40",
    )
    session.add(ctx.box)
    await session.flush()
    session.add(
        MarketplaceUnloadBoxLine(box_id=ctx.box.id, product_id=ctx.product.id, quantity=2)
    )
    await session.commit()
    return ctx


@asynccontextmanager
async def _api(session: AsyncSession, user: User) -> AsyncIterator[AsyncClient]:
    """Exercise route dependencies apart from authentication transport itself."""
    app = create_app()

    async def override_db() -> AsyncIterator[AsyncSession]:
        yield session

    app.dependency_overrides[get_db] = override_db
    app.dependency_overrides[get_current_user] = lambda: user
    app.dependency_overrides[get_effective_seller_id] = lambda: None
    try:
        async with AsyncClient(
            transport=ASGITransport(app=app), base_url="http://test"
        ) as client:
            yield client
    finally:
        app.dependency_overrides.clear()


@pytest.mark.asyncio
async def test_review_f1_actual_export_rejects_box_from_foreign_tenant_or_warehouse(
    db_session: AsyncSession,
) -> None:
    ctx = await _fixture(db_session)
    async with _api(db_session, ctx.user) as client:
        for invalid_box in (ctx.foreign_box, ctx.other_warehouse_box):
            ctx.box.warehouse_box_id = invalid_box.id
            await db_session.commit()

            response = await client.get(_EXPORT_PATH.format(request_id=ctx.request.id))

            assert response.status_code == 422, response.text
            assert str(ctx.box.id) in response.json()["detail"]
            assert invalid_box.internal_barcode.encode() not in response.content


@pytest.mark.asyncio
async def test_review_f2_actual_repeated_export_keeps_xlsx_row_order_when_sql_scan_reverses(
    db_session: AsyncSession,
) -> None:
    ctx = await _fixture(db_session)
    # Four rows expose an unordered SQL scan; all data remains unchanged.
    suffix = uuid.uuid4().hex[:8]
    second_product = Product(
        tenant_id=ctx.tenant.id,
        seller_id=ctx.seller.id,
        name="WB product two",
        sku_code=f"wb-product-two-{suffix}",
        wb_barcode="002345",
    )
    third_product = Product(
        tenant_id=ctx.tenant.id,
        seller_id=ctx.seller.id,
        name="WB product three",
        sku_code=f"wb-product-three-{suffix}",
        wb_barcode="003456",
    )
    fourth_product = Product(
        tenant_id=ctx.tenant.id,
        seller_id=ctx.seller.id,
        name="WB product four",
        sku_code=f"wb-product-four-{suffix}",
        wb_barcode="004567",
    )
    db_session.add_all((second_product, third_product, fourth_product))
    await db_session.flush()
    second_box = WarehouseBox(
        tenant_id=ctx.tenant.id,
        warehouse_id=ctx.warehouse.id,
        internal_barcode="WHB-SECOND-REVIEW",
    )
    db_session.add(second_box)
    await db_session.flush()
    second_unload_box = MarketplaceUnloadBox(
        request_id=ctx.request.id,
        warehouse_box_id=second_box.id,
        box_preset="30_20_30",
    )
    db_session.add(second_unload_box)
    await db_session.flush()
    db_session.add_all(
        (
            MarketplaceUnloadBoxLine(box_id=ctx.box.id, product_id=second_product.id, quantity=3),
            MarketplaceUnloadBoxLine(
                box_id=second_unload_box.id, product_id=third_product.id, quantity=1
            ),
            MarketplaceUnloadBoxLine(
                box_id=second_unload_box.id, product_id=fourth_product.id, quantity=4
            ),
        )
    )
    await db_session.commit()

    async with _api(db_session, ctx.user) as client:
        first = await client.get(_EXPORT_PATH.format(request_id=ctx.request.id))
        assert first.status_code == 200, first.text
        await db_session.execute(text("PRAGMA reverse_unordered_selects = ON"))
        second = await client.get(_EXPORT_PATH.format(request_id=ctx.request.id))
        await db_session.execute(text("PRAGMA reverse_unordered_selects = OFF"))

    assert second.status_code == 200, second.text
    assert _xlsx_rows(second.content) == _xlsx_rows(first.content)


@pytest.mark.asyncio
async def test_review_f3_actual_empty_wb_request_keeps_wb_context_without_false_marketplace_error(
    db_session: AsyncSession,
) -> None:
    ctx = await _fixture(db_session)
    empty_request = MarketplaceUnloadRequest(
        tenant_id=ctx.tenant.id,
        warehouse_id=ctx.warehouse.id,
        seller_id=ctx.seller.id,
        marketplace="wb",
        status="confirmed",
        document_number="WB-679-EMPTY",
    )
    db_session.add(empty_request)
    await db_session.commit()

    async with _api(db_session, ctx.user) as client:
        context = await client.get(f"/operations/marketplace-unload-requests/{empty_request.id}")
        export = await client.get(_EXPORT_PATH.format(request_id=empty_request.id))

    assert context.status_code == 200, context.text
    assert context.json()["marketplace"] == "wb"
    assert context.json()["boxes"] == []
    assert export.status_code == 422, export.text
    detail = export.json()["detail"].lower()
    assert "короб" in detail
    assert "только для отгрузки wb" not in detail
