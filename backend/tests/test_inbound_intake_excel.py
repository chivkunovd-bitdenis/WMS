from __future__ import annotations

import io
import time
import uuid
import zipfile
from typing import Any

import pytest
from httpx import AsyncClient
from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]

from app.db.session import SessionLocal
from app.models.product import Product
from app.models.product_barcode import ProductBarcode
from app.models.product_marketplace_link import ProductMarketplaceLink
from app.services.inbound_intake_excel_service import HEADERS


def xlsx(rows: list[tuple[Any, ...]], headers: tuple[str, ...] = HEADERS) -> bytes:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def malformed_sheet_xml(content: bytes) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content)) as source, zipfile.ZipFile(output, "w") as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                assert b"</sheetData>" in data
                data = data.replace(b"</sheetData>", b"<row>", 1)
            target.writestr(entry, data)
    return output.getvalue()


@pytest.mark.asyncio
async def test_seller_intake_excel_import_is_atomic_scoped_and_repeatable(
    async_client: AsyncClient,
) -> None:
    suffix = str(time.time_ns())
    registered = await async_client.post(
        "/auth/register",
        json={
            "organization_name": "Excel Intake",
            "slug": f"excel-intake-{suffix}",
            "admin_email": f"excel-admin-{suffix}@example.com",
            "password": "password123",
        },
    )
    assert registered.status_code == 200, registered.text
    admin = {"Authorization": f"Bearer {registered.json()['access_token']}"}
    warehouse = await async_client.post(
        "/warehouses", headers=admin, json={"name": "Warehouse", "code": f"ex-{suffix}"}
    )
    assert warehouse.status_code == 200, warehouse.text
    seller1 = await async_client.post("/sellers", headers=admin, json={"name": "First"})
    seller2 = await async_client.post("/sellers", headers=admin, json={"name": "Second"})
    sid1, sid2 = seller1.json()["id"], seller2.json()["id"]
    products: list[str] = []
    for code, barcode, size, seller in [
        ("J-36", "000036", "36", sid1),
        ("J-38", "000038", "38", sid1),
        ("OTHER", "000040", "40", sid2),
    ]:
        response = await async_client.post(
            "/products",
            headers=admin,
            json={
                "name": "Jacket",
                "sku_code": code,
                "wb_barcode": barcode,
                "wb_size": size,
                "seller_id": seller,
                "length_mm": 1,
                "width_mm": 1,
                "height_mm": 1,
            },
        )
        assert response.status_code == 200, response.text
        products.append(response.json()["id"])
    email = f"excel-seller-{suffix}@example.com"
    account = await async_client.post(
        "/auth/seller-accounts",
        headers=admin,
        json={"seller_id": sid1, "email": email, "password": "password123"},
    )
    assert account.status_code == 201, account.text
    login = await async_client.post("/auth/login", json={"email": email, "password": "password123"})
    assert login.status_code == 200, login.text
    seller = {"Authorization": f"Bearer {login.json()['access_token']}"}
    created = await async_client.post(
        "/operations/inbound-intake-requests",
        headers=seller,
        json={"warehouse_id": warehouse.json()["id"], "operation_type": "inbound"},
    )
    assert created.status_code == 201, created.text
    rid = created.json()["id"]
    endpoint = f"/operations/inbound-intake-requests/{rid}/lines/import-excel"

    template = await async_client.get(
        "/operations/inbound-intake-requests/excel-template", headers=seller
    )
    assert template.status_code == 200
    sheet = load_workbook(io.BytesIO(template.content)).active
    assert sheet is not None
    assert tuple(cell.value for cell in sheet[1]) == HEADERS
    assert sheet["C2"].number_format == "@"

    async def upload(content: bytes, name: str = "lines.xlsx") -> Any:
        return await async_client.post(
            endpoint,
            headers=seller,
            files={
                "file": (
                    name,
                    content,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

    valid = xlsx(
        [
            ("Jacket", "J-36", "000036", "36", 2),
            ("Jacket", "J-38", "000038", "38", 5),
            ("Jacket", "J-36", "000036", "36", 3),
        ]
    )
    first = await upload(valid)
    assert first.status_code == 200, first.text
    assert first.json() == {"imported": 2}

    async def quantities() -> dict[str, int]:
        detail = await async_client.get(
            f"/operations/inbound-intake-requests/{rid}", headers=seller
        )
        assert detail.status_code == 200
        return {line["product_id"]: line["expected_qty"] for line in detail.json()["lines"]}

    assert await quantities() == {products[0]: 5, products[1]: 5}
    repeat = await upload(valid)
    assert repeat.status_code == 200
    assert await quantities() == {products[0]: 5, products[1]: 5}

    bad_sku = await upload(
        xlsx(
            [
                ("Jacket", "J-36", "000036", "36", 10),
                ("Unknown", "NOT-FOUND", "000099", "99", 1),
            ]
        )
    )
    assert bad_sku.status_code == 422
    assert bad_sku.json()["detail"]["issues"][0]["row"] == 3
    assert "NOT-FOUND" in str(bad_sku.json())
    assert await quantities() == {products[0]: 5, products[1]: 5}

    other_seller = await upload(xlsx([("Other", "OTHER", "000040", "40", 1)]))
    assert other_seller.status_code == 422
    assert await quantities() == {products[0]: 5, products[1]: 5}

    unknown_barcode = await upload(xlsx([("Unknown", "", "000099", "", 1)]))
    assert unknown_barcode.status_code == 422
    assert unknown_barcode.json()["detail"]["issues"][0]["article"] == "000099"

    for bad in [
        xlsx([("Jacket", "J-36", "000036", "36", 0)]),
        xlsx([("Jacket", "J-36", "000036", "36", 1.5)]),
        xlsx([("Jacket", "J-36", "000036", "38", 1)]),
        xlsx([("Jacket", "J-36", "000036", "36", 1)], headers=("wrong",)),
        b"broken",
        malformed_sheet_xml(valid),
    ]:
        invalid = await upload(bad)
        assert invalid.status_code == 422, invalid.text
        assert await quantities() == {products[0]: 5, products[1]: 5}

    set_one = await upload(xlsx([("Jacket", "J-36", "000036", "36", 7)]))
    assert set_one.status_code == 200
    assert await quantities() == {products[0]: 7, products[1]: 5}

    async with SessionLocal() as session:
        product = await session.get(Product, uuid.UUID(products[0]))
        assert product is not None
        session.add(
            ProductBarcode(
                tenant_id=product.tenant_id,
                seller_id=uuid.UUID(sid1),
                product_id=product.id,
                barcode="000036ALT",
                source="wb",
            )
        )
        session.add(
            ProductMarketplaceLink(
                tenant_id=product.tenant_id,
                seller_id=uuid.UUID(sid1),
                product_id=product.id,
                marketplace="ozon",
                external_barcodes=["000036OZ"],
            )
        )
        await session.commit()
    alias = await upload(xlsx([("Jacket", "J-36", "000036ALT", "36", 9)]))
    assert alias.status_code == 200, alias.text
    assert await quantities() == {products[0]: 9, products[1]: 5}
    marketplace_alias = await upload(xlsx([("Jacket", "J-36", "000036OZ", "36", 11)]))
    assert marketplace_alias.status_code == 200, marketplace_alias.text
    assert await quantities() == {products[0]: 11, products[1]: 5}
