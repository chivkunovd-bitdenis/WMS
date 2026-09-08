from __future__ import annotations

import io
import json
import uuid
from typing import Any

import httpx
import pytest
from openpyxl import load_workbook  # type: ignore[import-untyped]
from sqlalchemy import func, select

from app.db.session import SessionLocal
from app.models.fbs_order import FbsOrder
from app.models.inbound_intake import InboundIntakeLine
from app.models.inventory_movement import InventoryMovement
from app.models.marking_code import MarkingCode, MarkingCodeEvent
from app.models.product import Product
from app.models.seller import Seller
from app.services import fbs_marking_service as wb
from app.services import inbound_intake_service as intake
from app.services import inbound_marking_service as svc
from app.services import marking_code_service as marking
from tests.test_inbound_intake_service_be01 import _auth_ids, _setup_request

CIS = "010460123456789021SERIAL1234567\x1d91KEY1\x1d92SIGNATURE+/="


async def _setup(client: httpx.AsyncClient) -> tuple[uuid.UUID, uuid.UUID, uuid.UUID, uuid.UUID]:
    tenant, user = await _auth_ids(client)
    req_id, product_id = await _setup_request(client, tenant, expected_qty=2)
    async with SessionLocal() as session:
        product = await session.get(Product, product_id)
        assert product is not None
        seller = Seller(tenant_id=tenant, name="Marking seller")
        session.add(seller)
        await session.flush()
        product.seller_id = seller.id
        await session.commit()
        req = await intake.begin_receiving(session, tenant, req_id, actor_user_id=user)
        line_id = req.lines[0].id
        await intake.set_line_actual_qty(session, tenant, req_id, line_id, actual_qty=2)
    return tenant, user, req_id, line_id


@pytest.mark.asyncio
async def test_attach_preserves_full_code_qty_and_idempotency(
    async_client: httpx.AsyncClient,
) -> None:
    tenant, user, req, line = await _setup(async_client)
    async with SessionLocal() as session:
        first = await svc.attach_code(
            session, tenant, req, line_id=line, cis_code=CIS, actor_user_id=user
        )
        again = await svc.attach_code(
            session, tenant, req, line_id=line, cis_code="]d2" + CIS, actor_user_id=user
        )
        assert first == again
        assert first["cis_code"] == CIS and first["cz_status"] == "pending"
        assert await session.scalar(select(func.count(MarkingCodeEvent.id))) == 1
        assert (await session.get(InboundIntakeLine, line)).actual_qty == 2
        assert await session.scalar(select(func.count(InventoryMovement.id))) == 0
        code = await session.get(MarkingCode, uuid.UUID(first["id"]))
        assert code.status == "applied" and code.source == "external_fbs"
        assert await marking.count_available_for_product(session, tenant, code.product_id) == 0
        assert await marking.is_unbound_received_code(session, code)
        with pytest.raises(intake.InboundIntakeError, match="request_not_found"):
            await svc.attach_code(
                session, uuid.uuid4(), req, line_id=line, cis_code=CIS, actor_user_id=user
            )


@pytest.mark.asyncio
async def test_capacity_and_existing_lifecycle_are_not_overwritten(
    async_client: httpx.AsyncClient,
) -> None:
    tenant, user, req, line = await _setup(async_client)
    async with SessionLocal() as session:
        first = await svc.attach_code(
            session, tenant, req, line_id=line, cis_code=CIS, actor_user_id=user
        )
        code = await session.get(MarkingCode, uuid.UUID(first["id"]))
        code.status = "shipped"
        await session.commit()
        # Repeating a receipt scan is read-idempotent even after later lifecycle progress.
        await svc.attach_code(session, tenant, req, line_id=line, cis_code=CIS, actor_user_id=user)
        assert code.status == "shipped"
        assert not await marking.is_unbound_received_code(session, code)
        await svc.attach_code(
            session,
            tenant,
            req,
            line_id=line,
            cis_code=CIS.replace("SERIAL", "SECOND"),
            actor_user_id=user,
        )
        with pytest.raises(intake.InboundIntakeError, match="marking_quantity_exceeded"):
            await svc.attach_code(
                session,
                tenant,
                req,
                line_id=line,
                cis_code=CIS.replace("SERIAL", "THIRDX"),
                actor_user_id=user,
            )
        await session.rollback()
        assert await session.scalar(select(func.count(MarkingCode.id))) == 2


@pytest.mark.parametrize(
    ("payload", "status", "reason"),
    [
        (
            {
                "outerStatus": "INTRODUCED",
                "checkResult": True,
                "codeResolveData": {"verified": True},
            },
            "introduced",
            "введён",
        ),
        ({"outerStatus": "APPLIED", "checkResult": False}, "problem", "не введён"),
        (
            {
                "codeFounded": False,
                "checkResult": False,
                "status": "not_found_dm",
                "codeResolveData": {"valid": True, "verified": False},
            },
            "problem",
            "не найден",
        ),
        (
            {
                "outerStatus": "INTRODUCED",
                "checkResult": True,
                "codeResolveData": {"verified": False},
            },
            "problem",
            "Криптоподпись",
        ),
        ({"outerStatus": "INTRODUCED", "checkResult": False}, "unavailable", "не подтвердил"),
        ({"outerStatus": "NEW_FUTURE_STATUS", "checkResult": True}, "unavailable", "не подтвердил"),
        ([], "unavailable", "не подтвердил"),
    ],
)
def test_check_contract_is_conservative(payload: Any, status: str, reason: str) -> None:
    result = svc.interpret_check(payload)
    assert result["status"] == status
    assert reason in result["reason"]


@pytest.mark.asyncio
@pytest.mark.parametrize("offline", [False, True])
async def test_background_check_is_idempotent_and_has_no_stock_effect(
    async_client: httpx.AsyncClient,
    monkeypatch: pytest.MonkeyPatch,
    offline: bool,
) -> None:
    tenant, user, req, line = await _setup(async_client)
    async with SessionLocal() as session:
        item = await svc.attach_code(
            session, tenant, req, line_id=line, cis_code=CIS, actor_user_id=user
        )
        job = await svc.schedule_check(session, tenant, req)
        assert job is not None
        assert await svc.schedule_check(session, tenant, req) is None
    calls = []

    async def fake_post(self: httpx.AsyncClient, url: str, **kwargs: Any) -> httpx.Response:
        calls.append(kwargs["json"]["code"])
        if offline:
            raise httpx.ConnectError("offline")
        return httpx.Response(
            200,
            request=httpx.Request("POST", url),
            json={"outerStatus": "APPLIED", "checkResult": False},
        )

    monkeypatch.setattr(httpx.AsyncClient, "post", fake_post)
    await svc.run_check_job(job)
    await svc.run_check_job(job)
    assert calls == [CIS]
    async with SessionLocal() as session:
        result = await svc.list_codes(session, tenant, req)
        assert not result["checking"]
        assert result["items"][0]["cz_status"] == ("unavailable" if offline else "problem")
        code = await session.get(MarkingCode, uuid.UUID(item["id"]))
        assert code.status == "applied"
        assert await session.scalar(select(func.count(InventoryMovement.id))) == 0
        assert await session.scalar(select(func.count(MarkingCodeEvent.id))) == 1


def test_xlsx_only_problems_keeps_codes_as_text() -> None:
    items = [
        {"article": "=1+1", "cis_code": CIS, "cz_status": status, "cz_reason": status}
        for status in ["introduced", "pending", "problem", "unavailable"]
    ]
    workbook = load_workbook(io.BytesIO(svc.export_problems(items)))
    sheet = workbook.active
    assert sheet.max_row == 3
    assert sheet["A2"].value == "=1+1" and sheet["A2"].data_type == "s"
    assert sheet["B2"].value.replace("\\u001d", "\x1d") == CIS


@pytest.mark.asyncio
async def test_wb_only_admits_unbound_received_code(async_client: httpx.AsyncClient) -> None:
    tenant, user, req, line = await _setup(async_client)
    async with SessionLocal() as session:
        item = await svc.attach_code(
            session, tenant, req, line_id=line, cis_code=CIS, actor_user_id=user
        )
        code = await session.get(MarkingCode, uuid.UUID(item["id"]))
        order = FbsOrder(tenant_id=tenant, seller_id=code.seller_id, product_id=code.product_id)
        claimed = await wb._claim_pool_code_if_present(
            session, tenant_id=tenant, order=order, cis_raw=CIS
        )
        assert claimed.id == code.id and code.status == "applied"
        event = await session.scalar(select(MarkingCodeEvent))
        original = event.meta_json
        event.meta_json = json.dumps({"source_process": "fbs"})
        await session.flush()
        with pytest.raises(wb.FbsMarkingError, match="duplicate_kiz"):
            await wb._claim_pool_code_if_present(
                session, tenant_id=tenant, order=order, cis_raw=CIS
            )
        event.meta_json = original
        for status in ["reserved", "void", "shipped", "introduced"]:
            code.status = status
            await session.flush()
            assert not await marking.is_unbound_received_code(session, code)


@pytest.mark.asyncio
async def test_failed_check_scheduling_cannot_undo_posting(
    async_client: httpx.AsyncClient,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app.models.user import User

    tenant, user_id, req, line = await _setup(async_client)
    async with SessionLocal() as session:
        await svc.attach_code(
            session, tenant, req, line_id=line, cis_code=CIS, actor_user_id=user_id
        )
        user = await session.get(User, user_id)
        email = user.email
    login = await async_client.post("/auth/login", json={"email": email, "password": "password123"})
    assert login.status_code == 200, login.text
    headers = {"Authorization": f"Bearer {login.json()['access_token']}"}

    async def fail_schedule(*args: Any, **kwargs: Any) -> None:
        raise RuntimeError("simulated scheduling failure")

    monkeypatch.setattr(svc, "schedule_check", fail_schedule)
    response = await async_client.post(
        f"/operations/inbound-intake-requests/{req}/complete-receiving", headers=headers
    )
    assert response.status_code == 200, response.text
    assert response.json()["status"] == "sorting"
    again = await async_client.post(
        f"/operations/inbound-intake-requests/{req}/complete-receiving", headers=headers
    )
    assert again.status_code == 409, again.text
    async with SessionLocal() as session:
        assert (
            await session.scalar(
                select(func.count(InventoryMovement.id)).where(
                    InventoryMovement.inbound_intake_line_id == line
                )
            )
            == 1
        )
        assert (await svc.list_codes(session, tenant, req))["items"][0][
            "cz_status"
        ] == "unavailable"


@pytest.mark.asyncio
@pytest.mark.parametrize("marketplace", ["wb", "ozon"])
async def test_receiving_code_binds_as_external_once_and_cannot_be_printed(
    async_client: httpx.AsyncClient,
    marketplace: str,
) -> None:
    from app.models.packaging_task import PackagingTaskLine
    from app.services.fbs_kiz_service import _prepare_code_for_binding
    from app.services.ozon_kiz_service import OzonKizError, _claim_or_create_code
    from tests.test_fbs_kiz import (
        _create_order,
        _create_supply,
        _register_ff_admin,
        _setup_seller_warehouse,
    )

    headers, suffix = await _register_ff_admin(async_client)
    seller, warehouse, tenant = await _setup_seller_warehouse(async_client, headers, suffix)
    supply = await _create_supply(
        tenant_id=tenant, seller_id=seller, warehouse_id=warehouse, suffix=suffix
    )
    seeded = await _create_order(
        tenant_id=tenant,
        seller_id=seller,
        warehouse_id=warehouse,
        supply_id=supply,
        suffix=suffix,
        wb_order_id=396001,
        sticker_code="RCPT396",
        wb_barcode="RCPT396",
        with_packaging=True,
    )
    async with SessionLocal() as session:
        req = await intake.create_request(session, tenant, warehouse_id=warehouse)
        req_id = req.id
        await intake.add_line(session, tenant, req_id, product_id=seeded.product_id, expected_qty=1)
        req = await intake.get_request(session, tenant, req_id)
        await session.refresh(req, attribute_names=["lines"])
        req.status = "receiving"
        req.lines[0].actual_qty = 1
        await session.commit()
        user = (await async_client.get("/auth/me", headers=headers)).json()
        item = await svc.attach_code(
            session,
            tenant,
            req_id,
            line_id=req.lines[0].id,
            cis_code=CIS,
            actor_user_id=uuid.UUID(user["id"]),
        )
        code = await session.get(MarkingCode, uuid.UUID(item["id"]))
        assert await marking.count_available_for_product(session, tenant, seeded.product_id) == 0
        order = await session.get(FbsOrder, seeded.order_id)
        line = await session.get(PackagingTaskLine, seeded.packaging_task_line_id)
        if marketplace == "wb":
            bound, from_pool = await _prepare_code_for_binding(session, tenant, order, CIS, line)
        else:
            bound, from_pool = await _claim_or_create_code(
                session, order, seeded.product_id, CIS, line
            )
        assert bound.id == code.id and not from_pool
        assert code.status == "applied" and code.packaging_task_line_id == line.id
        await session.commit()
        assert not await marking.is_unbound_received_code(session, code)
        with pytest.raises(OzonKizError, match="duplicate_kiz"):
            await _claim_or_create_code(session, order, seeded.product_id, CIS, line)
        assert await session.scalar(select(func.count(InventoryMovement.id))) == 0


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "source,status,error",
    [
        ("pool", "available", "marking_code_in_pool"),
        ("pool", "printed", "marking_code_in_pool"),
        ("external_fbs", "applied", "marking_code_already_used"),
        ("external_fbs", "shipped", "marking_code_already_used"),
        ("external_fbs", "void", "marking_code_already_used"),
    ],
)
async def test_existing_lifecycle_is_rejected_without_silent_reassignment(
    async_client: httpx.AsyncClient,
    source: str,
    status: str,
    error: str,
) -> None:
    tenant, user, req, line_id = await _setup(async_client)
    async with SessionLocal() as session:
        line = await session.get(InboundIntakeLine, line_id)
        product = await session.get(Product, line.product_id)
        code = MarkingCode(
            tenant_id=tenant,
            seller_id=product.seller_id,
            product_id=product.id,
            cis_code=CIS,
            source=source,
            status=status,
        )
        session.add(code)
        await session.commit()
        with pytest.raises(intake.InboundIntakeError, match=error):
            await svc.attach_code(
                session, tenant, req, line_id=line_id, cis_code=CIS, actor_user_id=user
            )
        assert code.status == status and code.source == source
        assert await session.scalar(select(func.count(MarkingCodeEvent.id))) == 0


@pytest.mark.asyncio
@pytest.mark.parametrize("same_code", [True, False])
async def test_concurrent_receipt_scan_serializes_capacity_and_duplicate(
    async_client: httpx.AsyncClient,
    same_code: bool,
) -> None:
    import asyncio

    from sqlalchemy import text

    from app.db.session import engine

    if engine.dialect.name != "postgresql":
        pytest.skip("PostgreSQL row locks required")
    tenant, user, req, line = await _setup(async_client)
    async with SessionLocal() as setup:
        await intake.set_line_actual_qty(setup, tenant, req, line, actual_qty=1)
    async with SessionLocal() as first, SessionLocal() as second:
        await svc._request(first, tenant, req, lock=True)
        second_pid = await second.scalar(text("SELECT pg_backend_pid()"))

        async def scan_second() -> str:
            try:
                result = await svc.attach_code(
                    second,
                    tenant,
                    req,
                    line_id=line,
                    cis_code=CIS if same_code else CIS.replace("SERIAL", "SECOND"),
                    actor_user_id=user,
                )
                return str(result["id"])
            except intake.InboundIntakeError as exc:
                await second.rollback()
                return exc.code

        pending = asyncio.create_task(scan_second())
        try:
            async with SessionLocal() as observer:
                async with asyncio.timeout(10):
                    while True:
                        waiting = await observer.scalar(
                            text("SELECT wait_event_type FROM pg_stat_activity WHERE pid = :pid"),
                            {"pid": second_pid},
                        )
                        await observer.rollback()
                        if waiting == "Lock":
                            break
                        assert not pending.done(), "Receipt scan bypassed request row lock"
                        await asyncio.sleep(0.01)
            result = await svc.attach_code(
                first, tenant, req, line_id=line, cis_code=CIS, actor_user_id=user
            )
            assert await asyncio.wait_for(pending, 10) == (
                result["id"] if same_code else "marking_quantity_exceeded"
            )
        finally:
            await first.rollback()
            if not pending.done():
                pending.cancel()
            await asyncio.gather(pending, return_exceptions=True)
    async with SessionLocal() as session:
        assert await session.scalar(select(func.count(MarkingCode.id))) == 1
        assert await session.scalar(select(func.count(MarkingCodeEvent.id))) == 1
        assert await session.scalar(select(func.count(InventoryMovement.id))) == 0
