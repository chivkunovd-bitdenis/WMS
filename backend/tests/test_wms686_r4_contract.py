"""WMS-686: контракт тестов FBO, редакция 4 (до реализации).

Источник ожиданий — `docs/requirements/WMS-686-decisions-2026-10-09.md` (решения
D0-D9, главный документ) и `docs/requirements/WMS-686.md`. Имя каждого теста
начинается с `test_r4_`; ссылка в столбце «Тест» — `путь::имя`.

Все сценарии идут через HTTP API, как работает веб: настоящие приёмка, отгрузка,
подбор, короба и «Завершить». Заменена только внешняя граница — фоновая проверка
кодов в Честном знаке после проведения приёмки.

Новые ручки (префикс `/operations/marketplace-unload-requests/{request_id}`):
- POST /marking-codes/scan, GET /marking-codes, DELETE /marking-codes/{id},
  POST /marking-codes/issue;
- POST /boxes/{box_id}/extract-all;
- GET и PUT /pass, GET /pass.xlsx;
- POST /ship принимает `acknowledge_marking` и при нехватке КИЗ отвечает 422
  `{code: "marking_incomplete", items}`;
- POST /boxes/attach при превышении плана отвечает 422
  `{code: "plan_limit_exceeded", message, items}`.

КИЗ строится в формате GS1 DataMatrix, который принимает приёмка
(`inbound_marking_service.normalize_scanned_code`): `01` + GTIN-14 + `21` +
серийный номер + GS + `93` + хвост. GTIN-14 — это `0` + EAN-13 товара.
"""

from __future__ import annotations

import io
import json
import uuid
from typing import Any, cast

import pytest
from httpx import AsyncClient, Response
from inbound_box_intake_helpers import (  # type: ignore[import-not-found]
    fulfill_inbound_via_box_scans,
    post_primary_accept,
    set_planned_boxes,
)
from openpyxl import load_workbook  # type: ignore[import-untyped]
from sqlalchemy import func, select
from test_marketplace_unload_and_discrepancy_acts import (  # type: ignore[import-not-found]
    _patch_mp_planned_date,
    _patch_packaging_instructions,
    _seller_wb_mp_warehouse,
)
from test_marketplace_unload_pick_from_container import (  # type: ignore[import-not-found]
    _loose_balance_id,
)

from app.db.session import SessionLocal
from app.models.inventory_balance import InventoryBalance
from app.models.inventory_movement import InventoryMovement
from app.models.marking_code import MarkingCode
from app.models.storage_location import StorageLocation
from app.services.inbound_marking_service import normalize_scanned_code
from app.services.marking_code_service import extract_gtin_from_cis, is_unbound_received_code
from app.services.sorting_location_service import SORTING_LOCATION_CODE

BASE = "/operations/marketplace-unload-requests"
INBOUND = "/operations/inbound-intake-requests"
GS = "\x1d"


def _ean13(prefix12: str) -> str:
    """EAN-13 с верной контрольной цифрой: ШК товара, из которого строится GTIN КИЗ."""
    total = sum(int(digit) * (3 if index % 2 else 1) for index, digit in enumerate(prefix12))
    return prefix12 + str((10 - total % 10) % 10)


EAN_A = _ean13("460686100001")
EAN_B = _ean13("460686100002")
EAN_C = _ean13("460686100003")
EAN_D = _ean13("460686100004")


def _kiz(ean13: str, serial: str) -> str:
    """КИЗ товара с данным EAN-13; формат тот же, что принимает приёмка."""
    code = f"010{ean13}21{serial}{GS}93{serial[-4:]}"
    assert normalize_scanned_code(code) == code
    assert extract_gtin_from_cis(code) == f"0{ean13}"
    return code


def _long_kiz(ean13: str, serial: str) -> tuple[str, str]:
    """Длинный КИЗ: ключ проверки `91` и криптохвост `92` из 88 символов (> 128 знаков)."""
    tail = (serial * 10)[:88]
    code = f"010{ean13}21{serial}{GS}91EE06{GS}92{tail}"
    assert len(tail) == 88 and len(code) > 128
    assert normalize_scanned_code(code) == code
    assert extract_gtin_from_cis(code) == f"0{ean13}"
    return code, tail


def _head(code: str) -> str:
    """Неизменяемая часть КИЗ (`01…21серийный`) для сравнения с сохранённым кодом."""
    return code.split(GS, 1)[0]


# --------------------------------------------------------------------------- данные


async def _org(
    client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> tuple[dict[str, str], str, str, int]:
    """Новая организация ФФ, склад, селлер с подключённым WB и склад WB для отгрузки."""
    suffix = uuid.uuid4().hex[:10]
    reg = await client.post(
        "/auth/register",
        json={
            "organization_name": "WMS-686 R4",
            "slug": f"wms686r4-{suffix}",
            "admin_email": f"wms686r4-{suffix}@example.com",
            "password": "password123",
        },
    )
    assert reg.status_code in (200, 201), reg.text
    h = {"Authorization": f"Bearer {reg.json()['access_token']}"}
    wh = await client.post(
        "/warehouses", headers=h, json={"name": "W", "code": f"w-{uuid.uuid4().hex[:10]}"}
    )
    assert wh.status_code in (200, 201), wh.text
    sid, wb_wid = await _seller_wb_mp_warehouse(client, h, monkeypatch)
    return h, str(wh.json()["id"]), str(sid), int(wb_wid)


async def _cell(client: AsyncClient, h: dict[str, str], wid: str, code: str) -> str:
    loc = await client.post(f"/warehouses/{wid}/locations", headers=h, json={"code": code})
    assert loc.status_code == 200, loc.text
    return str(loc.json()["id"])


async def _product(
    client: AsyncClient,
    h: dict[str, str],
    sid: str,
    *,
    barcode: str,
    name: str = "Футболка 48",
    honest_sign: bool = False,
    sku: str | None = None,
) -> str:
    created = await client.post(
        "/products",
        headers=h,
        json={
            "name": name,
            "sku_code": sku or f"SKU-{uuid.uuid4().hex[:8]}",
            "length_mm": 1,
            "width_mm": 1,
            "height_mm": 1,
            "seller_id": sid,
            "wb_barcode": barcode,
            "requires_honest_sign": honest_sign,
        },
    )
    assert created.status_code in (200, 201), created.text
    pid = str(created.json()["id"])
    await _patch_packaging_instructions(client, h, pid)
    return pid


async def _receive(
    client: AsyncClient,
    h: dict[str, str],
    *,
    wid: str,
    pid: str,
    qty: int,
    loc_id: str,
    kiz: tuple[str, ...] = (),
) -> str:
    """Настоящая приёмка россыпью в ячейку; КИЗ сканируются в приёмке как в вебе."""
    created = await client.post(INBOUND, headers=h, json={"warehouse_id": wid})
    assert created.status_code == 201, created.text
    rid = str(created.json()["id"])
    line = await client.post(
        f"{INBOUND}/{rid}/lines",
        headers=h,
        json={"product_id": pid, "expected_qty": qty, "storage_location_id": loc_id},
    )
    assert line.status_code == 201, line.text
    await set_planned_boxes(client, INBOUND, rid, h)
    submit = await client.post(f"{INBOUND}/{rid}/submit", headers=h)
    assert submit.status_code == 200, submit.text
    await post_primary_accept(client, INBOUND, rid, h)
    await fulfill_inbound_via_box_scans(client, h, rid, line.json()["sku_code"], qty)
    for code in kiz:
        scanned = await client.post(
            f"{INBOUND}/{rid}/marking-codes/scan",
            headers=h,
            json={"line_id": line.json()["id"], "cis_code": code},
        )
        assert scanned.status_code == 200, scanned.text
    verify = await client.post(f"{INBOUND}/{rid}/verify", headers=h)
    assert verify.status_code == 200, verify.text
    post = await client.post(f"{INBOUND}/{rid}/post", headers=h)
    assert post.status_code == 200, post.text
    return rid


async def _box_in_cell(
    client: AsyncClient,
    h: dict[str, str],
    *,
    wid: str,
    loc_id: str,
    contents: dict[str, int],
) -> tuple[str, str]:
    """Короб склада в ячейке; товар перекладывается в него из россыпи этой ячейки."""
    box = await client.post(f"/warehouses/{wid}/sorting-objects", headers=h, json={"kind": "box"})
    assert box.status_code == 201, box.text
    box_id, box_barcode = str(box.json()["id"]), str(box.json()["barcode"])
    placed = await client.post(
        f"/warehouses/{wid}/map/move",
        headers=h,
        json={"kind": "box", "id": box_id, "to_kind": "cell", "to_id": loc_id, "qty": 1},
    )
    assert placed.status_code == 200, placed.text
    for pid, qty in contents.items():
        moved = await client.post(
            f"/warehouses/{wid}/map/move",
            headers=h,
            json={
                "kind": "product",
                "id": await _loose_balance_id(loc_id, pid),
                "to_kind": "box",
                "to_id": box_id,
                "qty": qty,
            },
        )
        assert moved.status_code == 200, moved.text
    return box_id, box_barcode


async def _shipment(
    client: AsyncClient,
    h: dict[str, str],
    *,
    wid: str,
    sid: str,
    wb_wid: int,
    lines: dict[str, int],
) -> str:
    """Утверждённая отгрузка FBO на WB; возвращает id."""
    created = await client.post(
        BASE, headers=h, json={"warehouse_id": wid, "seller_id": sid, "wb_mp_warehouse_id": wb_wid}
    )
    assert created.status_code == 201, created.text
    mid = str(created.json()["id"])
    for pid, qty in lines.items():
        added = await client.post(
            f"{BASE}/{mid}/lines", headers=h, json={"product_id": pid, "quantity": qty}
        )
        assert added.status_code in (200, 201), added.text
    await _patch_mp_planned_date(client, h, mid)
    submitted = await client.post(f"{BASE}/{mid}/submit", headers=h)
    assert submitted.status_code == 200, submitted.text
    assert submitted.json()["status"] == "confirmed"
    return mid


async def _seller_headers(
    client: AsyncClient, admin_h: dict[str, str], sid: str
) -> dict[str, str]:
    email = f"wms686r4-seller-{uuid.uuid4().hex[:10]}@example.com"
    account = await client.post(
        "/auth/seller-accounts",
        headers=admin_h,
        json={"seller_id": sid, "email": email, "password": "password123"},
    )
    assert account.status_code == 201, account.text
    login = await client.post("/auth/login", json={"email": email, "password": "password123"})
    assert login.status_code == 200, login.text
    return {"Authorization": f"Bearer {login.json()['access_token']}"}


class Ctx:
    """Отгрузка с одним товаром, лежащим россыпью в ячейке."""

    h: dict[str, str]
    wid: str
    sid: str
    wb_wid: int
    loc_id: str
    pid: str
    sku: str
    mid: str


async def _loose(
    client: AsyncClient,
    monkeypatch: pytest.MonkeyPatch,
    *,
    qty: int = 4,
    plan: int = 3,
    honest_sign: bool = True,
) -> Ctx:
    ctx = Ctx()
    ctx.h, ctx.wid, ctx.sid, ctx.wb_wid = await _org(client, monkeypatch)
    ctx.loc_id = await _cell(client, ctx.h, ctx.wid, "A-1-1")
    ctx.sku = f"SKU-{uuid.uuid4().hex[:8]}"
    ctx.pid = await _product(
        client, ctx.h, ctx.sid, barcode=EAN_A, honest_sign=honest_sign, sku=ctx.sku
    )
    await _receive(client, ctx.h, wid=ctx.wid, pid=ctx.pid, qty=qty, loc_id=ctx.loc_id)
    ctx.mid = await _shipment(
        client, ctx.h, wid=ctx.wid, sid=ctx.sid, wb_wid=ctx.wb_wid, lines={ctx.pid: plan}
    )
    return ctx


# ----------------------------------------------------------------------- действия


async def _pick_scan(
    client: AsyncClient,
    h: dict[str, str],
    mid: str,
    barcode: str,
    *,
    pid: str | None = None,
    loc_id: str | None = None,
    box_id: str | None = None,
) -> Response:
    """Скан на вкладке «Подбор» в том виде, в каком его шлёт веб."""
    body: dict[str, Any] = {"barcode": barcode}
    if pid is not None:
        body["product_id"] = pid
    if loc_id is not None:
        body["storage_location_id"] = loc_id
    if box_id is not None:
        body["container_kind"] = "box"
        body["container_id"] = box_id
    return await client.post(f"{BASE}/{mid}/pick/scan", headers=h, json=body)


async def _pick_unit(client: AsyncClient, ctx: Ctx, count: int = 1) -> None:
    """Подобрать `count` штук ШК товара россыпью из ячейки."""
    for _ in range(count):
        scanned = await _pick_scan(
            client, ctx.h, ctx.mid, EAN_A, pid=ctx.pid, loc_id=ctx.loc_id
        )
        assert scanned.status_code == 200, scanned.text
        assert scanned.json()["kind"] == "product"


async def _pick_set(
    client: AsyncClient,
    h: dict[str, str],
    mid: str,
    *,
    pid: str,
    loc_id: str,
    qty: int,
    box_id: str | None = None,
) -> Response:
    body: dict[str, Any] = {"product_id": pid, "storage_location_id": loc_id, "quantity": qty}
    if box_id is not None:
        body["container_kind"] = "box"
        body["container_id"] = box_id
    return await client.post(f"{BASE}/{mid}/pick/set", headers=h, json=body)


async def _kiz_scan(
    client: AsyncClient,
    h: dict[str, str],
    mid: str,
    code: str,
    *,
    pid: str | None = None,
    mutation_id: str | None = None,
) -> Response:
    body: dict[str, Any] = {"code": code}
    if pid is not None:
        body["product_id"] = pid
    if mutation_id is not None:
        body["mutation_id"] = mutation_id
    return await client.post(f"{BASE}/{mid}/marking-codes/scan", headers=h, json=body)


async def _kiz_items(client: AsyncClient, h: dict[str, str], mid: str) -> list[dict[str, Any]]:
    got = await client.get(f"{BASE}/{mid}/marking-codes", headers=h)
    assert got.status_code == 200, got.text
    return cast(list[dict[str, Any]], got.json()["items"])


async def _kiz_heads(client: AsyncClient, h: dict[str, str], mid: str) -> list[str]:
    return sorted(_head(str(item["cis_code"])) for item in await _kiz_items(client, h, mid))


async def _attach(client: AsyncClient, h: dict[str, str], mid: str, barcode: str) -> Response:
    return await client.post(
        f"{BASE}/{mid}/boxes/attach", headers=h, json={"barcode": barcode, "box_preset": "60_40_40"}
    )


async def _ship(
    client: AsyncClient, h: dict[str, str], mid: str, *, acknowledge_marking: bool = False
) -> Response:
    body = {"acknowledge_marking": True} if acknowledge_marking else {}
    try:
        return await client.post(f"{BASE}/{mid}/ship", headers=h, json=body)
    except Exception as exc:  # отказ сервера без HTTP-ответа — тоже провал «Завершить»
        pytest.fail(f"«Завершить» не вернуло ответ: {exc!r}")


async def _detail(client: AsyncClient, h: dict[str, str], mid: str) -> dict[str, Any]:
    got = await client.get(f"{BASE}/{mid}", headers=h)
    assert got.status_code == 200, got.text
    return cast(dict[str, Any], got.json())


async def _picked(client: AsyncClient, h: dict[str, str], mid: str, pid: str) -> int:
    """«Подобрано» по товару на вкладке «Подбор» (не путать с «В коробах»)."""
    options = await client.get(f"{BASE}/{mid}/pick-options", headers=h)
    assert options.status_code == 200, options.text
    return int(next(row for row in options.json() if row["product_id"] == pid)["picked_qty"])


async def _line_id(client: AsyncClient, h: dict[str, str], mid: str, pid: str) -> str:
    detail = await _detail(client, h, mid)
    return str(next(row for row in detail["lines"] if row["product_id"] == pid)["id"])


async def _import_pool(
    client: AsyncClient,
    h: dict[str, str],
    *,
    sid: str,
    pid: str,
    sku: str,
    codes: list[str],
) -> None:
    """Пул кодов товара через существующий импорт «Честного знака»."""
    body = "cis,sku_code\n" + "\n".join(f"{code},{sku}" for code in codes)
    imported = await client.post(
        "/operations/marking-codes/import",
        headers=h,
        data={
            "seller_id": sid,
            "pools_json": json.dumps([{"title": "WMS-686 R4 пул", "product_ids": [pid]}]),
        },
        files=[("files", ("codes.csv", body.encode(), "text/csv"))],
    )
    assert imported.status_code == 200, imported.text
    assert imported.json()["accepted_count"] == len(codes), imported.text


async def _pool_available(client: AsyncClient, h: dict[str, str], sid: str, pid: str) -> int:
    inventory = await client.get(
        "/operations/marking-codes/inventory", headers=h, params={"seller_id": sid}
    )
    assert inventory.status_code == 200, inventory.text
    return int(
        next(row for row in inventory.json()["rows"] if row["product_id"] == pid)[
            "available_count"
        ]
    )


# ------------------------------------------------------------------- остаток в БД


async def _stock(
    pid: str,
    *,
    loc_id: str | None = None,
    container_id: str | None = None,
    loose: bool = False,
) -> int:
    """Остаток товара: весь, по ячейке, по таре или россыпью в ячейке."""
    stmt = select(func.coalesce(func.sum(InventoryBalance.quantity), 0)).where(
        InventoryBalance.product_id == uuid.UUID(pid)
    )
    if loc_id is not None:
        stmt = stmt.where(InventoryBalance.storage_location_id == uuid.UUID(loc_id))
    if container_id is not None:
        stmt = stmt.where(InventoryBalance.container_id == uuid.UUID(container_id))
    if loose:
        stmt = stmt.where(InventoryBalance.container_id.is_(None))
    async with SessionLocal() as session:
        return int(await session.scalar(stmt) or 0)


async def _sorting_stock(wid: str, pid: str) -> int:
    """Сколько штук товара лежит на системной ячейке «Сортировка» склада."""
    stmt = (
        select(func.coalesce(func.sum(InventoryBalance.quantity), 0))
        .join(StorageLocation, StorageLocation.id == InventoryBalance.storage_location_id)
        .where(
            InventoryBalance.product_id == uuid.UUID(pid),
            StorageLocation.warehouse_id == uuid.UUID(wid),
            StorageLocation.code == SORTING_LOCATION_CODE,
        )
    )
    async with SessionLocal() as session:
        return int(await session.scalar(stmt) or 0)


async def _movements(pid: str) -> list[tuple[str, int]]:
    """Все движения остатка товара (тип, изменение) в порядке записи."""
    stmt = (
        select(InventoryMovement.movement_type, InventoryMovement.quantity_delta)
        .where(InventoryMovement.product_id == uuid.UUID(pid))
        .order_by(InventoryMovement.created_at, InventoryMovement.id)
    )
    async with SessionLocal() as session:
        return [(str(kind), int(delta)) for kind, delta in (await session.execute(stmt)).all()]


# =============================================================== перенос короба целиком


@pytest.mark.asyncio
async def test_r4_inb_box_moved_whole_keeps_its_barcode_in_shipment_and_wb_xlsx(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """D4.1, D6.1: короб приёмки INB целиком — тот же ШК в ответе, в отгрузке и в XLSX для WB."""
    h, wid, sid, wb_wid = await _org(async_client, monkeypatch)
    loc_id = await _cell(async_client, h, wid, "A-1-1")
    pid = await _product(async_client, h, sid, barcode=EAN_A)

    created = await async_client.post(INBOUND, headers=h, json={"warehouse_id": wid})
    assert created.status_code == 201, created.text
    rid = str(created.json()["id"])
    line = await async_client.post(
        f"{INBOUND}/{rid}/lines", headers=h, json={"product_id": pid, "expected_qty": 4}
    )
    assert line.status_code == 201, line.text
    await set_planned_boxes(async_client, INBOUND, rid, h)
    submitted = await async_client.post(f"{INBOUND}/{rid}/submit", headers=h)
    assert submitted.status_code == 200, submitted.text
    await post_primary_accept(async_client, INBOUND, rid, h)
    intake = (await async_client.get(f"{INBOUND}/{rid}", headers=h)).json()
    inb_id = str(intake["boxes"][0]["id"])
    inb_barcode = str(intake["boxes"][0]["internal_barcode"])
    assert inb_barcode.startswith("INB-")
    await fulfill_inbound_via_box_scans(async_client, h, rid, line.json()["sku_code"], 4)
    verified = await async_client.post(f"{INBOUND}/{rid}/verify", headers=h)
    assert verified.status_code == 200, verified.text
    putaway = await async_client.post(
        f"{INBOUND}/{rid}/boxes/{inb_id}/putaway", headers=h, json={"storage_location_id": loc_id}
    )
    assert putaway.status_code == 200, putaway.text
    assert await _stock(pid, container_id=inb_id) == 4
    mid = await _shipment(async_client, h, wid=wid, sid=sid, wb_wid=wb_wid, lines={pid: 4})

    attached = await _attach(async_client, h, mid, inb_barcode)
    assert attached.status_code == 201, attached.text
    assert sum(int(row["quantity"]) for row in attached.json()["lines"]) == 4
    assert attached.json()["internal_barcode"] == inb_barcode, "ШК короба приёмки потерян"
    detail = await _detail(async_client, h, mid)
    assert [box["internal_barcode"] for box in detail["boxes"]] == [inb_barcode]

    exported = await async_client.get(f"{BASE}/{mid}/wb-fbw-packaging.xlsx", headers=h)
    assert exported.status_code == 200, exported.text
    sheet = load_workbook(io.BytesIO(exported.content)).active
    assert [tuple(cell.value for cell in row) for row in sheet.iter_rows(min_row=2, max_col=3)] == [
        (EAN_A, 4, inb_barcode)
    ]


@pytest.mark.asyncio
async def test_r4_whole_box_over_remaining_plan_is_refused_with_items_and_changes_nothing(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """D1.7: план 6, подобрано 2, в коробе 5 → отказ целиком с перечнем; данные не тронуты."""
    h, wid, sid, wb_wid = await _org(async_client, monkeypatch)
    loc_id = await _cell(async_client, h, wid, "A-1-1")
    pid = await _product(async_client, h, sid, barcode=EAN_A, name="Футболка 48")
    await _receive(async_client, h, wid=wid, pid=pid, qty=7, loc_id=loc_id)
    k1, k1_barcode = await _box_in_cell(
        async_client, h, wid=wid, loc_id=loc_id, contents={pid: 5}
    )
    mid = await _shipment(async_client, h, wid=wid, sid=sid, wb_wid=wb_wid, lines={pid: 6})
    taken = await _pick_set(async_client, h, mid, pid=pid, loc_id=loc_id, qty=2)
    assert taken.status_code == 200, taken.text
    assert await _picked(async_client, h, mid, pid) == 2
    movements = await _movements(pid)

    for _ in range(2):  # повтор даёт тот же отказ и тоже ничего не меняет
        refused = await _attach(async_client, h, mid, k1_barcode)
        assert refused.status_code == 422, refused.text
        detail = refused.json()["detail"]
        assert isinstance(detail, dict), f"detail должен быть объектом: {detail!r}"
        assert detail["code"] == "plan_limit_exceeded"
        assert isinstance(detail["message"], str) and detail["message"].strip()
        assert len(detail["items"]) == 1
        item = detail["items"][0]
        assert item["product_id"] == pid
        assert item["product_name"] == "Футболка 48"
        assert (item["in_box"], item["remaining"]) == (5, 4)

        assert await _picked(async_client, h, mid, pid) == 2
        assert (await _detail(async_client, h, mid))["boxes"] == []
        assert await _stock(pid, container_id=k1) == 5
        assert await _sorting_stock(wid, pid) == 2
        assert await _stock(pid) == 7
        assert await _movements(pid) == movements


@pytest.mark.asyncio
async def test_r4_mixed_box_with_one_position_over_plan_is_refused_whole(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """D1.7: короб A x3 + B x2 при плане A 3, B 1 — отказ всему коробу, A тоже не уходит."""
    h, wid, sid, wb_wid = await _org(async_client, monkeypatch)
    loc_id = await _cell(async_client, h, wid, "A-1-1")
    pid_a = await _product(async_client, h, sid, barcode=EAN_A, name="Футболка 48")
    pid_b = await _product(async_client, h, sid, barcode=EAN_B, name="Футболка 50")
    await _receive(async_client, h, wid=wid, pid=pid_a, qty=3, loc_id=loc_id)
    await _receive(async_client, h, wid=wid, pid=pid_b, qty=2, loc_id=loc_id)
    box_id, box_barcode = await _box_in_cell(
        async_client, h, wid=wid, loc_id=loc_id, contents={pid_a: 3, pid_b: 2}
    )
    mid = await _shipment(
        async_client, h, wid=wid, sid=sid, wb_wid=wb_wid, lines={pid_a: 3, pid_b: 1}
    )

    refused = await _attach(async_client, h, mid, box_barcode)
    assert refused.status_code == 422, refused.text
    detail = refused.json()["detail"]
    assert isinstance(detail, dict), f"detail должен быть объектом: {detail!r}"
    assert detail["code"] == "plan_limit_exceeded"
    by_product = {str(item["product_id"]): item for item in detail["items"]}
    assert set(by_product) == {pid_b}, detail["items"]
    assert by_product[pid_b]["product_name"] == "Футболка 50"
    assert (by_product[pid_b]["in_box"], by_product[pid_b]["remaining"]) == (2, 1)

    assert (await _detail(async_client, h, mid))["boxes"] == []
    picked_a = await _picked(async_client, h, mid, pid_a)
    picked_b = await _picked(async_client, h, mid, pid_b)
    assert (picked_a, picked_b) == (0, 0)
    assert await _stock(pid_a, container_id=box_id) == 3
    assert await _stock(pid_b, container_id=box_id) == 2
    assert (await _sorting_stock(wid, pid_a), await _sorting_stock(wid, pid_b)) == (0, 0)


@pytest.mark.asyncio
async def test_r4_whole_box_exactly_filling_remaining_plan_still_attaches(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """D1.7, граница: план 7, подобрано 2, в коробе 5 → перенос проходит, подобрано 7."""
    h, wid, sid, wb_wid = await _org(async_client, monkeypatch)
    loc_id = await _cell(async_client, h, wid, "A-1-1")
    pid = await _product(async_client, h, sid, barcode=EAN_A)
    await _receive(async_client, h, wid=wid, pid=pid, qty=7, loc_id=loc_id)
    k1, k1_barcode = await _box_in_cell(
        async_client, h, wid=wid, loc_id=loc_id, contents={pid: 5}
    )
    mid = await _shipment(async_client, h, wid=wid, sid=sid, wb_wid=wb_wid, lines={pid: 7})
    taken = await _pick_set(async_client, h, mid, pid=pid, loc_id=loc_id, qty=2)
    assert taken.status_code == 200, taken.text

    attached = await _attach(async_client, h, mid, k1_barcode)
    assert attached.status_code == 201, attached.text
    assert sum(int(row["quantity"]) for row in attached.json()["lines"]) == 5
    assert await _picked(async_client, h, mid, pid) == 7
    assert await _stock(pid, container_id=k1) == 0
    assert await _sorting_stock(wid, pid) == 7
    assert await _stock(pid) == 7


# ================================================================== подбор: тара


@pytest.mark.asyncio
async def test_r4_pick_scan_of_container_barcode_is_container_even_with_container_selected(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """D1.4: ШК второй тары при выбранной первой и повтор ШК той же тары — kind=container."""
    h, wid, sid, wb_wid = await _org(async_client, monkeypatch)
    loc_id = await _cell(async_client, h, wid, "A-1-1")
    pid = await _product(async_client, h, sid, barcode=EAN_A)
    await _receive(async_client, h, wid=wid, pid=pid, qty=5, loc_id=loc_id)
    k1, k1_barcode = await _box_in_cell(
        async_client, h, wid=wid, loc_id=loc_id, contents={pid: 3}
    )
    k2, k2_barcode = await _box_in_cell(
        async_client, h, wid=wid, loc_id=loc_id, contents={pid: 2}
    )
    mid = await _shipment(async_client, h, wid=wid, sid=sid, wb_wid=wb_wid, lines={pid: 3})

    first = await _pick_scan(async_client, h, mid, k1_barcode)
    assert first.status_code == 200, first.text
    assert (first.json()["kind"], first.json()["container_id"]) == ("container", k1)

    # K1 выбрана — ШК K2 заменяет источник, а не уходит в поиск товара.
    second = await _pick_scan(async_client, h, mid, k2_barcode, loc_id=loc_id, box_id=k1)
    assert second.status_code == 200, f"ШК второй тары не распознан: {second.text}"
    assert (second.json()["kind"], second.json()["container_id"]) == ("container", k2)

    # Повторный ШК той же выбранной тары — тоже тара, а не «штрихкод не найден».
    again = await _pick_scan(async_client, h, mid, k2_barcode, loc_id=loc_id, box_id=k2)
    assert again.status_code == 200, f"повтор ШК тары не распознан: {again.text}"
    assert (again.json()["kind"], again.json()["container_id"]) == ("container", k2)

    # Источник действительно K2: штука берётся из неё, K1 не тронута.
    unit = await _pick_scan(async_client, h, mid, EAN_A, pid=pid, loc_id=loc_id, box_id=k2)
    assert unit.status_code == 200, unit.text
    assert unit.json()["kind"] == "product"
    assert (await _stock(pid, container_id=k2), await _stock(pid, container_id=k1)) == (1, 3)


# =================================================================== КИЗ на строке


@pytest.mark.asyncio
async def test_r4_kiz_scan_links_to_shipment_line_without_new_unit_and_replay_is_already_linked(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """D1.4, D3.0, D3.3: КИЗ после ШК — привязка к строке без +1; повтор; сверх подобранного."""
    ctx = await _loose(async_client, monkeypatch, qty=4, plan=3)
    kiz1, kiz2 = _kiz(EAN_A, "R4SERIAL00001"), _kiz(EAN_A, "R4SERIAL00002")
    line_id = await _line_id(async_client, ctx.h, ctx.mid, ctx.pid)
    await _pick_unit(async_client, ctx)
    movements = await _movements(ctx.pid)

    first = await _kiz_scan(async_client, ctx.h, ctx.mid, kiz1)  # товар определяется по GTIN
    assert first.status_code == 200, first.text
    body = first.json()
    assert body["marking_code_id"]
    assert str(body["cis_code"]).startswith(_head(kiz1))
    assert (body["product_id"], body["line_id"]) == (ctx.pid, line_id)
    assert body["already_linked"] is False
    assert (body["kiz_count"], body["picked_qty"]) == (1, 1)
    assert await _picked(async_client, ctx.h, ctx.mid, ctx.pid) == 1
    assert await _movements(ctx.pid) == movements
    detail_line = (await _detail(async_client, ctx.h, ctx.mid))["lines"][0]
    assert (detail_line["kiz_count"], detail_line["requires_honest_sign"]) == (1, True)

    again = await _kiz_scan(async_client, ctx.h, ctx.mid, kiz1, pid=ctx.pid)
    assert again.status_code == 200, again.text
    assert again.json()["already_linked"] is True
    assert again.json()["marking_code_id"] == body["marking_code_id"]
    assert (again.json()["kiz_count"], again.json()["picked_qty"]) == (1, 1)
    assert await _kiz_heads(async_client, ctx.h, ctx.mid) == [_head(kiz1)]

    # Подобрана одна штука, КИЗ уже есть: второй код — отказ, ничего не меняется.
    extra = await _kiz_scan(async_client, ctx.h, ctx.mid, kiz2, pid=ctx.pid)
    assert extra.status_code == 422, extra.text
    assert extra.json()["detail"] == "marking_quantity_exceeded"
    assert await _kiz_heads(async_client, ctx.h, ctx.mid) == [_head(kiz1)]
    assert await _picked(async_client, ctx.h, ctx.mid, ctx.pid) == 1
    assert await _movements(ctx.pid) == movements

    # Следующая штука подобрана — второй КИЗ принимается и считается.
    await _pick_unit(async_client, ctx)
    second = await _kiz_scan(async_client, ctx.h, ctx.mid, kiz2, pid=ctx.pid)
    assert second.status_code == 200, second.text
    assert (second.json()["kiz_count"], second.json()["picked_qty"]) == (2, 2)
    assert await _kiz_heads(async_client, ctx.h, ctx.mid) == sorted([_head(kiz1), _head(kiz2)])


@pytest.mark.asyncio
async def test_r4_kiz_of_other_product_other_shipment_and_garbage_are_refused(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """D1.4: код другого товара, код другой отгрузки, мусор и неизвестный товар — отказ."""
    ctx = await _loose(async_client, monkeypatch, qty=4, plan=2)
    await _product(async_client, ctx.h, ctx.sid, barcode=EAN_B, name="Футболка 50")
    y = await _shipment(
        async_client, ctx.h, wid=ctx.wid, sid=ctx.sid, wb_wid=ctx.wb_wid, lines={ctx.pid: 2}
    )
    kiz1 = _kiz(EAN_A, "R4SERIAL10001")
    await _pick_unit(async_client, ctx)
    movements = await _movements(ctx.pid)

    other_product = await _kiz_scan(
        async_client, ctx.h, ctx.mid, _kiz(EAN_B, "R4SERIAL10002"), pid=ctx.pid
    )
    assert other_product.status_code == 422, other_product.text
    assert other_product.json()["detail"] == "marking_code_other_product"

    garbage = await _kiz_scan(async_client, ctx.h, ctx.mid, "NOT-A-KIZ", pid=ctx.pid)
    assert garbage.status_code == 422, garbage.text
    assert garbage.json()["detail"] == "marking_code_invalid"

    unknown = await _kiz_scan(async_client, ctx.h, ctx.mid, _kiz(EAN_C, "R4SERIAL10003"))
    assert unknown.status_code == 422, unknown.text
    assert unknown.json()["detail"] == "marking_product_unknown"
    assert await _kiz_items(async_client, ctx.h, ctx.mid) == []
    assert await _movements(ctx.pid) == movements  # отказы ничего не списали и не подобрали

    in_x = await _kiz_scan(async_client, ctx.h, ctx.mid, kiz1, pid=ctx.pid)
    assert in_x.status_code == 200, in_x.text
    unit_y = await _pick_scan(async_client, ctx.h, y, EAN_A, pid=ctx.pid, loc_id=ctx.loc_id)
    assert unit_y.status_code == 200, unit_y.text
    in_y = await _kiz_scan(async_client, ctx.h, y, kiz1, pid=ctx.pid)
    assert in_y.status_code == 422, in_y.text
    assert in_y.json()["detail"] == "marking_code_other_shipment"

    assert await _kiz_items(async_client, ctx.h, y) == []
    assert await _kiz_heads(async_client, ctx.h, ctx.mid) == [_head(kiz1)]
    assert await _picked(async_client, ctx.h, y, ctx.pid) == 1
    assert await _picked(async_client, ctx.h, ctx.mid, ctx.pid) == 1


@pytest.mark.asyncio
async def test_r4_linked_code_is_occupied_for_fbs_unbound_received_check(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """D3.1: код приёмки, привязанный к отгрузке FBO, уже не «свободный принятый» для FBS."""

    async def no_honest_sign_check(job_id: uuid.UUID) -> None:
        del job_id  # внешняя граница: проверка кода в Честном знаке после приёмки

    monkeypatch.setattr(
        "app.services.inbound_marking_service.run_check_job", no_honest_sign_check
    )
    h, wid, sid, wb_wid = await _org(async_client, monkeypatch)
    loc_id = await _cell(async_client, h, wid, "A-1-1")
    pid = await _product(async_client, h, sid, barcode=EAN_A, honest_sign=True)
    kiz1 = _kiz(EAN_A, "R4SERIAL20001")
    rid = await _receive(async_client, h, wid=wid, pid=pid, qty=3, loc_id=loc_id, kiz=(kiz1,))
    intake = await async_client.get(f"{INBOUND}/{rid}", headers=h)
    intake_number = intake.json()["document_number"]

    async def received_code_is_free() -> bool:
        async with SessionLocal() as session:
            code = await session.scalar(
                select(MarkingCode).where(MarkingCode.cis_code.startswith(_head(kiz1)))
            )
            assert code is not None
            return await is_unbound_received_code(session, code)

    assert await received_code_is_free() is True  # принятый в приёмке и ещё нигде не занятый

    mid = await _shipment(async_client, h, wid=wid, sid=sid, wb_wid=wb_wid, lines={pid: 2})
    unit = await _pick_scan(async_client, h, mid, EAN_A, pid=pid, loc_id=loc_id)
    assert unit.status_code == 200, unit.text
    linked = await _kiz_scan(async_client, h, mid, kiz1, pid=pid)
    assert linked.status_code == 200, linked.text
    assert linked.json()["already_linked"] is False
    items = await _kiz_items(async_client, h, mid)
    assert [item["intake_document_number"] for item in items] == [intake_number]

    assert await received_code_is_free() is False, "код привязан к отгрузке, но считается свободным"


@pytest.mark.asyncio
async def test_r4_long_kiz_with_gs_and_88_char_crypto_tail_is_accepted_whole(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """D1.5: КИЗ длиннее 128 знаков (GS, ключ 91, хвост 92 из 88 символов) принимается целиком."""
    ctx = await _loose(async_client, monkeypatch, qty=3, plan=2)
    code, tail = _long_kiz(EAN_A, "R4LONG0000001")
    await _pick_unit(async_client, ctx)

    scanned = await _kiz_scan(async_client, ctx.h, ctx.mid, code, pid=ctx.pid)
    assert scanned.status_code == 200, f"длинный КИЗ отклонён: {scanned.status_code} {scanned.text}"
    assert scanned.json()["already_linked"] is False
    assert len(str(scanned.json()["cis_code"])) > 128
    assert str(scanned.json()["cis_code"]).endswith(tail)
    items = await _kiz_items(async_client, ctx.h, ctx.mid)
    assert len(items) == 1
    assert str(items[0]["cis_code"]).endswith(tail), "криптохвост КИЗ обрезан"

    replay = await _kiz_scan(async_client, ctx.h, ctx.mid, code, pid=ctx.pid)
    assert replay.status_code == 200, replay.text
    assert replay.json()["already_linked"] is True

    # Остальные поля скана тоже принимают длинную строку (до 512 знаков): отказ, если он
    # есть, должен быть по существу, а не «строка слишком длинная» (ошибка валидации).
    on_pick = await _pick_scan(async_client, ctx.h, ctx.mid, code, pid=ctx.pid, loc_id=ctx.loc_id)
    assert not isinstance(on_pick.json().get("detail"), list), on_pick.text
    box = await async_client.post(
        f"{BASE}/{ctx.mid}/boxes", headers=ctx.h, json={"box_preset": "60_40_40"}
    )
    assert box.status_code == 201, box.text
    on_box = await async_client.post(
        f"{BASE}/{ctx.mid}/boxes/{box.json()['id']}/scan", headers=ctx.h, json={"barcode": code}
    )
    assert not isinstance(on_box.json().get("detail"), list), on_box.text


@pytest.mark.asyncio
async def test_r4_delete_marking_code_unlinks_it_and_keeps_picked_units(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """D1.3: ✕ отвязывает ошибочный код; подобранное не меняется; код можно привязать снова."""
    ctx = await _loose(async_client, monkeypatch, qty=3, plan=2)
    kiz1 = _kiz(EAN_A, "R4SERIAL30001")
    await _pick_unit(async_client, ctx)
    first = await _kiz_scan(async_client, ctx.h, ctx.mid, kiz1, pid=ctx.pid)
    assert first.status_code == 200, first.text
    code_id = first.json()["marking_code_id"]
    movements = await _movements(ctx.pid)

    removed = await async_client.delete(
        f"{BASE}/{ctx.mid}/marking-codes/{code_id}", headers=ctx.h
    )
    assert removed.status_code == 200, removed.text
    assert removed.json() == {"removed": True}
    assert await _kiz_items(async_client, ctx.h, ctx.mid) == []
    assert await _picked(async_client, ctx.h, ctx.mid, ctx.pid) == 1
    assert await _movements(ctx.pid) == movements

    repeated = await async_client.delete(
        f"{BASE}/{ctx.mid}/marking-codes/{code_id}", headers=ctx.h
    )
    assert repeated.status_code == 200, repeated.text
    assert repeated.json() == {"removed": False}

    again = await _kiz_scan(async_client, ctx.h, ctx.mid, kiz1, pid=ctx.pid)
    assert again.status_code == 200, again.text
    assert again.json()["already_linked"] is False
    assert again.json()["kiz_count"] == 1


@pytest.mark.asyncio
async def test_r4_cancel_shipment_unlinks_its_kiz_for_reuse(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """D3.5: отмена отгрузки отвязывает её КИЗ; код привязывается в другой отгрузке."""
    ctx = await _loose(async_client, monkeypatch, qty=4, plan=2)
    x = ctx.mid
    y = await _shipment(
        async_client, ctx.h, wid=ctx.wid, sid=ctx.sid, wb_wid=ctx.wb_wid, lines={ctx.pid: 1}
    )
    kiz1 = _kiz(EAN_A, "R4SERIAL40001")
    unit_y = await _pick_scan(async_client, ctx.h, y, EAN_A, pid=ctx.pid, loc_id=ctx.loc_id)
    assert unit_y.status_code == 200, unit_y.text
    in_y = await _kiz_scan(async_client, ctx.h, y, kiz1, pid=ctx.pid)
    assert in_y.status_code == 200, in_y.text

    cancelled = await async_client.post(f"{BASE}/{y}/cancel", headers=ctx.h)
    assert cancelled.status_code == 200, cancelled.text
    assert cancelled.json()["status"] == "cancelled"
    assert await _kiz_items(async_client, ctx.h, y) == []

    await _pick_unit(async_client, ctx)  # теперь в отгрузке X
    in_x = await _kiz_scan(async_client, ctx.h, x, kiz1, pid=ctx.pid)
    assert in_x.status_code == 200, in_x.text
    assert in_x.json()["already_linked"] is False
    assert await _kiz_heads(async_client, ctx.h, x) == [_head(kiz1)]


# ============================================================ выдача из пула, «Допечатать»


@pytest.mark.asyncio
async def test_r4_issue_takes_missing_codes_from_pool_replay_is_same_and_short_pool_is_reported(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """R21: issue выдаёт S-K кодов из пула и привязывает; повтор; shortage; пустой пул."""
    h, wid, sid, wb_wid = await _org(async_client, monkeypatch)
    loc_id = await _cell(async_client, h, wid, "A-1-1")
    sku_a, sku_b = f"SKU-{uuid.uuid4().hex[:8]}", f"SKU-{uuid.uuid4().hex[:8]}"
    sku_d = f"SKU-{uuid.uuid4().hex[:8]}"
    pid = await _product(async_client, h, sid, barcode=EAN_A, honest_sign=True, sku=sku_a)
    pid_short = await _product(
        async_client, h, sid, barcode=EAN_B, name="Мало в пуле", honest_sign=True, sku=sku_b
    )
    pid_empty = await _product(
        async_client, h, sid, barcode=EAN_D, name="Без пула", honest_sign=True, sku=sku_d
    )
    products = ((pid, EAN_A, 3), (pid_short, EAN_B, 2), (pid_empty, EAN_D, 2))
    for product_id, _, count in products:
        await _receive(async_client, h, wid=wid, pid=product_id, qty=count, loc_id=loc_id)
    pool = [_kiz(EAN_A, f"R4POOL0000{n}") for n in range(1, 4)]
    pool_short = [_kiz(EAN_B, "R4POOLSHORT1")]
    await _import_pool(async_client, h, sid=sid, pid=pid, sku=sku_a, codes=pool)
    await _import_pool(async_client, h, sid=sid, pid=pid_short, sku=sku_b, codes=pool_short)
    assert await _pool_available(async_client, h, sid, pid) == 3
    mid = await _shipment(
        async_client, h, wid=wid, sid=sid, wb_wid=wb_wid, lines={p: n for p, _, n in products}
    )
    for product_id, barcode, count in products:
        for _ in range(count):
            unit = await _pick_scan(async_client, h, mid, barcode, pid=product_id, loc_id=loc_id)
            assert unit.status_code == 200, unit.text
    linked = await _kiz_scan(async_client, h, mid, _kiz(EAN_A, "R4EXTERNAL001"), pid=pid)
    assert linked.status_code == 200, linked.text  # K = 1 из S = 3

    async def issue(product_id: str, key: str, **extra: Any) -> Response:
        return await async_client.post(
            f"{BASE}/{mid}/marking-codes/issue",
            headers=h,
            json={"product_id": product_id, "mutation_id": key, **extra},
        )

    key = str(uuid.uuid4())
    issued = await issue(pid, key)
    assert issued.status_code == 200, issued.text
    pool_heads = {_head(code) for code in pool}
    first_heads = sorted(_head(str(item["cis_code"])) for item in issued.json()["items"])
    assert len(first_heads) == 2 and set(first_heads) <= pool_heads
    assert issued.json()["shortage"] == 0
    assert await _pool_available(async_client, h, sid, pid) == 1
    assert len(await _kiz_items(async_client, h, mid)) == 3
    line = (await _detail(async_client, h, mid))["lines"]
    assert next(row for row in line if row["product_id"] == pid)["kiz_count"] == 3

    replay = await issue(pid, key)  # потеря ответа и повтор: те же коды, не новые
    assert replay.status_code == 200, replay.text
    assert sorted(_head(str(item["cis_code"])) for item in replay.json()["items"]) == first_heads
    assert await _pool_available(async_client, h, sid, pid) == 1
    assert len(await _kiz_items(async_client, h, mid)) == 3
    mismatch = await issue(pid, key, quantity=1)
    assert mismatch.status_code == 409, mismatch.text
    assert mismatch.json()["detail"] == "mutation_payload_mismatch"

    nothing = await issue(pid, str(uuid.uuid4()))  # S - K = 0
    assert nothing.status_code == 422, nothing.text
    assert nothing.json()["detail"] == "nothing_to_issue"

    short = await issue(pid_short, str(uuid.uuid4()))  # нужно 2, в пуле 1
    assert short.status_code == 200, short.text
    assert len(short.json()["items"]) == 1 and short.json()["shortage"] == 1
    assert _head(str(short.json()["items"][0]["cis_code"])) == _head(pool_short[0])

    empty = await issue(pid_empty, str(uuid.uuid4()))  # нужно 2, в пуле 0
    assert empty.status_code == 422, empty.text
    assert empty.json()["detail"] == "marking_pool_empty"
    assert len(await _kiz_items(async_client, h, mid)) == 4


@pytest.mark.asyncio
async def test_r4_decreasing_picked_below_kiz_count_unlinks_latest_codes_without_blocking(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """R19: уменьшили «подобрано» ниже K — отвязаны последние коды, K = S, ответ их называет."""
    ctx = await _loose(async_client, monkeypatch, qty=3, plan=3)
    kiz1, kiz2 = _kiz(EAN_A, "R4SERIAL50001"), _kiz(EAN_A, "R4SERIAL50002")
    for code in (kiz1, kiz2):
        await _pick_unit(async_client, ctx)
        linked = await _kiz_scan(async_client, ctx.h, ctx.mid, code, pid=ctx.pid)
        assert linked.status_code == 200, linked.text
    assert await _picked(async_client, ctx.h, ctx.mid, ctx.pid) == 2

    lowered = await _pick_set(
        async_client, ctx.h, ctx.mid, pid=ctx.pid, loc_id=ctx.loc_id, qty=1
    )
    assert lowered.status_code == 200, lowered.text
    assert [_head(str(code)) for code in lowered.json()["unlinked_marking_codes"]] == [_head(kiz2)]
    assert await _picked(async_client, ctx.h, ctx.mid, ctx.pid) == 1
    assert await _kiz_heads(async_client, ctx.h, ctx.mid) == [_head(kiz1)]


@pytest.mark.asyncio
async def test_r4_issue_before_picking_defaults_to_plan_minus_linked(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """R21: S = 0 — выдача по умолчанию P - K (печать до подбора); затем нечего выдавать."""
    ctx = await _loose(async_client, monkeypatch, qty=4, plan=4)
    pool = [_kiz(EAN_A, f"R4PRE0000{n}") for n in range(1, 6)]
    await _import_pool(async_client, ctx.h, sid=ctx.sid, pid=ctx.pid, sku=ctx.sku, codes=pool)
    assert await _picked(async_client, ctx.h, ctx.mid, ctx.pid) == 0

    async def issue(**extra: Any) -> Response:
        return await async_client.post(
            f"{BASE}/{ctx.mid}/marking-codes/issue",
            headers=ctx.h,
            json={"product_id": ctx.pid, "mutation_id": str(uuid.uuid4()), **extra},
        )

    issued = await issue()
    assert issued.status_code == 200, issued.text
    assert len(issued.json()["items"]) == 4 and issued.json()["shortage"] == 0
    assert len(await _kiz_items(async_client, ctx.h, ctx.mid)) == 4
    assert await _pool_available(async_client, ctx.h, ctx.sid, ctx.pid) == 1
    nothing = await issue()
    assert nothing.status_code == 422, nothing.text
    assert nothing.json()["detail"] == "nothing_to_issue"
    assert len(await _kiz_items(async_client, ctx.h, ctx.mid)) == 4


@pytest.mark.asyncio
async def test_r4_issue_ceiling_is_plan_while_scan_stays_limited_by_picked(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """R21, R19: потолок выдачи P - K; скан ограничен S; правка подбора выданное не трогает."""
    ctx = await _loose(async_client, monkeypatch, qty=6, plan=6)
    pool = [_kiz(EAN_A, f"R4CEIL0000{n}") for n in range(1, 7)]
    await _import_pool(async_client, ctx.h, sid=ctx.sid, pid=ctx.pid, sku=ctx.sku, codes=pool)
    for n in (1, 2):
        await _pick_unit(async_client, ctx)
        linked = await _kiz_scan(
            async_client, ctx.h, ctx.mid, _kiz(EAN_A, f"R4CEXT0000{n}"), pid=ctx.pid
        )
        assert linked.status_code == 200, linked.text  # S = 2, K = 2

    async def issue(**extra: Any) -> Response:
        return await async_client.post(
            f"{BASE}/{ctx.mid}/marking-codes/issue",
            headers=ctx.h,
            json={"product_id": ctx.pid, "mutation_id": str(uuid.uuid4()), **extra},
        )

    by_default = await issue()  # S > 0 и S = K: по умолчанию выдавать нечего
    assert by_default.status_code == 422, by_default.text
    assert by_default.json()["detail"] == "nothing_to_issue"
    explicit = await issue(quantity=3)  # явное число выше S, но в пределах P - K = 4
    assert explicit.status_code == 200, explicit.text
    assert len(explicit.json()["items"]) == 3
    assert len(await _kiz_items(async_client, ctx.h, ctx.mid)) == 5
    clipped = await issue(quantity=20)  # выше потолка P - K = 1 — уменьшается до потолка
    assert clipped.status_code == 200, clipped.text
    assert len(clipped.json()["items"]) == 1
    assert len(await _kiz_items(async_client, ctx.h, ctx.mid)) == 6
    full = await issue(quantity=1)  # K = P
    assert full.status_code == 422, full.text
    assert full.json()["detail"] == "nothing_to_issue"

    # Привязка сканом ограничена подобранным: K = 6 > S = 2.
    extra = await _kiz_scan(async_client, ctx.h, ctx.mid, _kiz(EAN_A, "R4CEXT0009"), pid=ctx.pid)
    assert extra.status_code == 422, extra.text
    assert extra.json()["detail"] == "marking_quantity_exceeded"

    # Коды выданы заранее (K > S до правки) — уменьшение подбора их не отвязывает.
    lowered = await _pick_set(async_client, ctx.h, ctx.mid, pid=ctx.pid, loc_id=ctx.loc_id, qty=1)
    assert lowered.status_code == 200, lowered.text
    assert lowered.json().get("unlinked_marking_codes", []) == []
    assert len(await _kiz_items(async_client, ctx.h, ctx.mid)) == 6


# ============================================================== извлечь всё из короба


@pytest.mark.asyncio
async def test_r4_extract_all_empties_box_keeps_picked_and_kiz_and_repeat_changes_nothing(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """R34-R35: «Извлечь всё» — B по коробу 0; подобрано, КИЗ и остаток не меняются; повтор."""
    h, wid, sid, wb_wid = await _org(async_client, monkeypatch)
    loc_id = await _cell(async_client, h, wid, "A-1-1")
    pid = await _product(async_client, h, sid, barcode=EAN_A, honest_sign=True)
    await _receive(async_client, h, wid=wid, pid=pid, qty=5, loc_id=loc_id)
    _, k1_barcode = await _box_in_cell(
        async_client, h, wid=wid, loc_id=loc_id, contents={pid: 3}
    )
    mid = await _shipment(async_client, h, wid=wid, sid=sid, wb_wid=wb_wid, lines={pid: 5})
    attached = await _attach(async_client, h, mid, k1_barcode)
    assert attached.status_code == 201, attached.text
    box_id = str(attached.json()["id"])
    # Второй короб с тем же товаром: извлечение из первого его не затрагивает.
    second = await async_client.post(
        f"{BASE}/{mid}/boxes", headers=h, json={"box_preset": "60_40_40"}
    )
    assert second.status_code == 201, second.text
    second_id = str(second.json()["id"])
    for _ in range(2):
        unit = await _pick_scan(async_client, h, mid, EAN_A, pid=pid, loc_id=loc_id)
        assert unit.status_code == 200, unit.text
    packed = await async_client.post(
        f"{BASE}/{mid}/boxes/{second_id}/scan",
        headers=h,
        json={"barcode": EAN_A, "pick_from_storage": False, "quantity": 2},
    )
    assert packed.status_code == 200, packed.text
    kizs = [_kiz(EAN_A, f"R4EXTRACT000{n}") for n in (1, 2)]
    for code in kizs:
        linked = await _kiz_scan(async_client, h, mid, code, pid=pid)
        assert linked.status_code == 200, linked.text
    assert await _picked(async_client, h, mid, pid) == 5

    async def snapshot() -> tuple[object, ...]:
        detail = await _detail(async_client, h, mid)
        boxes = {
            str(box["id"]): sorted((str(x["product_id"]), int(x["quantity"])) for x in box["lines"])
            for box in detail["boxes"]
        }
        return (
            boxes,
            await _picked(async_client, h, mid, pid),
            await _kiz_heads(async_client, h, mid),
            await _stock(pid),
            await _sorting_stock(wid, pid),
            await _movements(pid),
        )

    before = await snapshot()
    assert before[0] == {box_id: [(pid, 3)], second_id: [(pid, 2)]}
    cleared = await async_client.post(f"{BASE}/{mid}/boxes/{box_id}/extract-all", headers=h)
    assert cleared.status_code == 200, cleared.text
    assert cleared.json()["lines"] == []
    after = await snapshot()
    assert after[0] == {box_id: [], second_id: [(pid, 2)]}, "состав не очищен / задет второй короб"
    assert after[1:] == before[1:], "извлечение изменило подобранное, КИЗ или остаток"
    assert after[2] == sorted(_head(code) for code in kizs)

    repeated = await async_client.post(f"{BASE}/{mid}/boxes/{box_id}/extract-all", headers=h)
    assert repeated.status_code == 200, repeated.text
    assert repeated.json()["lines"] == []
    assert await snapshot() == after


# ============================================================= «Завершить» без отказа 500


async def _marked_and_plain_box(
    client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> tuple[dict[str, str], str, str, str, str]:
    """Отгрузка из одного короба: товар с ЧЗ (3 шт.) и товар без ЧЗ (2 шт.) целиком."""
    h, wid, sid, wb_wid = await _org(client, monkeypatch)
    loc_id = await _cell(client, h, wid, "A-1-1")
    pid = await _product(client, h, sid, barcode=EAN_A, honest_sign=True)
    plain = await _product(client, h, sid, barcode=EAN_B, name="Без ЧЗ", honest_sign=False)
    await _receive(client, h, wid=wid, pid=pid, qty=3, loc_id=loc_id)
    await _receive(client, h, wid=wid, pid=plain, qty=2, loc_id=loc_id)
    _, k1_barcode = await _box_in_cell(
        client, h, wid=wid, loc_id=loc_id, contents={pid: 3, plain: 2}
    )
    mid = await _shipment(client, h, wid=wid, sid=sid, wb_wid=wb_wid, lines={pid: 3, plain: 2})
    attached = await _attach(client, h, mid, k1_barcode)
    assert attached.status_code == 201, attached.text
    return h, mid, pid, plain, k1_barcode


@pytest.mark.asyncio
async def test_r4_ship_with_missing_kiz_is_422_marking_incomplete_and_acknowledge_ships(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """R40: ЧЗ-товар целым коробом без КИЗ — 422 marking_incomplete (не 500); с подтверждением."""
    h, mid, pid, plain, _ = await _marked_and_plain_box(async_client, monkeypatch)
    movements = await _movements(pid)

    refused = await _ship(async_client, h, mid)
    assert refused.status_code == 422, f"{refused.status_code} {refused.text}"
    detail = refused.json()["detail"]
    assert isinstance(detail, dict), f"detail должен быть объектом: {detail!r}"
    assert detail["code"] == "marking_incomplete"
    assert len(detail["items"]) == 1, "в перечне только товар с ЧЗ, у которого K < N"
    item = detail["items"][0]
    assert (item["product_id"], item["linked"], item["quantity"]) == (pid, 0, 3)
    assert item["product_name"] == "Футболка 48"
    assert (await _detail(async_client, h, mid))["status"] != "shipped"
    assert (await _stock(pid), await _stock(plain)) == (3, 2)
    assert await _movements(pid) == movements

    shipped = await _ship(async_client, h, mid, acknowledge_marking=True)
    assert shipped.status_code == 200, shipped.text
    assert shipped.json()["status"] == "shipped"
    assert (await _stock(pid), await _stock(plain)) == (0, 0)
    written_off = [delta for kind, delta in await _movements(pid) if kind == "marketplace_unload"]
    assert written_off == [-3]


@pytest.mark.asyncio
async def test_r4_ship_with_all_kiz_linked_needs_no_acknowledgement(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """R40: K = количеству в коробах — «Завершить» проходит без подтверждения и без упаковки."""
    h, mid, pid, plain, _ = await _marked_and_plain_box(async_client, monkeypatch)
    for n in (1, 2):
        linked = await _kiz_scan(async_client, h, mid, _kiz(EAN_A, f"R4SHIP0000{n}"), pid=pid)
        assert linked.status_code == 200, linked.text

    partial = await _ship(async_client, h, mid)  # K = 2 из 3 — всё ещё неполно
    assert partial.status_code == 422, f"{partial.status_code} {partial.text}"
    assert partial.json()["detail"]["code"] == "marking_incomplete"
    item = partial.json()["detail"]["items"][0]
    assert (item["product_id"], item["linked"], item["quantity"]) == (pid, 2, 3)

    last = await _kiz_scan(async_client, h, mid, _kiz(EAN_A, "R4SHIP00003"), pid=pid)
    assert last.status_code == 200, last.text
    assert last.json()["kiz_count"] == 3
    shipped = await _ship(async_client, h, mid)
    assert shipped.status_code == 200, shipped.text
    assert shipped.json()["status"] == "shipped"
    assert (await _stock(pid), await _stock(plain)) == (0, 0)


# ======================================================================= селлер


@pytest.mark.asyncio
async def test_r4_seller_sees_boxes_and_marking_of_own_shipment_but_cannot_write(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """R22, R47: селлер в GET /{id} видит короба, подбор и КИЗ как ФФ; любая запись — 403."""
    h, wid, sid, wb_wid = await _org(async_client, monkeypatch)
    loc_id = await _cell(async_client, h, wid, "A-1-1")
    pid = await _product(async_client, h, sid, barcode=EAN_A, honest_sign=True)
    await _receive(async_client, h, wid=wid, pid=pid, qty=3, loc_id=loc_id)
    _, k1_barcode = await _box_in_cell(
        async_client, h, wid=wid, loc_id=loc_id, contents={pid: 3}
    )
    mid = await _shipment(async_client, h, wid=wid, sid=sid, wb_wid=wb_wid, lines={pid: 3})
    attached = await _attach(async_client, h, mid, k1_barcode)
    assert attached.status_code == 201, attached.text
    kiz1 = _kiz(EAN_A, "R4SELLER0001")
    linked = await _kiz_scan(async_client, h, mid, kiz1, pid=pid)
    assert linked.status_code == 200, linked.text
    seller_h = await _seller_headers(async_client, h, sid)

    ff_detail = await _detail(async_client, h, mid)
    detail = await _detail(async_client, seller_h, mid)
    assert [box["internal_barcode"] for box in detail["boxes"]] == [k1_barcode]
    assert [
        (str(line["product_id"]), int(line["quantity"])) for line in detail["boxes"][0]["lines"]
    ] == [(pid, 3)]
    assert detail["pick_allocations"] == ff_detail["pick_allocations"] != []
    assert [(row["picked_qty"], row["kiz_count"]) for row in detail["lines"]] == [(3, 1)]
    assert detail.get("linked_packaging_task") is None
    assert await _kiz_heads(async_client, seller_h, mid) == [_head(kiz1)]

    code_id = linked.json()["marking_code_id"]
    writes: list[Response] = [
        await _kiz_scan(async_client, seller_h, mid, _kiz(EAN_A, "R4SELLER0002"), pid=pid),
        await async_client.delete(f"{BASE}/{mid}/marking-codes/{code_id}", headers=seller_h),
        await async_client.post(
            f"{BASE}/{mid}/marking-codes/issue",
            headers=seller_h,
            json={"product_id": pid, "mutation_id": str(uuid.uuid4())},
        ),
        await async_client.post(
            f"{BASE}/{mid}/boxes/{attached.json()['id']}/extract-all", headers=seller_h
        ),
        await _ship(async_client, seller_h, mid, acknowledge_marking=True),
    ]
    assert [response.status_code for response in writes] == [403] * len(writes)
    assert await _kiz_heads(async_client, h, mid) == [_head(kiz1)]
    assert (await _detail(async_client, h, mid))["status"] != "shipped"
    assert [box["lines"] for box in (await _detail(async_client, h, mid))["boxes"]] != [[]]


# ====================================================================== пропуск


PASS_DETAILS = {
    "driver_last_name": "Иванов",
    "driver_first_name": "Пётр",
    "driver_phone": "0079001234567",
    "car_brand": "ГАЗель Next",
    "car_number": "007АВ777",
    "cargo_type": "box",
    "cargo_places_count": 3,
    "arrival_date": "2026-10-12",
}


@pytest.mark.asyncio
async def test_r4_pass_ff_saves_seller_reads_and_xlsx_keeps_text_until_shipped(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """R43-R45: ФФ вносит пропуск; селлер читает, не пишет; XLSX текстом; после ship - 409."""
    ctx = await _loose(async_client, monkeypatch, qty=2, plan=2, honest_sign=False)
    pass_url = f"{BASE}/{ctx.mid}/pass"
    seller_h = await _seller_headers(async_client, ctx.h, ctx.sid)

    empty = await async_client.get(pass_url, headers=ctx.h)
    assert empty.status_code == 200, empty.text
    assert empty.json() == {"pass_details": None, "editable": True}
    not_filled = await async_client.get(f"{pass_url}.xlsx", headers=ctx.h)
    assert not_filled.status_code == 404, not_filled.text
    assert not_filled.json()["detail"] == "pass_not_filled"
    seller_empty = await async_client.get(pass_url, headers=seller_h)
    assert seller_empty.status_code == 200, seller_empty.text
    assert seller_empty.json() == {"pass_details": None, "editable": False}

    bad = await async_client.put(
        pass_url,
        headers=ctx.h,
        json={"pass_details": {**PASS_DETAILS, "car_number": "  "}},
    )
    assert bad.status_code == 422, bad.text
    assert (await async_client.get(pass_url, headers=ctx.h)).json()["pass_details"] is None

    saved = await async_client.put(pass_url, headers=ctx.h, json={"pass_details": PASS_DETAILS})
    assert saved.status_code == 200, saved.text
    assert saved.json()["pass_details"] == PASS_DETAILS
    assert saved.json()["editable"] is True
    assert (await _detail(async_client, ctx.h, ctx.mid))["pass_details"] == PASS_DETAILS

    seen = await async_client.get(pass_url, headers=seller_h)
    assert seen.status_code == 200, seen.text
    assert seen.json() == {"pass_details": PASS_DETAILS, "editable": False}
    assert (await _detail(async_client, seller_h, ctx.mid))["pass_details"] == PASS_DETAILS
    denied = await async_client.put(
        pass_url, headers=seller_h, json={"pass_details": {**PASS_DETAILS, "car_brand": "Чужая"}}
    )
    assert denied.status_code == 403, denied.text
    assert (await async_client.get(pass_url, headers=ctx.h)).json()["pass_details"] == PASS_DETAILS

    async def sheet_texts(headers: dict[str, str]) -> dict[str, str]:
        got = await async_client.get(f"{pass_url}.xlsx", headers=headers)
        assert got.status_code == 200, got.text
        assert got.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        sheet = load_workbook(io.BytesIO(got.content)).active
        assert sheet.title == "Пропуск"
        header, values = list(sheet.iter_rows(min_row=1, max_row=2))
        assert sheet.max_row == 2
        return {
            str(h_cell.value): v_cell.value
            for h_cell, v_cell in zip(header, values, strict=True)
            if v_cell.data_type == "s"
        }

    for headers in (ctx.h, seller_h):
        texts = await sheet_texts(headers)
        assert texts["Госномер"] == PASS_DETAILS["car_number"], "госномер записан не текстом"
        assert texts["Телефон"] == PASS_DETAILS["driver_phone"], "телефон записан не текстом"

    # «Завершить» не требует пропуска; после проведения пропуск только для чтения.
    await _pick_unit(async_client, ctx, count=2)
    box = await async_client.post(
        f"{BASE}/{ctx.mid}/boxes", headers=ctx.h, json={"box_preset": "60_40_40"}
    )
    assert box.status_code == 201, box.text
    packed = await async_client.post(
        f"{BASE}/{ctx.mid}/boxes/{box.json()['id']}/scan",
        headers=ctx.h,
        json={"barcode": EAN_A, "pick_from_storage": False, "quantity": 2},
    )
    assert packed.status_code == 200, packed.text
    shipped = await _ship(async_client, ctx.h, ctx.mid)
    assert shipped.status_code == 200, shipped.text
    assert shipped.json()["status"] == "shipped"

    late = await async_client.put(
        pass_url, headers=ctx.h, json={"pass_details": {**PASS_DETAILS, "car_brand": "Другая"}}
    )
    assert late.status_code == 409, late.text
    assert late.json()["detail"] == "pass_not_editable"
    after = await async_client.get(pass_url, headers=ctx.h)
    assert after.json() == {"pass_details": PASS_DETAILS, "editable": False}
    assert (await sheet_texts(ctx.h))["Госномер"] == PASS_DETAILS["car_number"]


@pytest.mark.asyncio
async def test_r4_shipment_without_pass_ships_and_pass_phone_is_optional_for_wb(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """R43: пустой пропуск не блокирует «Завершить»; для WB телефон необязателен."""
    with_pass = await _loose(async_client, monkeypatch, qty=1, plan=1, honest_sign=False)
    no_phone = {key: value for key, value in PASS_DETAILS.items() if key != "driver_phone"}
    saved = await async_client.put(
        f"{BASE}/{with_pass.mid}/pass", headers=with_pass.h, json={"pass_details": no_phone}
    )
    assert saved.status_code == 200, saved.text
    assert saved.json()["pass_details"].get("driver_phone") in (None, "")

    ctx = await _loose(async_client, monkeypatch, qty=1, plan=1, honest_sign=False)
    assert (await _detail(async_client, ctx.h, ctx.mid))["pass_details"] is None
    await _pick_unit(async_client, ctx)
    box = await async_client.post(
        f"{BASE}/{ctx.mid}/boxes", headers=ctx.h, json={"box_preset": "60_40_40"}
    )
    assert box.status_code == 201, box.text
    packed = await async_client.post(
        f"{BASE}/{ctx.mid}/boxes/{box.json()['id']}/scan",
        headers=ctx.h,
        json={"barcode": EAN_A, "pick_from_storage": False},
    )
    assert packed.status_code == 200, packed.text
    shipped = await _ship(async_client, ctx.h, ctx.mid)
    assert shipped.status_code == 200, shipped.text
    assert shipped.json()["status"] == "shipped"
    assert shipped.json()["pass_details"] is None
