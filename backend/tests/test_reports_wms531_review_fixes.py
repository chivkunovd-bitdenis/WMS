"""WMS-531 · регресс на находки перекрёстного ревью Astra (раунд 1).

Каждый тест воспроизводит ровно тот сценарий, что описан в
`docs/reviews/artifacts/wms-531/review-astra-1.md`, под соответствующим F-id.
F8 — продуктовая развилка, отдельного regression-теста не требует (описана
в требованиях лишь частично); её код зафиксирован тестом на C1 в
`test_reports_wms531.py` (историческое сторно с подтверждённой связью).
"""

from __future__ import annotations

import io
import time
import uuid
from datetime import UTC, datetime, timedelta

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook

from app.db.session import SessionLocal
from app.models.fbs_order import FbsOrder
from app.models.fbs_shipment_reversal_ledger import FbsShipmentReversalLedger
from app.models.inventory_balance import InventoryBalance
from tests.test_reports_wms531 import (
    MSK_NOON,
    PERIOD,
    _movement,
    _org,
    _product,
    _seller,
    _warehouse_location,
)


@pytest.mark.asyncio
async def test_f1_excel_builds_in_roughly_linear_time_at_scale(
    async_client: AsyncClient,
) -> None:
    """F1 (P1): раньше `sheet.max_row` после каждой строки делал сборку файла
    квадратичной (замер ревью: 20.225 с на 4000 строк). Порог ниже — грубый,
    с большим запасом на нагрузку среды, но квадратичная реализация всё
    равно не уложится: по формуле ревью 6000 строк — это уже около 45 с,
    линейная — доли секунды даже при заметно более медленном хосте.
    """
    headers, tenant_id, _user_id = await _org(async_client, name="Wms531f1")
    seller_id = await _seller(async_client, headers, "F1 seller")
    warehouse_id, location_id = await _warehouse_location(async_client, headers, name="f1wh")

    movement_count = 6000
    async with SessionLocal() as session:
        product_id = await _product(
            session, tenant_id=tenant_id, seller_id=seller_id, name="F1 product", sku="F1-1"
        )
        batch = []
        for i in range(movement_count):
            delta = 1 if i % 2 == 0 else -1
            batch.append(_movement(
                tenant_id=tenant_id, product_id=product_id, seller_id=seller_id,
                warehouse_id=warehouse_id, location_id=location_id, quantity_delta=delta,
                movement_type="inbound_intake" if delta > 0 else "marketplace_unload",
                created_at=datetime(2026, 9, 1, 0, 0, tzinfo=UTC) + timedelta(minutes=i),
            ))
        session.add_all(batch)
        session.add(InventoryBalance(
            tenant_id=tenant_id, product_id=product_id, storage_location_id=location_id, quantity=0,
        ))
        await session.commit()

    started = time.monotonic()
    response = await async_client.get(
        "/reports/inventory/export.xlsx", headers=headers,
        params={**PERIOD, "group_by": "product"},
    )
    elapsed = time.monotonic() - started
    assert response.status_code == 200, response.text
    message = f"export took {elapsed:.1f}s for {movement_count} rows — quadratic regression?"
    assert elapsed < 45, message

    workbook = load_workbook(io.BytesIO(response.content))
    sheet = workbook.active
    assert sheet is not None
    # Индекс 4 — «Дата» (с колонкой «Селлер» впереди): заполнена только у
    # строк движения, не у строки товара/селлера/итого.
    movement_rows = [row for row in sheet.iter_rows(min_row=2) if row[4].value is not None]
    assert len(movement_rows) == movement_count


@pytest.mark.asyncio
async def test_f2_second_ozon_position_document_found_when_ledger_predates_period(
    async_client: AsyncClient,
) -> None:
    """F2: журнал готовится ДО фактической передачи (иногда на дни раньше);
    ограничивать поиск позиции по created_at журнала терял вторую и
    следующие позиции. Якорь — created_at движения-первой позиции."""
    headers, tenant_id, _user_id = await _org(async_client, name="Wms531f2")
    seller_id = await _seller(async_client, headers, "F2 seller")
    warehouse_id, location_id = await _warehouse_location(async_client, headers, name="f2wh")
    async with SessionLocal() as session:
        p1 = await _product(
            session, tenant_id=tenant_id, seller_id=seller_id, name="P1", sku="F2-P1"
        )
        p2 = await _product(
            session, tenant_id=tenant_id, seller_id=seller_id, name="P2", sku="F2-P2"
        )
        order = FbsOrder(
            tenant_id=tenant_id, seller_id=seller_id, marketplace="ozon", wb_order_id=-42,
            external_order_id="F2-PROBE", product_id=p1, warehouse_id=warehouse_id,
            created_at_wb=datetime(2026, 8, 1, tzinfo=UTC),
            deadline_at=datetime(2026, 8, 2, tzinfo=UTC),
            mapping_status="mapped", reserve_status="no_stock",
        )
        session.add(order)
        await session.flush()
        # Обе строки списания — 10.09 (внутри периода), а журнал подготовлен
        # 31.08 — за 10 дней до фактического списания.
        m1 = _movement(
            tenant_id=tenant_id, product_id=p1, seller_id=seller_id,
            warehouse_id=warehouse_id, location_id=location_id, quantity_delta=-1,
            movement_type="fbs_shipment", created_at=datetime(2026, 9, 10, MSK_NOON, tzinfo=UTC),
        )
        m2 = _movement(
            tenant_id=tenant_id, product_id=p2, seller_id=seller_id,
            warehouse_id=warehouse_id, location_id=location_id, quantity_delta=-1,
            movement_type="fbs_shipment", created_at=datetime(2026, 9, 10, MSK_NOON, tzinfo=UTC),
        )
        session.add_all([m1, m2])
        await session.flush()
        session.add(FbsShipmentReversalLedger(
            tenant_id=tenant_id, fbs_order_id=order.id, product_id=p1,
            storage_location_id=location_id, quantity=2, shipment_movement_id=m1.id,
            created_at=datetime(2026, 8, 31, tzinfo=UTC),
            ozon_positions_json=[
                {"product_id": str(p1), "movement_id": str(m1.id)},
                {"product_id": str(p2), "movement_id": str(m2.id)},
            ],
        ))
        await session.commit()

    resp = await async_client.get(
        "/reports/inventory/movements", headers=headers,
        params={**PERIOD, "product_id": str(p2)},
    )
    assert resp.status_code == 200, resp.text
    rows = resp.json()["rows"]
    assert len(rows) == 1
    assert rows[0]["document"] is not None, "P2's document was lost (ledger predates period)"
    assert rows[0]["document"]["number"] == "Заказ Ozon №F2-PROBE"


@pytest.mark.asyncio
async def test_f3_split_second_position_inherits_document_via_group(
    async_client: AsyncClient,
) -> None:
    """F3: P2 списан ДВУМЯ строками (M2, M3) с общим transfer_group_id; журнал
    напрямую знает только M1 (P1) и M2 (позиция P2) через JSON — M3 не связан
    ни прямым id, ни позицией, только группой с уже разрешённым M2."""
    headers, tenant_id, _user_id = await _org(async_client, name="Wms531f3")
    seller_id = await _seller(async_client, headers, "F3 seller")
    warehouse_id, location_id = await _warehouse_location(async_client, headers, name="f3wh")
    async with SessionLocal() as session:
        p1 = await _product(
            session, tenant_id=tenant_id, seller_id=seller_id, name="P1", sku="F3-P1"
        )
        p2 = await _product(
            session, tenant_id=tenant_id, seller_id=seller_id, name="P2", sku="F3-P2"
        )
        order = FbsOrder(
            tenant_id=tenant_id, seller_id=seller_id, marketplace="ozon", wb_order_id=-43,
            external_order_id="F3-PROBE", product_id=p1, warehouse_id=warehouse_id,
            created_at_wb=datetime(2026, 9, 1, tzinfo=UTC),
            deadline_at=datetime(2026, 9, 2, tzinfo=UTC),
            mapping_status="mapped", reserve_status="no_stock",
        )
        session.add(order)
        await session.flush()
        m1 = _movement(
            tenant_id=tenant_id, product_id=p1, seller_id=seller_id,
            warehouse_id=warehouse_id, location_id=location_id, quantity_delta=-1,
            movement_type="fbs_shipment",
        )
        group = uuid.uuid4()
        m2 = _movement(
            tenant_id=tenant_id, product_id=p2, seller_id=seller_id,
            warehouse_id=warehouse_id, location_id=location_id, quantity_delta=-1,
            movement_type="fbs_shipment", transfer_group_id=group,
        )
        m3 = _movement(
            tenant_id=tenant_id, product_id=p2, seller_id=seller_id,
            warehouse_id=warehouse_id, location_id=location_id, quantity_delta=-1,
            movement_type="fbs_shipment", transfer_group_id=group,
        )
        session.add_all([m1, m2, m3])
        await session.flush()
        session.add(FbsShipmentReversalLedger(
            tenant_id=tenant_id, fbs_order_id=order.id, product_id=p1,
            storage_location_id=location_id, quantity=2, shipment_movement_id=m1.id,
            ozon_positions_json=[
                {"product_id": str(p1), "movement_id": str(m1.id)},
                {"product_id": str(p2), "movement_id": str(m2.id)},
            ],
        ))
        await session.commit()

    resp = await async_client.get(
        "/reports/inventory/movements", headers=headers,
        params={**PERIOD, "product_id": str(p2)},
    )
    assert resp.status_code == 200, resp.text
    documents_by_id = {row["id"]: row["document"] for row in resp.json()["rows"]}
    assert documents_by_id[str(m2.id)]["number"] == "Заказ Ozon №F3-PROBE"
    assert documents_by_id[str(m3.id)] is not None, "M3 did not inherit the group's document"
    assert documents_by_id[str(m3.id)]["number"] == "Заказ Ozon №F3-PROBE"


@pytest.mark.asyncio
async def test_f4_movements_expansion_respects_search(async_client: AsyncClient) -> None:
    """F4: сводка по виду с поиском отвечала «приход 3» (только ALPHA), а
    раскрытие того же вида без параметра search возвращало оба товара
    (сумма 10) — раскрытие обязано отвечать тому же фильтру, что сводка."""
    headers, tenant_id, _user_id = await _org(async_client, name="Wms531f4")
    seller_id = await _seller(async_client, headers, "F4 seller")
    warehouse_id, location_id = await _warehouse_location(async_client, headers, name="f4wh")
    async with SessionLocal() as session:
        alpha_id = await _product(
            session, tenant_id=tenant_id, seller_id=seller_id, name="ALPHA widget", sku="F4-ALPHA"
        )
        beta_id = await _product(
            session, tenant_id=tenant_id, seller_id=seller_id, name="BETA widget", sku="F4-BETA"
        )
        session.add(_movement(
            tenant_id=tenant_id, product_id=alpha_id, seller_id=seller_id,
            warehouse_id=warehouse_id, location_id=location_id, quantity_delta=3,
            movement_type="inbound_intake",
        ))
        session.add(_movement(
            tenant_id=tenant_id, product_id=beta_id, seller_id=seller_id,
            warehouse_id=warehouse_id, location_id=location_id, quantity_delta=7,
            movement_type="inbound_intake",
        ))
        await session.commit()

    summary = await async_client.get(
        "/reports/inventory", headers=headers,
        params={**PERIOD, "group_by": "operation", "search": "ALPHA"},
    )
    assert summary.status_code == 200, summary.text
    summary_rows = summary.json()["rows"]
    assert summary_rows == [
        {"operation": "Приёмка", "in_qty": 3, "out_qty": 0, "net": 3}
    ]

    expanded = await async_client.get(
        "/reports/inventory/movements", headers=headers,
        params={**PERIOD, "operation": "Приёмка", "search": "ALPHA"},
    )
    assert expanded.status_code == 200, expanded.text
    expanded_rows = expanded.json()["rows"]
    assert sum(row["quantity"] for row in expanded_rows) == 3
    assert all(row["product_name"] == "ALPHA widget" for row in expanded_rows)


@pytest.mark.asyncio
async def test_f5_overview_daily_series_uses_current_seller_not_movement_snapshot(
    async_client: AsyncClient,
) -> None:
    """F5: движение снимает seller_id в момент записи; товар мог с тех пор
    сменить владельца (передача между селлерами). Дневной ряд и сравнение с
    прошлым периодом обязаны фильтроваться по ТЕКУЩЕМУ владельцу товара —
    так же, как уже фильтруются плитки и таблицы (R10)."""
    headers, tenant_id, _user_id = await _org(async_client, name="Wms531f5")
    seller_a = await _seller(async_client, headers, "F5 A")
    seller_b = await _seller(async_client, headers, "F5 B")
    warehouse_id, location_id = await _warehouse_location(async_client, headers, name="f5wh")
    async with SessionLocal() as session:
        # Товар СЕЙЧАС принадлежит A, но движение записано со снимком B —
        # ровно сценарий передачи между селлерами до объединения истории.
        product_id = await _product(
            session, tenant_id=tenant_id, seller_id=seller_a, name="F5 product", sku="F5-1"
        )
        session.add(_movement(
            tenant_id=tenant_id, product_id=product_id, seller_id=seller_b,
            warehouse_id=warehouse_id, location_id=location_id, quantity_delta=7,
            movement_type="inbound_intake",
        ))
        session.add(InventoryBalance(
            tenant_id=tenant_id, product_id=product_id, storage_location_id=location_id, quantity=7,
        ))
        await session.commit()

    overview_a = await async_client.get(
        "/reports/overview", headers=headers, params={**PERIOD, "seller_id": str(seller_a)}
    )
    assert overview_a.status_code == 200, overview_a.text
    body_a = overview_a.json()
    assert body_a["in_qty"] == 7
    daily_in_a = sum(day["in_qty"] for day in body_a["daily"])
    assert daily_in_a == 7, "daily series still scoped by the movement's stale seller snapshot"

    overview_b = await async_client.get(
        "/reports/overview", headers=headers, params={**PERIOD, "seller_id": str(seller_b)}
    )
    body_b = overview_b.json()
    assert body_b["in_qty"] == 0
    assert sum(day["in_qty"] for day in body_b["daily"]) == 0


@pytest.mark.asyncio
async def test_f6_ungrouped_location_movement_flags_incomplete_transfer(
    async_client: AsyncClient,
) -> None:
    """F6: движение расположения без transfer_group_id вообще не может быть
    парой — раньше такие строки не попадали даже в кандидаты проверки и
    расхождение «было + приход - расход != стало» оставалось непомеченным."""
    headers, tenant_id, _user_id = await _org(async_client, name="Wms531f6")
    seller_id = await _seller(async_client, headers, "F6 seller")
    warehouse_id, location_id = await _warehouse_location(async_client, headers, name="f6wh")
    async with SessionLocal() as session:
        product_id = await _product(
            session, tenant_id=tenant_id, seller_id=seller_id, name="F6 product", sku="F6-1"
        )
        session.add(_movement(
            tenant_id=tenant_id, product_id=product_id, seller_id=seller_id,
            warehouse_id=warehouse_id, location_id=location_id, quantity_delta=-5,
            movement_type="stock_transfer_out", transfer_group_id=None,
        ))
        session.add(InventoryBalance(
            tenant_id=tenant_id, product_id=product_id, storage_location_id=location_id, quantity=5,
        ))
        await session.commit()

    resp = await async_client.get(
        "/reports/inventory", headers=headers, params={**PERIOD, "group_by": "product"}
    )
    row = resp.json()["rows"][0]
    assert row["integrity_error"] is True

    overview = await async_client.get("/reports/overview", headers=headers, params=PERIOD)
    assert overview.json()["has_incomplete_transfer"] is True


@pytest.mark.asyncio
async def test_f7_seller_cabinet_excel_has_no_unnamed_totals_row(
    async_client: AsyncClient,
) -> None:
    """F7: include_seller=False раньше всё равно писал безымянную строку
    уровня 0 с четырьмя итоговыми числами перед товаром — в кабинете селлера
    верхним уровнем должен быть сам товар (или вид), без лишней обёртки."""
    headers, tenant_id, _user_id = await _org(async_client, name="Wms531f7")
    seller_id = await _seller(async_client, headers, "F7 seller")
    warehouse_id, location_id = await _warehouse_location(async_client, headers, name="f7wh")
    async with SessionLocal() as session:
        product_id = await _product(
            session, tenant_id=tenant_id, seller_id=seller_id, name="F7 product", sku="F7-1"
        )
        session.add(_movement(
            tenant_id=tenant_id, product_id=product_id, seller_id=seller_id,
            warehouse_id=warehouse_id, location_id=location_id, quantity_delta=9,
            movement_type="inbound_intake",
        ))
        session.add(_movement(
            tenant_id=tenant_id, product_id=product_id, seller_id=seller_id,
            warehouse_id=warehouse_id, location_id=location_id, quantity_delta=-2,
            movement_type="marketplace_unload",
        ))
        session.add(InventoryBalance(
            tenant_id=tenant_id, product_id=product_id, storage_location_id=location_id, quantity=7,
        ))
        await session.commit()

    email = f"f7-seller-{uuid.uuid4().hex[:10]}@example.com"
    created = await async_client.post(
        "/auth/seller-accounts", headers=headers,
        json={"seller_id": str(seller_id), "email": email, "password": "password123"},
    )
    assert created.status_code in (200, 201), created.text
    login = await async_client.post("/auth/login", json={"email": email, "password": "password123"})
    assert login.status_code == 200, login.text
    seller_headers = {"Authorization": f"Bearer {login.json()['access_token']}"}

    response = await async_client.get(
        "/reports/inventory/export.xlsx", headers=seller_headers,
        params={**PERIOD, "group_by": "product"},
    )
    assert response.status_code == 200, response.text
    workbook = load_workbook(io.BytesIO(response.content))
    sheet = workbook.active
    assert sheet is not None

    data_rows = list(sheet.iter_rows(min_row=2, max_row=sheet.max_row - 1))  # без «Итого»
    assert len(data_rows) == 3  # товар + 2 движения, БЕЗ отдельной строки селлера
    product_row = data_rows[0]
    assert product_row[0].value == "F7 product"
    assert sheet.row_dimensions[product_row[0].row].outlineLevel == 0
    assert sheet.row_dimensions[product_row[0].row].hidden in (False, None)
    for movement_row in data_rows[1:]:
        assert sheet.row_dimensions[movement_row[0].row].outlineLevel == 1


@pytest.mark.asyncio
async def test_f9_excel_document_text_carries_the_document_kind_prefix(
    async_client: AsyncClient,
) -> None:
    """F9: ячейка «Документ» несла голый номер («№000042») без слова
    «Приёмка» — в JSON это нормально (фронт форматирует сам), но Excel —
    конечный текст и обязан включать подпись вида документа (R2, R12)."""
    headers, tenant_id, _user_id = await _org(async_client, name="Wms531f9")
    seller_id = await _seller(async_client, headers, "F9 seller")
    warehouse_id, location_id = await _warehouse_location(async_client, headers, name="f9wh")
    async with SessionLocal() as session:
        from app.models.inbound_intake import InboundIntakeLine, InboundIntakeRequest

        product_id = await _product(
            session, tenant_id=tenant_id, seller_id=seller_id, name="F9 product", sku="F9-1"
        )
        intake = InboundIntakeRequest(
            tenant_id=tenant_id, seller_id=seller_id, warehouse_id=warehouse_id,
            status="done", operation_type="inbound", display_number="№000042",
        )
        session.add(intake)
        await session.flush()
        line = InboundIntakeLine(request_id=intake.id, product_id=product_id, expected_qty=9)
        session.add(line)
        await session.flush()
        session.add(_movement(
            tenant_id=tenant_id, product_id=product_id, seller_id=seller_id,
            warehouse_id=warehouse_id, location_id=location_id, quantity_delta=9,
            movement_type="inbound_intake", inbound_intake_line_id=line.id,
        ))
        session.add(InventoryBalance(
            tenant_id=tenant_id, product_id=product_id, storage_location_id=location_id, quantity=9,
        ))
        await session.commit()

    response = await async_client.get(
        "/reports/inventory/export.xlsx", headers=headers,
        params={**PERIOD, "group_by": "product"},
    )
    assert response.status_code == 200, response.text
    workbook = load_workbook(io.BytesIO(response.content))
    sheet = workbook.active
    assert sheet is not None
    document_texts = {
        row[6].value for row in sheet.iter_rows(min_row=2) if row[6].value is not None
    }
    assert "Приёмка №000042" in document_texts
    assert "№000042" not in document_texts


@pytest.mark.asyncio
async def test_f10_text_fields_are_not_written_as_excel_formulas(
    async_client: AsyncClient,
) -> None:
    """F10: товар с именем `=1+1` открывался в Excel как ФОРМУЛА (data_type
    "f"), а не как текст названия. Такое значение обязано остаться строкой."""
    headers, tenant_id, _user_id = await _org(async_client, name="Wms531f10")
    seller_id = await _seller(async_client, headers, "F10 seller")
    warehouse_id, location_id = await _warehouse_location(async_client, headers, name="f10wh")
    async with SessionLocal() as session:
        product_id = await _product(
            session, tenant_id=tenant_id, seller_id=seller_id, name="=1+1", sku="F10-1"
        )
        session.add(_movement(
            tenant_id=tenant_id, product_id=product_id, seller_id=seller_id,
            warehouse_id=warehouse_id, location_id=location_id, quantity_delta=3,
            movement_type="inbound_intake",
        ))
        session.add(InventoryBalance(
            tenant_id=tenant_id, product_id=product_id, storage_location_id=location_id, quantity=3,
        ))
        await session.commit()

    response = await async_client.get(
        "/reports/inventory/export.xlsx", headers=headers,
        params={**PERIOD, "group_by": "product"},
    )
    assert response.status_code == 200, response.text
    workbook = load_workbook(io.BytesIO(response.content))
    sheet = workbook.active
    assert sheet is not None
    product_name_cells = [
        cell for row in sheet.iter_rows(min_row=2) for cell in row if cell.value == "=1+1"
    ]
    assert product_name_cells, "product name cell not found by literal value"
    for cell in product_name_cells:
        assert cell.data_type == "s", f"cell {cell.coordinate} became a formula, not text"
