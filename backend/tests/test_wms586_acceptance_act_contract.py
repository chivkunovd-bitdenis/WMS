"""WMS-586 A1-A7/A10: real API, persisted fixtures, both file formats."""

from __future__ import annotations

import asyncio
import io
import re
import uuid
from datetime import UTC, date, datetime
from urllib.parse import unquote

import fitz
import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import select

from app.core.roles import FULFILLMENT_SELLER
from app.db.session import SessionLocal
from app.models import Base
from app.models.inbound_intake import InboundIntakeLine, InboundIntakeRequest
from app.models.product import Product
from app.models.user import User

BASE = "/operations/inbound-intake-requests"
MIMES = {
    "pdf": "application/pdf",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
HEADERS = ["№", "Товар", "Артикул продавца", "SKU", "ШК", "План", "Факт", "Расхождение"]


async def seed(
    client: AsyncClient,
    values=None,
    *,
    status="sorting",
    number="№000586",
    created="2026-09-29T12:00:00+00:00",
    operation="inbound",
):
    registration = await client.post(
        "/auth/register",
        json={
            "organization_name": "WMS586 contract",
            "slug": f"act-{uuid.uuid4().hex}",
            "admin_email": f"act-{uuid.uuid4().hex}@example.com",
            "password": "password123",
        },
    )
    assert registration.status_code == 200
    headers = {"Authorization": f"Bearer {registration.json()['access_token']}"}
    seller = await client.post("/sellers", headers=headers, json={"name": "ИП Контракт"})
    assert seller.status_code == 201, seller.text
    warehouses = (await client.get("/warehouses", headers=headers)).json()
    warehouse_id = (
        warehouses[0]["id"]
        if warehouses
        else (
            await client.post(
                "/warehouses", headers=headers, json={"name": "Склад", "code": "MAIN"}
            )
        ).json()["id"]
    )
    created_response = await client.post(
        BASE, headers=headers, json={"warehouse_id": warehouse_id, "seller_id": seller.json()["id"]}
    )
    assert created_response.status_code == 201, created_response.text
    request_id = uuid.UUID(created_response.json()["id"])
    values = [(10, 8), (5, 7), (4, 4), (0, 3)] if values is None else values
    expected = []
    async with SessionLocal() as session:
        req = await session.get(InboundIntakeRequest, request_id)
        req.status = status
        req.operation_type = operation
        req.display_number = number
        req.document_number = "DOC-FALLBACK"
        req.created_at = datetime.fromisoformat(created)
        req.planned_delivery_date = date(2026, 10, 2)
        for i, (plan, fact) in enumerate(values):
            name = "=1+1 худи" if i == 0 else f"Товар-{i:03d} длинное кириллическое название"
            if len(values) > 50:
                name = (name + " зимняя коллекция утеплённая хлопковая одежда " * 3).strip()
            sku, vendor, barcode = f"00{i:04d}", f"000АРТ-{i:03d}", f"046000000{i:04d}"
            if len(values) > 50:
                vendor += "-зимняя-коллекция-длинный-артикул-селлера"
            product = Product(
                tenant_id=req.tenant_id,
                seller_id=req.seller_id,
                name=name,
                sku_code=sku,
                wb_vendor_code=vendor,
                wb_barcode=barcode,
            )
            session.add(product)
            await session.flush()
            session.add(
                InboundIntakeLine(
                    request_id=req.id,
                    product_id=product.id,
                    expected_qty=plan,
                    actual_qty=fact,
                    posted_qty=99,
                    added_by_fulfillment=plan == 0,
                )
            )
            expected.append(
                [i + 1, name, vendor, sku, barcode, plan, fact or 0, (fact or 0) - plan]
            )
        await session.commit()
    return headers, request_id, expected


async def download(client, headers, request_id, fmt):
    response = await client.get(f"{BASE}/{request_id}/acceptance-act.{fmt}", headers=headers)
    assert response.status_code == 200, (fmt, response.status_code, response.text[:200])
    assert response.headers["content-type"].startswith(MIMES[fmt])
    disposition = unquote(response.headers["content-disposition"])
    assert "attachment" in disposition and f".{fmt}" in disposition
    assert "000586" in disposition and "29.09.2026" in disposition
    return response.content


def workbook(content):
    sheet = load_workbook(io.BytesIO(content)).active
    assert [cell.value for cell in sheet[3]] == HEADERS
    assert all(cell.value is None for cell in sheet[2])
    return sheet


def pdf(content):
    assert content.startswith(b"%PDF-"), "must be an actual PDF, not an HTML/JSON error"
    doc = fitz.open(stream=content, filetype="pdf")
    assert len(doc) > 0
    assert all(page.get_text().strip() for page in doc), "empty/trailing PDF page"
    return doc


def assert_pdf_rows(doc, expected):
    """Locate numeric columns from headers, then read each row at its ordinal baseline.

    This checks numbers against source rows, not merely against the other export.
    Wrapped product text is checked separately, without requiring a particular layout.
    """
    text = re.sub(r"\s+", "", "\n".join(page.get_text() for page in doc))
    name_positions = []
    for row in expected:
        name = re.sub(r"\s+", "", str(row[1]))
        assert text.count(name) == 1, (row[0], "missing or duplicated product name")
        name_positions.append(text.index(name))
        for value in row[1:5]:
            assert re.sub(r"\s+", "", str(value)) in text, (row[0], value)
    assert name_positions == sorted(name_positions), "product names must follow source row order"
    for column in (2, 4):
        positions = [
            text.index(re.sub(r"\s+", "", str(row[column]))) for row in expected if row[column]
        ]
        assert positions == sorted(positions), (
            column,
            "product identifiers must follow source rows",
        )
    found = {}
    totals = []
    red_diffs = set()
    column_headings = {}
    for page in doc:
        words = page.get_text("words")
        headings = {str(w[4]): w for w in words if w[4] in {"Товар", "План", "Факт", "Расхождение"}}
        if set(headings) == {"Товар", "План", "Факт", "Расхождение"}:
            column_headings = headings
        assert column_headings, "table headers must identify numeric columns"
        headings = column_headings
        left = headings["Товар"][0]
        centers = [(headings[k][0] + headings[k][2]) / 2 for k in ("План", "Факт", "Расхождение")]
        boundaries = [(centers[0] + centers[1]) / 2, (centers[1] + centers[2]) / 2]
        for w in words:
            ordinal = str(w[4])
            is_total = ordinal == "Итого"
            if not is_total and not (w[0] < left and ordinal.isdigit()):
                continue
            y = (w[1] + w[3]) / 2
            cells = [[], [], []]
            for n in words:
                if abs((n[1] + n[3]) / 2 - y) > 3 or n[0] < centers[0] - 25:
                    continue
                raw = str(n[4]).replace("\u2212", "-").replace("\u2013", "-")
                if re.fullmatch(r"[+-]?\d+", raw):
                    x = (n[0] + n[2]) / 2
                    col = 0 if x < boundaries[0] else 1 if x < boundaries[1] else 2
                    cells[col].append(raw)
            assert all(len(c) == 1 for c in cells), (ordinal, cells)
            numbers = [int(c[0]) for c in cells]
            if is_total:
                totals.append(numbers)
            else:
                index = int(ordinal)
                assert index not in found, f"duplicate row {index}"
                assert 1 <= index <= len(expected), (index, "unexpected product row")
                row_words = [str(n[4]) for n in words if abs((n[1] + n[3]) / 2 - y) < 3]
                assert str(expected[index - 1][3]) in row_words, (index, "SKU in wrong row")
                found[index] = numbers
                if numbers[2] != 0:
                    assert cells[2][0].startswith(("+", "-")), "signed nonzero discrepancy"
                    for block in page.get_text("dict")["blocks"]:
                        for line in block.get("lines", []):
                            for span in line["spans"]:
                                if (
                                    span["bbox"][0] > boundaries[1]
                                    and abs((span["bbox"][1] + span["bbox"][3]) / 2 - y) < 4
                                ):
                                    color = int(span["color"])
                                    if (color >> 16 & 255) > (color >> 8 & 255) + 40:
                                        red_diffs.add(index)
    assert found == {row[0]: row[5:8] for row in expected}
    sums = [sum(row[i] for row in expected) for i in (5, 6, 7)]
    assert totals == [sums]
    assert red_diffs == {row[0] for row in expected if row[7] != 0}


@pytest.mark.asyncio
@pytest.mark.parametrize("fmt", ["xlsx", "pdf"])
async def test_plan_fact_text_and_totals(async_client, fmt):
    headers, rid, expected = await seed(async_client)
    content = await download(async_client, headers, rid, fmt)
    if fmt == "xlsx":
        sheet = workbook(content)
        assert sheet["A1"].value == "Приёмка №000586 от 29.09.2026"
        assert [
            list(row) for row in sheet.iter_rows(min_row=4, max_row=7, values_only=True)
        ] == expected
        assert [cell.value for cell in sheet[8]][5:] == [19, 22, 3]
        for row in range(4, 8):
            assert all(sheet.cell(row, col).data_type == "s" for col in range(2, 6))
            assert sheet.cell(row, 8).number_format == "+0;-0;0"
            if sheet.cell(row, 8).value:
                assert sheet.cell(row, 8).font.color.rgb.endswith("B00020")
    else:
        doc = pdf(content)
        assert "Приёмка №000586 от 29.09.2026" in " ".join(doc[0].get_text().split())
        assert_pdf_rows(doc, expected)


@pytest.mark.asyncio
@pytest.mark.parametrize("fmt", ["xlsx", "pdf"])
async def test_zero_none_and_alternative_counters(async_client, fmt):
    headers, rid, expected = await seed(async_client, [(10, 0), (5, None), (4, 2)])
    # posted_qty is 99 on all persisted rows. A container/effective count also
    # differs: product amount in a box must not replace the received actual_qty.
    from app.models.inbound_intake import InboundIntakeBox, InboundIntakeBoxLine

    async with SessionLocal() as session:
        req = await session.get(InboundIntakeRequest, rid)
        line = (
            await session.scalars(
                select(InboundIntakeLine).where(InboundIntakeLine.request_id == rid)
            )
        ).first()
        box = InboundIntakeBox(
            tenant_id=req.tenant_id,
            request_id=rid,
            box_number=1,
            internal_barcode="INB-COUNTER-CONTRACT",
        )
        session.add(box)
        await session.flush()
        session.add(InboundIntakeBoxLine(box_id=box.id, product_id=line.product_id, quantity=20))
        product = await session.get(Product, line.product_id)
        product.wb_vendor_code = None
        product.wb_barcode = None
        expected[0][2], expected[0][4] = "", ""
        await session.commit()
    content = await download(async_client, headers, rid, fmt)
    if fmt == "pdf":
        assert_pdf_rows(pdf(content), expected)
    else:
        actual = [
            list(row) for row in workbook(content).iter_rows(min_row=4, max_row=6, values_only=True)
        ]
        actual[0][2], actual[0][4] = actual[0][2] or "", actual[0][4] or ""
        assert actual == expected


@pytest.mark.asyncio
@pytest.mark.parametrize("fmt", ["xlsx", "pdf"])
@pytest.mark.parametrize(
    "status", ["sorting", "verified", "done", "posted", "draft", "submitted", "receiving"]
)
async def test_closed_statuses_and_empty_act(async_client, fmt, status):
    headers, rid, _ = await seed(async_client, [], status=status)
    response = await async_client.get(f"{BASE}/{rid}/acceptance-act.{fmt}", headers=headers)
    if status in {"draft", "submitted", "receiving"}:
        assert response.status_code == 409
        assert response.json() == {"detail": "reception_not_closed"}
    else:
        assert response.status_code == 200, response.text
        if fmt == "xlsx":
            assert [c.value for c in workbook(response.content)[4]][5:] == [0, 0, 0]
        else:
            assert_pdf_rows(pdf(response.content), [])


@pytest.mark.asyncio
@pytest.mark.parametrize("fmt", ["xlsx", "pdf"])
@pytest.mark.parametrize(
    "number,document,operation,created,title",
    [
        (
            "№000586",
            "other",
            "inbound",
            datetime(2026, 9, 28, 21, 30, tzinfo=UTC),
            "Приёмка №000586 от 29.09.2026",
        ),
        (
            "№000999",
            "other",
            "inbound",
            datetime(2026, 9, 20, 12, tzinfo=UTC),
            "Приёмка №000999 от 20.09.2026",
        ),
        (
            None,
            "RET-42",
            "return",
            datetime(2026, 9, 28, 12, tzinfo=UTC),
            "Возврат RET-42 от 28.09.2026",
        ),
        (None, None, "inbound", datetime(2026, 9, 28, 12, tzinfo=UTC), None),
    ],
    ids=["moscow_boundary", "other_document", "return_document_number", "uuid_fallback"],
)
async def test_identity_dates_return_and_fallbacks(
    async_client, fmt, number, document, operation, created, title
):
    headers_a, rid_a, _ = await seed(async_client, [])
    headers, rid, _ = await seed(async_client, [])
    title = title or f"Приёмка {rid} от 28.09.2026"
    async with SessionLocal() as session:
        req = await session.get(InboundIntakeRequest, rid)
        req.display_number, req.document_number = number, document
        req.operation_type, req.created_at = operation, created
        await session.commit()
    for selected_headers, selected_rid, expected_title in [
        (headers_a, rid_a, "Приёмка №000586 от 29.09.2026"),
        (headers, rid, title),
        (headers_a, rid_a, "Приёмка №000586 от 29.09.2026"),
    ]:
        response = await async_client.get(
            f"{BASE}/{selected_rid}/acceptance-act.{fmt}", headers=selected_headers
        )
        assert response.status_code == 200, response.text
        actual = (
            workbook(response.content)["A1"].value
            if fmt == "xlsx"
            else " ".join(pdf(response.content)[0].get_text().split())
        )
        assert expected_title in actual
        disposition = unquote(response.headers["content-disposition"])
        assert expected_title.split(" от ")[1] in disposition


@pytest.mark.asyncio
@pytest.mark.parametrize("fmt", ["xlsx", "pdf"])
async def test_tenant_seller_and_missing_scope(async_client, fmt):
    headers_a, rid_a, _ = await seed(async_client)
    headers_b, rid_b, _ = await seed(async_client)
    for headers, rid in [(headers_a, rid_b), (headers_b, rid_a), (headers_a, uuid.uuid4())]:
        response = await async_client.get(f"{BASE}/{rid}/acceptance-act.{fmt}", headers=headers)
        assert response.status_code == 404
        assert response.json() == {"detail": "request_not_found"}
    seller_b = await async_client.post(
        "/sellers", headers=headers_a, json={"name": "Другой селлер"}
    )
    me = (await async_client.get("/auth/me", headers=headers_a)).json()
    async with SessionLocal() as session:
        user = await session.get(User, uuid.UUID(me["id"]))
        req = await session.get(InboundIntakeRequest, rid_a)
        own_seller = req.seller_id
        user.role, user.seller_id = FULFILLMENT_SELLER, uuid.UUID(seller_b.json()["id"])
        await session.commit()
    response = await async_client.get(f"{BASE}/{rid_a}/acceptance-act.{fmt}", headers=headers_a)
    assert response.status_code == 404
    assert response.json() == {"detail": "request_not_found"}
    async with SessionLocal() as session:
        user = await session.get(User, uuid.UUID(me["id"]))
        user.seller_id = own_seller
        await session.commit()
    assert (
        await async_client.get(f"{BASE}/{rid_a}/acceptance-act.{fmt}", headers=headers_a)
    ).status_code == 200


async def snapshot():
    # All persisted tables, including status/lines/stock/reserves/billing/acts.
    async with SessionLocal() as session:
        return {
            table.name: sorted(
                repr(tuple(row)) for row in (await session.execute(select(table))).all()
            )
            for table in Base.metadata.sorted_tables
        }


@pytest.mark.asyncio
@pytest.mark.parametrize("fmt", ["xlsx", "pdf"])
async def test_repeated_concurrent_download_is_read_only(async_client, fmt):
    headers, rid, _ = await seed(async_client)
    before = await snapshot()
    contents = [await download(async_client, headers, rid, fmt)]
    contents += await asyncio.gather(*[download(async_client, headers, rid, fmt) for _ in range(2)])
    contents.append(await download(async_client, headers, rid, fmt))
    assert await snapshot() == before
    if fmt == "xlsx":
        normalized = [list(workbook(c).values) for c in contents]
    else:
        normalized = [[p.get_text() for p in pdf(c)] for c in contents]
    assert all(value == normalized[0] for value in normalized)


@pytest.mark.asyncio
@pytest.mark.parametrize("fmt", ["xlsx", "pdf"])
async def test_321_long_rows_no_loss_or_empty_tail(async_client, fmt):
    values = [(i % 11, i % 13) for i in range(321)]
    headers, rid, expected = await seed(async_client, values)
    content = await download(async_client, headers, rid, fmt)
    if fmt == "xlsx":
        sheet = workbook(content)
        assert sheet.max_row == 325
        assert [
            list(row) for row in sheet.iter_rows(min_row=4, max_row=324, values_only=True)
        ] == expected
        assert [c.value for c in sheet[325]][5:] == [sum(r[i] for r in expected) for i in (5, 6, 7)]
    else:
        doc = pdf(content)
        assert len(doc) > 1
        assert "Итого" in doc[-1].get_text()
        assert_pdf_rows(doc, expected)
