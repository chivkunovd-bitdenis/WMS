"""WMS-679: contract for the local FBW packaging XLSX before implementation.

The module under test is deliberately a local read-only boundary.  It does not
describe or call a Wildberries API: WB packaging write and WB label data remain
unconfirmed dependencies of this task.
"""

from __future__ import annotations

import uuid
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook  # type: ignore[import-untyped]
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.inventory_balance import InventoryBalance
from app.models.marketplace_unload import (
    MarketplaceUnloadBox,
    MarketplaceUnloadBoxLine,
    MarketplaceUnloadRequest,
)
from app.models.marketplace_unload_reservation import MarketplaceUnloadReservation
from app.models.product import Product
from app.models.product_barcode import ProductBarcode
from app.models.seller import Seller
from app.models.storage_location import StorageLocation
from app.models.tenant import Tenant
from app.models.warehouse import Warehouse
from app.models.warehouse_box import WarehouseBox


# This import is intentionally inside the helper: today it is the expected RED
# failure (the export does not exist).  The implementation must expose this
# local service contract; it must not replace it with an unverified WB call.
async def _export(session: AsyncSession, tenant_id: uuid.UUID, request_id: uuid.UUID):
    from app.services.marketplace_unload_wb_fbw_export_service import (
        export_wb_fbw_packaging_xlsx,
    )

    return await export_wb_fbw_packaging_xlsx(
        session, tenant_id=tenant_id, request_id=request_id
    )


class Ctx:
    tenant: Tenant
    other_tenant: Tenant
    warehouse: Warehouse
    seller: Seller
    request: MarketplaceUnloadRequest
    first_box: WarehouseBox
    legacy_box: WarehouseBox
    primary: Product
    fallback: Product
    shelf_life: Product


async def _fixture(session: AsyncSession) -> Ctx:
    """A two-tenant, three-box composition; no integration credentials exist here."""
    suffix = uuid.uuid4().hex[:8]
    ctx = Ctx()
    ctx.tenant = Tenant(name="WMS-679", slug=f"wms679-{suffix}")
    ctx.other_tenant = Tenant(name="Other", slug=f"wms679-other-{suffix}")
    session.add_all((ctx.tenant, ctx.other_tenant))
    await session.flush()
    ctx.warehouse = Warehouse(tenant_id=ctx.tenant.id, name="WMS", code=f"w679-{suffix}")
    ctx.seller = Seller(tenant_id=ctx.tenant.id, name="WB seller")
    session.add_all((ctx.warehouse, ctx.seller))
    await session.flush()

    ctx.primary = Product(
        tenant_id=ctx.tenant.id,
        seller_id=ctx.seller.id,
        name="Основной WB-код",
        sku_code=f"primary-{suffix}",
        wb_barcode="001234",
    )
    ctx.fallback = Product(
        tenant_id=ctx.tenant.id,
        seller_id=ctx.seller.id,
        name="Единственный вариантный код",
        sku_code=f"fallback-{suffix}",
        wb_barcode=None,
    )
    ctx.shelf_life = Product(
        tenant_id=ctx.tenant.id,
        seller_id=ctx.seller.id,
        name="Товар со сроком",
        sku_code=f"shelf-{suffix}",
        wb_barcode="=WB-001",
        wb_shelf_life="12 месяцев",
    )
    session.add_all((ctx.primary, ctx.fallback, ctx.shelf_life))
    await session.flush()
    session.add(
        ProductBarcode(
            tenant_id=ctx.tenant.id,
            seller_id=ctx.seller.id,
            product_id=ctx.fallback.id,
            barcode="0000456",
            source="wb",
        )
    )

    ctx.request = MarketplaceUnloadRequest(
        tenant_id=ctx.tenant.id,
        warehouse_id=ctx.warehouse.id,
        seller_id=ctx.seller.id,
        marketplace="wb",
        status="confirmed",
        document_number="WB-679",
    )
    ctx.first_box = WarehouseBox(
        tenant_id=ctx.tenant.id,
        warehouse_id=ctx.warehouse.id,
        internal_barcode="WHB-000000000001",
    )
    ctx.legacy_box = WarehouseBox(
        tenant_id=ctx.tenant.id,
        warehouse_id=ctx.warehouse.id,
        internal_barcode="INB-ABCDEF123456",
    )
    empty_box = WarehouseBox(
        tenant_id=ctx.tenant.id,
        warehouse_id=ctx.warehouse.id,
        internal_barcode="WHB-000000000003",
    )
    session.add_all((ctx.request, ctx.first_box, ctx.legacy_box, empty_box))
    await session.flush()
    first = MarketplaceUnloadBox(
        request_id=ctx.request.id, warehouse_box_id=ctx.first_box.id, box_preset="60_40_40"
    )
    legacy = MarketplaceUnloadBox(
        request_id=ctx.request.id, warehouse_box_id=ctx.legacy_box.id, box_preset="30_20_30"
    )
    empty = MarketplaceUnloadBox(
        request_id=ctx.request.id, warehouse_box_id=empty_box.id, box_preset="30_20_30"
    )
    session.add_all((first, legacy, empty))
    await session.flush()
    session.add_all(
        (
            MarketplaceUnloadBoxLine(box_id=first.id, product_id=ctx.primary.id, quantity=2),
            MarketplaceUnloadBoxLine(box_id=first.id, product_id=ctx.fallback.id, quantity=3),
            MarketplaceUnloadBoxLine(box_id=legacy.id, product_id=ctx.primary.id, quantity=1),
            MarketplaceUnloadBoxLine(box_id=legacy.id, product_id=ctx.shelf_life.id, quantity=4),
        )
    )

    location = StorageLocation(
        tenant_id=ctx.tenant.id,
        warehouse_id=ctx.warehouse.id,
        code=f"A-{suffix}",
        barcode=f"A-{suffix}",
    )
    session.add(location)
    await session.flush()
    session.add(
        InventoryBalance(
            tenant_id=ctx.tenant.id,
            product_id=ctx.primary.id,
            storage_location_id=location.id,
            quantity=20,
            quantity_unpacked=20,
            quantity_packed=0,
        )
    )
    await session.commit()
    return ctx


def _workbook(content: bytes):
    return load_workbook(BytesIO(content), data_only=False)


async def _read_snapshot(session: AsyncSession, ctx: Ctx) -> tuple[object, ...]:
    box_rows = list(
        (
            await session.execute(
                select(
                    MarketplaceUnloadBox.id,
                    MarketplaceUnloadBox.warehouse_box_id,
                    WarehouseBox.internal_barcode,
                )
                .join(WarehouseBox, WarehouseBox.id == MarketplaceUnloadBox.warehouse_box_id)
                .where(MarketplaceUnloadBox.request_id == ctx.request.id)
                .order_by(MarketplaceUnloadBox.id)
            )
        ).all()
    )
    line_rows = list(
        (
            await session.execute(
                select(
                    MarketplaceUnloadBoxLine.box_id,
                    MarketplaceUnloadBoxLine.product_id,
                    MarketplaceUnloadBoxLine.quantity,
                )
                .join(MarketplaceUnloadBox)
                .where(MarketplaceUnloadBox.request_id == ctx.request.id)
                .order_by(MarketplaceUnloadBoxLine.id)
            )
        ).all()
    )
    stock = list(
        (
            await session.execute(
                select(
                    InventoryBalance.product_id,
                    InventoryBalance.quantity,
                    InventoryBalance.quantity_unpacked,
                    InventoryBalance.quantity_packed,
                ).order_by(InventoryBalance.id)
            )
        ).all()
    )
    reserves = list((await session.execute(select(MarketplaceUnloadReservation.id))).all())
    status = await session.scalar(
        select(MarketplaceUnloadRequest.status).where(MarketplaceUnloadRequest.id == ctx.request.id)
    )
    return box_rows, line_rows, stock, reserves, status


@pytest.mark.asyncio
async def test_c1_c2_c3_c4_c6_wb_export_is_exact_string_xlsx_from_existing_boxes(
    db_session: AsyncSession,
) -> None:
    ctx = await _fixture(db_session)

    result = await _export(db_session, ctx.tenant.id, ctx.request.id)
    workbook = _workbook(result.content)
    assert workbook.sheetnames == ["Sheet1"]
    sheet = workbook["Sheet1"]
    assert [sheet.cell(1, column).value for column in range(1, 6)] == [
        "Баркод товара",
        "Кол-во товаров",
        "ШК короба",
        "Срок годности",
        "ШК короба для печати в стороннем сервисе",
    ]
    assert sheet.max_column == 5
    rows = [tuple(cell.value for cell in row) for row in sheet.iter_rows(min_row=2, max_col=5)]
    assert rows == [
        ("001234", 2, "WHB-000000000001", None, None),
        ("0000456", 3, "WHB-000000000001", None, None),
        ("001234", 1, "INB-ABCDEF123456", None, None),
        ("=WB-001", 4, "INB-ABCDEF123456", None, None),
    ]
    assert all(sheet.cell(row, 1).data_type == "s" for row in range(2, 6))
    assert all(sheet.cell(row, 3).data_type == "s" for row in range(2, 6))
    assert all(sheet.cell(row, 2).data_type == "n" for row in range(2, 6))
    assert all(isinstance(sheet.cell(row, 2).value, int) for row in range(2, 6))
    assert all(sheet.cell(row, 4).value is None for row in range(2, 6))
    assert all(sheet.cell(row, 5).value is None for row in range(2, 6))
    assert any("фактическую дату партии" in warning.lower() for warning in result.warnings)


@pytest.mark.asyncio
async def test_c5_invalid_box_or_ambiguous_variant_returns_actionable_error_without_partial_file(
    db_session: AsyncSession,
) -> None:
    ctx = await _fixture(db_session)
    ambiguous = Product(
        tenant_id=ctx.tenant.id,
        seller_id=ctx.seller.id,
        name="Неоднозначный вариант",
        sku_code=f"ambiguous-{uuid.uuid4().hex[:8]}",
    )
    db_session.add(ambiguous)
    await db_session.flush()
    db_session.add_all(
        (
            ProductBarcode(
                tenant_id=ctx.tenant.id, seller_id=ctx.seller.id, product_id=ambiguous.id,
                barcode="AMB-1", source="wb",
            ),
            ProductBarcode(
                tenant_id=ctx.tenant.id, seller_id=ctx.seller.id, product_id=ambiguous.id,
                barcode="AMB-2", source="wb",
            ),
            MarketplaceUnloadBoxLine(
                box_id=(await db_session.scalars(select(MarketplaceUnloadBox.id).where(
                    MarketplaceUnloadBox.request_id == ctx.request.id
                ))).first(),
                product_id=ambiguous.id,
                quantity=1,
            ),
        )
    )
    await db_session.commit()
    before = await _read_snapshot(db_session, ctx)

    with pytest.raises(Exception, match="Неоднозначный вариант"):
        await _export(db_session, ctx.tenant.id, ctx.request.id)

    assert await _read_snapshot(db_session, ctx) == before


@pytest.mark.asyncio
async def test_c7_c8_repeated_exports_are_read_only_and_each_file_is_a_complete_snapshot(
    db_session: AsyncSession,
) -> None:
    ctx = await _fixture(db_session)
    before = await _read_snapshot(db_session, ctx)

    first, second = await _export(db_session, ctx.tenant.id, ctx.request.id), await _export(
        db_session, ctx.tenant.id, ctx.request.id
    )

    assert first.content == second.content
    assert await _read_snapshot(db_session, ctx) == before
    assert [row[1] for row in _workbook(first.content).active.iter_rows(min_row=2, max_col=5)]


@pytest.mark.asyncio
async def test_c11_c12_export_is_wb_only_and_cannot_cross_tenant_boundary(
    db_session: AsyncSession,
) -> None:
    ctx = await _fixture(db_session)
    ctx.request.marketplace = "ozon"
    await db_session.commit()
    with pytest.raises(Exception, match="WB"):
        await _export(db_session, ctx.tenant.id, ctx.request.id)

    ctx.request.marketplace = "wb"
    await db_session.commit()
    with pytest.raises(Exception, match=r"not_found|не найдена|доступ"):
        await _export(db_session, ctx.other_tenant.id, ctx.request.id)


def test_c14_anonymized_official_template_is_the_only_excel_fixture() -> None:
    evidence = Path(__file__).resolve().parents[2] / "docs/evidence/WMS-679"
    workbook = load_workbook(evidence / "wb-fbw-shk-excel-template-anonymized.xlsx")
    assert workbook.sheetnames == ["Sheet1"]
    assert [workbook.active.cell(1, column).value for column in range(1, 6)] == [
        "Баркод товара",
        "Кол-во товаров",
        "ШК короба",
        "Срок годности",
        "ШК короба для печати в стороннем сервисе",
    ]
    assert "3dc986bd74ca7397cd7ab6ebb02793cd4b3e017a317d517b5cbd13915a1640c7" in (
        evidence / "README.md"
    ).read_text()
