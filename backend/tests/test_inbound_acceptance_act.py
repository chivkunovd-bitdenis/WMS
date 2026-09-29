# ruff: noqa: E501
from __future__ import annotations

import io
import uuid
from urllib.parse import unquote

import pytest
from openpyxl import load_workbook  # type: ignore[import-untyped]
from sqlalchemy import select

from app.db.session import SessionLocal
from app.models.inbound_intake import InboundIntakeLine, InboundIntakeRequest

BASE = "/operations/inbound-intake-requests"


async def _setup(async_client) -> tuple[dict[str, str], str, list[str]]:
    reg = await async_client.post("/auth/register", json={
        "organization_name": "Акт приёмки", "slug": f"act-{uuid.uuid4().hex}",
        "admin_email": f"act-{uuid.uuid4().hex}@example.com", "password": "password123",
    })
    headers = {"Authorization": f"Bearer {reg.json()['access_token']}"}
    seller = (await async_client.post("/sellers", headers=headers, json={"name": "ИП Акт"})).json()["id"]
    warehouses = (await async_client.get("/warehouses", headers=headers)).json()
    warehouse = warehouses[0]["id"] if warehouses else (
        await async_client.post("/warehouses", headers=headers, json={"name": "Основной", "code": "MAIN"})
    ).json()["id"]
    products = []
    for name, sku, barcode in (("Футболка M", "TS-M", "0460000000011"), ("=1+1 худи", "HD-L", "4680000000028")):
        response = await async_client.post("/products", headers=headers, json={
            "name": name, "sku_code": sku, "seller_id": seller, "wb_barcode": barcode, "wb_vendor_code": f"V-{sku}",
        })
        assert response.status_code in (200, 201), response.text
        products.append(response.json()["id"])
    created = await async_client.post(BASE, headers=headers, json={"warehouse_id": warehouse, "seller_id": seller})
    assert created.status_code == 201, created.text
    request_id = created.json()["id"]
    for product_id, qty in zip(products, (10, 5), strict=True):
        line = await async_client.post(f"{BASE}/{request_id}/lines", headers=headers,
                                       json={"product_id": product_id, "expected_qty": qty})
        assert line.status_code == 201, line.text
    return headers, request_id, products


async def _close_reception(request_id: str, facts: dict[int, int]) -> None:
    async with SessionLocal() as session:
        request = await session.get(InboundIntakeRequest, uuid.UUID(request_id))
        assert request is not None
        request.status = "sorting"
        lines = (await session.scalars(
            select(InboundIntakeLine).where(InboundIntakeLine.request_id == request.id)
        )).all()
        for line in lines:
            line.actual_qty = facts[line.expected_qty]
        await session.commit()


@pytest.mark.asyncio
async def test_wms586_acceptance_act_has_plan_fact_and_difference(async_client) -> None:
    headers, request_id, _ = await _setup(async_client)
    await _close_reception(request_id, {10: 8, 5: 7})

    response = await async_client.get(f"{BASE}/{request_id}/acceptance-act.xlsx", headers=headers)
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    disposition = unquote(response.headers["content-disposition"])
    assert "Акт приёмки" in disposition and ".xlsx" in disposition

    sheet = load_workbook(io.BytesIO(response.content)).active
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    assert rows[0][0].startswith("Приёмка №") and " от " in rows[0][0]
    assert rows[1] == [None] * len(rows[1])  # сразу после «номер и дата» — таблица, без шапок
    header_index = next(i for i, row in enumerate(rows) if row[0] == "№")
    assert rows[header_index] == ["№", "Товар", "Артикул продавца", "SKU", "ШК", "План", "Факт", "Расхождение"]
    body = {row[1]: row for row in rows[header_index + 1:] if row[1] and row[1] != "Итого"}
    assert body["Футболка M"][5:8] == [10, 8, -2]
    # Название вида «=…» остаётся текстом, ШК не теряет ведущий ноль.
    assert body["=1+1 худи"][5:8] == [5, 7, 2]
    assert body["Футболка M"][4] == "0460000000011"
    total = next(row for row in rows if row[1] == "Итого")
    assert total[5:8] == [15, 15, 0]


@pytest.mark.asyncio
async def test_wms586_acceptance_act_refused_before_reception_is_closed(async_client) -> None:
    headers, request_id, _ = await _setup(async_client)
    response = await async_client.get(f"{BASE}/{request_id}/acceptance-act.xlsx", headers=headers)
    assert response.status_code == 409
    assert response.json()["detail"] == "reception_not_closed"
