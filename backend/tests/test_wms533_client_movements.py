"""WMS-533 client movement API: exact FBS units and full Excel traversal."""

from __future__ import annotations

import io
import time
import uuid
from datetime import UTC, datetime

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import event, select

from app.db.session import SessionLocal, engine
from app.models.fbs_order import FbsOrder, FbsOrderMarking, FbsOrderProduct
from app.models.fbs_shipment_reversal_ledger import FbsShipmentReversalLedger
from app.models.inventory_movement import InventoryMovement
from app.models.marketplace_unload import MarketplaceUnloadRequest
from app.models.product import Product
from app.models.user import User
from app.services.tokens import decode_access_token

PERIOD = {"date_from": "2026-09-01T00:00:00Z", "date_to": "2026-10-01T00:00:00Z"}
AT = datetime(2026, 9, 12, 12, tzinfo=UTC)


async def _org(
    async_client: AsyncClient, *, name: str
) -> tuple[dict[str, str], uuid.UUID, uuid.UUID]:
    suffix = str(time.time_ns())
    registered = await async_client.post(
        "/auth/register",
        json={
            "organization_name": name,
            "slug": f"{name.lower()}-{suffix}",
            "admin_email": f"{name.lower()}-{suffix}@example.com",
            "password": "password123",
        },
    )
    assert registered.status_code == 200, registered.text
    token = str(registered.json()["access_token"])
    claims = decode_access_token(token)
    return (
        {"Authorization": f"Bearer {token}"},
        uuid.UUID(str(claims["tenant_id"])),
        uuid.UUID(str(claims["sub"])),
    )


async def _seller(async_client: AsyncClient, headers: dict[str, str], name: str) -> uuid.UUID:
    response = await async_client.post("/sellers", headers=headers, json={"name": name})
    assert response.status_code in (200, 201), response.text
    return uuid.UUID(response.json()["id"])


async def _warehouse_location(
    async_client: AsyncClient, headers: dict[str, str], *, name: str
) -> tuple[uuid.UUID, uuid.UUID]:
    suffix = str(time.time_ns())
    warehouse = await async_client.post(
        "/warehouses", headers=headers, json={"name": name, "code": f"{name}-{suffix}"}
    )
    assert warehouse.status_code == 200, warehouse.text
    warehouse_id = uuid.UUID(warehouse.json()["id"])
    location = await async_client.post(
        f"/warehouses/{warehouse_id}/locations",
        headers=headers,
        json={"code": f"A-{suffix}"},
    )
    assert location.status_code == 200, location.text
    return warehouse_id, uuid.UUID(location.json()["id"])


@pytest.mark.asyncio
async def test_client_report_units_filters_cursor_and_excel(async_client: AsyncClient) -> None:
    headers, tenant_id, _ = await _org(async_client, name="ClientReport")
    seller_id = await _seller(async_client, headers, "Seller")
    warehouse_id, location_id = await _warehouse_location(async_client, headers, name="ReportWH")
    other_warehouse, other_location = await _warehouse_location(
        async_client, headers, name="OtherWH"
    )
    ids: dict[str, uuid.UUID] = {}
    async with SessionLocal() as session:
        for sku in ("0007", "0008"):
            product = Product(
                tenant_id=tenant_id,
                seller_id=seller_id,
                name="=1+1" if sku == "0007" else f"Item {sku}",
                sku_code=sku,
                wb_barcode=f"00{sku}",
                wb_size="42" if sku == "0007" else "44",
            )
            session.add(product)
            await session.flush()
            ids[sku] = product.id
        wb_order = FbsOrder(
            tenant_id=tenant_id,
            seller_id=seller_id,
            marketplace="wb",
            wb_order_id=701,
            product_id=ids["0007"],
            warehouse_id=warehouse_id,
            created_at_wb=AT,
            deadline_at=AT,
            mapping_status="mapped",
            reserve_status="no_stock",
        )
        ozon_order = FbsOrder(
            tenant_id=tenant_id,
            seller_id=seller_id,
            marketplace="ozon",
            wb_order_id=-702,
            external_order_id="OZ-702",
            product_id=ids["0007"],
            warehouse_id=warehouse_id,
            created_at_wb=AT,
            deadline_at=AT,
            mapping_status="mapped",
            reserve_status="no_stock",
        )
        session.add_all([wb_order, ozon_order])
        await session.flush()
        first_position = FbsOrderProduct(
            order_id=ozon_order.id, product_id=ids["0007"], quantity=2, position_index=0
        )
        second_position = FbsOrderProduct(
            order_id=ozon_order.id, product_id=ids["0008"], quantity=1, position_index=1
        )
        session.add_all([first_position, second_position])
        await session.flush()
        gs = "0101234567890121\x1dSERIAL"
        session.add_all(
            [
                FbsOrderMarking(
                    tenant_id=tenant_id,
                    order_id=wb_order.id,
                    kind="sgtin",
                    value="WB-KIZ",
                    meta_status="accepted",
                ),
                FbsOrderMarking(
                    tenant_id=tenant_id,
                    order_id=ozon_order.id,
                    order_product_id=first_position.id,
                    kind="sgtin",
                    value=gs,
                    meta_status="accepted",
                ),
                FbsOrderMarking(
                    tenant_id=tenant_id,
                    order_id=ozon_order.id,
                    order_product_id=first_position.id,
                    kind="sgtin",
                    value="OZ-2",
                    meta_status="accepted",
                ),
                FbsOrderMarking(
                    tenant_id=tenant_id,
                    order_id=ozon_order.id,
                    order_product_id=second_position.id,
                    kind="sgtin",
                    value="OZ-3",
                    meta_status="accepted",
                ),
                FbsOrderMarking(
                    tenant_id=tenant_id,
                    order_id=ozon_order.id,
                    order_product_id=second_position.id,
                    kind="sgtin",
                    value="REJECTED",
                    meta_status="rejected",
                ),
                FbsOrderMarking(
                    tenant_id=tenant_id,
                    order_id=ozon_order.id,
                    kind="sgtin",
                    value="UNBOUND",
                    meta_status="accepted",
                ),
            ]
        )
        await session.flush()

        def movement(
            sku: str,
            qty: int,
            kind: str,
            *,
            wh: uuid.UUID = warehouse_id,
            loc: uuid.UUID = location_id,
            at: datetime = AT,
        ) -> InventoryMovement:
            value = InventoryMovement(
                tenant_id=tenant_id,
                product_id=ids[sku],
                seller_id=seller_id,
                warehouse_id=wh,
                storage_location_id=loc,
                quantity_delta=qty,
                movement_type=kind,
                created_at=at,
            )
            session.add(value)
            return value

        wb_move = movement("0007", -1, "fbs_shipment")
        oz_move1 = movement("0007", -1, "fbs_shipment")
        oz_split = movement("0007", -1, "fbs_shipment")
        oz_group = uuid.uuid4()
        oz_move1.transfer_group_id = oz_group
        oz_split.transfer_group_id = oz_group
        oz_move2 = movement("0008", -1, "fbs_shipment")
        reversal = movement("0007", 1, "fbs_shipment")
        move_internal = movement("0007", -4, "stock_transfer_out")
        move_other = movement("0007", 9, "inbound_intake", wh=other_warehouse, loc=other_location)
        unload = MarketplaceUnloadRequest(
            tenant_id=tenant_id,
            warehouse_id=warehouse_id,
            seller_id=seller_id,
            marketplace="ozon",
            status="shipped",
        )
        session.add(unload)
        await session.flush()
        fbo = movement("0008", -3, "marketplace_unload")
        fbo.marketplace_unload_request_id = unload.id
        await session.flush()
        session.add_all(
            [
                FbsShipmentReversalLedger(
                    tenant_id=tenant_id,
                    fbs_order_id=wb_order.id,
                    product_id=ids["0007"],
                    storage_location_id=location_id,
                    quantity=1,
                    shipment_movement_id=wb_move.id,
                    reversal_movement_id=reversal.id,
                ),
                FbsShipmentReversalLedger(
                    tenant_id=tenant_id,
                    fbs_order_id=ozon_order.id,
                    product_id=ids["0007"],
                    storage_location_id=location_id,
                    quantity=3,
                    shipment_movement_id=oz_move1.id,
                    ozon_positions_json=[
                        {
                            "product_id": str(ids["0007"]),
                            "movement_id": str(oz_move1.id),
                            "quantity": 2,
                        },
                        {
                            "product_id": str(ids["0008"]),
                            "movement_id": str(oz_move2.id),
                            "quantity": 1,
                        },
                    ],
                ),
            ]
        )
        await session.commit()

    query = {**PERIOD, "warehouse_id": str(warehouse_id), "limit": 2}
    rows = []
    cursor = None
    ledger_queries: list[tuple[str, object]] = []

    def capture_ledger_query(_conn, _cursor, statement, parameters, _context, _many):
        if "FROM fbs_shipment_reversal_ledger" in statement:
            ledger_queries.append((statement, parameters))

    event.listen(engine.sync_engine, "before_cursor_execute", capture_ledger_query)
    try:
        while True:
            response = await async_client.get(
                "/reports/client-movements",
                headers=headers,
                params={**query, **({"cursor": cursor} if cursor else {})},
            )
            assert response.status_code == 200, response.text
            rows.extend(response.json()["rows"])
            cursor = response.json()["next_cursor"]
            if cursor is None:
                break
    finally:
        event.remove(engine.sync_engine, "before_cursor_execute", capture_ledger_query)
    assert ledger_queries
    assert all(
        ("json_each" in statement or "jsonb_array_elements" in statement)
        and "shipment_movement_id IN" in statement
        and len(parameters) < 1000
        for statement, parameters in ledger_queries
    )
    assert len(rows) == 7
    assert len({row["id"] for row in rows}) == 7
    assert move_other.id.hex not in {row["movement_id"].replace("-", "") for row in rows}
    assert sorted(
        row["kiz"] for row in rows if row["movement_id"] in {str(oz_move1.id), str(oz_split.id)}
    ) == sorted([gs, "OZ-2"])
    assert next(row for row in rows if row["movement_id"] == str(oz_move2.id))["kiz"] == "OZ-3"
    assert next(row for row in rows if row["movement_id"] == str(fbo.id))["marketplace"] == "ozon"
    assert next(row for row in rows if row["movement_id"] == str(fbo.id))["kiz"] is None
    assert (
        next(row for row in rows if row["movement_id"] == str(move_internal.id))["marketplace"]
        is None
    )
    assert (
        next(row for row in rows if row["movement_id"] == str(reversal.id))["quantity_delta"] == 1
    )
    assert all(
        row["quantity_delta"] == -1
        for row in rows
        if row["movement_id"] in {str(oz_move1.id), str(oz_split.id)}
    )

    response = await async_client.get(
        "/reports/client-movements",
        headers=headers,
        params={
            **PERIOD,
            "shk": "000008",
            "warehouse_id": str(warehouse_id),
            "marketplace": "ozon",
        },
    )
    assert response.status_code == 200
    assert {row["movement_id"] for row in response.json()["rows"]} == {
        str(oz_move2.id), str(fbo.id)
    }
    assert {row["shk"] for row in response.json()["rows"]} == {"000008"}
    assert {row["size"] for row in response.json()["rows"]} == {"44"}
    sku_response = await async_client.get(
        "/reports/client-movements", headers=headers, params={**PERIOD, "sku": "0008"}
    )
    assert sku_response.status_code == 200
    assert {row["sku"] for row in sku_response.json()["rows"]} == {"0008"}
    for path in ("/reports/client-movements", "/reports/client-movements/export.xlsx"):
        conflict = await async_client.get(
            path, headers=headers, params={**PERIOD, "sku": "0008", "shk": "000008"}
        )
        assert conflict.status_code == 422
    assert (
        await async_client.get(
            "/reports/client-movements",
            headers=headers,
            params={**PERIOD, "date_from": "2026-09-01T00:00:00"},
        )
    ).status_code == 422

    exported = await async_client.get(
        "/reports/client-movements/export.xlsx",
        headers=headers,
        params={**PERIOD, "warehouse_id": str(warehouse_id)},
    )
    assert exported.status_code == 200, exported.text[:100]
    book = load_workbook(io.BytesIO(exported.content), read_only=True)
    assert book.sheetnames == ["WB", "Ozon", "Общие"]
    wb_rows = list(book["WB"].values)[1:]
    oz_rows = list(book["Ozon"].values)[1:]
    general_rows = list(book["Общие"].values)[1:]
    assert len(wb_rows) == 2 and len(oz_rows) == 4 and len(general_rows) == 1
    assert all(
        cells[0].value == "=1+1" and cells[0].data_type == "s"
        for cells in book["WB"].iter_rows(min_row=2, min_col=8, max_col=8)
    )
    assert any(row[6] == "0007" and "\\u001d" in row[-1] for row in oz_rows)
    assert all(row[6] in {"0007", "0008"} for row in oz_rows)


@pytest.mark.asyncio
async def test_client_report_full_page_boundary_and_scope(async_client: AsyncClient) -> None:
    headers, tenant_id, _ = await _org(async_client, name="ClientReportMany")
    seller_id = await _seller(async_client, headers, "Scope A")
    second_seller = await _seller(async_client, headers, "Scope B")
    warehouse_id, location_id = await _warehouse_location(async_client, headers, name="ManyWH")
    at_start = datetime(2026, 9, 1, tzinfo=UTC)
    at_end = datetime(2026, 10, 1, tzinfo=UTC)
    async with SessionLocal() as session:
        product = Product(
            tenant_id=tenant_id,
            seller_id=seller_id,
            name="Many",
            sku_code="001",
            wb_barcode="00009876",
            wb_size="XL",
        )
        other = Product(
            tenant_id=tenant_id,
            seller_id=second_seller,
            name="Other",
            sku_code="002",
            wb_barcode=None,
        )
        session.add_all([product, other])
        await session.flush()
        for index in range(205):
            session.add(
                InventoryMovement(
                    tenant_id=tenant_id,
                    product_id=product.id,
                    seller_id=seller_id,
                    warehouse_id=warehouse_id,
                    storage_location_id=location_id,
                    quantity_delta=1,
                    movement_type="inbound_intake",
                    created_at=at_start if index == 0 else AT,
                )
            )
        session.add(
            InventoryMovement(
                tenant_id=tenant_id,
                product_id=product.id,
                seller_id=seller_id,
                warehouse_id=warehouse_id,
                storage_location_id=location_id,
                quantity_delta=1,
                movement_type="inbound_intake",
                created_at=at_end,
            )
        )
        session.add(
            InventoryMovement(
                tenant_id=tenant_id,
                product_id=other.id,
                seller_id=second_seller,
                warehouse_id=warehouse_id,
                storage_location_id=location_id,
                quantity_delta=1,
                movement_type="inbound_intake",
                created_at=AT,
            )
        )
        await session.commit()
    assert (await async_client.get("/reports/client-movements", params=PERIOD)).status_code == 401
    response = await async_client.get(
        "/reports/client-movements", headers=headers, params={**PERIOD, "shk": "00009876"}
    )
    assert response.status_code == 200, response.text
    assert len(response.json()["rows"]) == 200
    assert response.json()["next_cursor"] is not None
    second = await async_client.get(
        "/reports/client-movements",
        headers=headers,
        params={**PERIOD, "shk": "00009876", "cursor": response.json()["next_cursor"]},
    )
    assert second.status_code == 200, second.text
    assert len(second.json()["rows"]) == 5
    assert second.json()["next_cursor"] is None
    assert all(
        row["shk"] == "00009876" and row["size"] == "XL"
        for row in response.json()["rows"] + second.json()["rows"]
    )
    export = await async_client.get(
        "/reports/client-movements/export.xlsx",
        headers=headers,
        params={**PERIOD, "shk": "00009876"},
    )
    assert export.status_code == 200
    book = load_workbook(io.BytesIO(export.content), read_only=True)
    general_rows = list(book["Общие"].values)
    assert len(general_rows) - 1 == 205
    assert general_rows[0][8:10] == ("shk", "size")
    assert general_rows[1][8:10] == ("00009876", "XL")

    email = f"client-report-{uuid.uuid4().hex[:10]}@example.com"
    created = await async_client.post(
        "/auth/seller-accounts",
        headers=headers,
        json={"seller_id": str(seller_id), "email": email, "password": "password123"},
    )
    assert created.status_code in (200, 201), created.text
    login = await async_client.post("/auth/login", json={"email": email, "password": "password123"})
    seller_headers = {"Authorization": f"Bearer {login.json()['access_token']}"}
    scoped = await async_client.get(
        "/reports/client-movements",
        headers=seller_headers,
        params={**PERIOD, "limit": 1},
    )
    assert scoped.status_code == 200, scoped.text
    assert scoped.json()["rows"][0]["sku"] == "001"
    assert scoped.json()["next_cursor"] is not None
    async with SessionLocal() as session:
        user = await session.scalar(select(User).where(User.email == email))
        assert user is not None
        user.seller_id = None
        await session.commit()
    no_scope = await async_client.get(
        "/reports/client-movements", headers=seller_headers, params=PERIOD
    )
    assert no_scope.status_code == 403
    tenant_headers, _, _ = await _org(async_client, name="ClientReportOtherTenant")
    isolated = await async_client.get(
        "/reports/client-movements", headers=tenant_headers, params=PERIOD
    )
    assert isolated.status_code == 200
    assert isolated.json()["rows"] == []
