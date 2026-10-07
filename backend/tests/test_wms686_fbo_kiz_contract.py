"""WMS-686: контракт подбора и упаковки FBO с необязательным КИЗ (до реализации).

Источник ожиданий — `docs/requirements/WMS-686.md` (редакция 3): требования
R1-R33, интерфейс контракта (раздел 9) и сценарии проверок C (раздел 10.1). Имя
каждого теста начинается с идентификатора проверки, чтобы ссылка `путь::имя` была
однозначной.

На базе до реализации (прогон 08.10.2026: 6 passed, 31 failed):
- GREEN (сохраняемое поведение): c03, c06, c11, c13, c32, c51;
- RED (функции нет): все остальные — падают на проверке ответа (AssertionError),
  а не на подготовке. c50/c51 помечены postgresql_concurrency и в целевом CI
  дополнительно идут на PostgreSQL.

Все сценарии идут через HTTP API, как работает веб: настоящие приёмка,
отгрузка, подбор, короба и «Завершить». Заменена только внешняя граница —
фоновая проверка кодов в Честном знаке после проведения приёмки (C26).

КИЗ строится в формате GS1 DataMatrix, который уже принимает приёмка
(`inbound_marking_service.normalize_scanned_code`): `01` + GTIN-14 + `21` +
серийный номер + GS + `93` + хвост. GTIN-14 — это `0` + EAN-13 товара, как его
сопоставляет `marking_code_service._gtin_lookup_variants`.
"""

from __future__ import annotations

import asyncio
import csv
import io
import json
import uuid
from typing import Any, cast

import pytest
from httpx import AsyncClient, Response
from inbound_box_intake_helpers import (  # type: ignore[import-not-found]
    fulfill_inbound_via_box_scans,
    post_primary_accept,
    set_planned_boxes,
)
from openpyxl import load_workbook  # type: ignore[import-untyped]
from sqlalchemy import func, select
from test_marketplace_unload_and_discrepancy_acts import (  # type: ignore[import-not-found]
    _patch_mp_planned_date,
    _patch_packaging_instructions,
    _seller_wb_mp_warehouse,
)
from test_marketplace_unload_pick_from_container import (  # type: ignore[import-not-found]
    _loose_balance_id,
)

from app.db.session import SessionLocal
from app.models.inventory_balance import InventoryBalance
from app.models.inventory_movement import InventoryMovement
from app.models.storage_location import StorageLocation
from app.services.inbound_marking_service import normalize_scanned_code
from app.services.marking_code_service import extract_gtin_from_cis
from app.services.sorting_location_service import SORTING_LOCATION_CODE

BASE = "/operations/marketplace-unload-requests"
INBOUND = "/operations/inbound-intake-requests"
GS = "\x1d"


def _ean13(prefix12: str) -> str:
    """EAN-13 с верной контрольной цифрой: ШК товара, из которого строится GTIN КИЗ."""
    total = sum(int(digit) * (3 if index % 2 else 1) for index, digit in enumerate(prefix12))
    return prefix12 + str((10 - total % 10) % 10)


EAN_A = _ean13("460686000001")
EAN_B = _ean13("460686000002")


def _kiz(ean13: str, serial: str) -> str:
    """КИЗ товара с данным EAN-13; формат тот же, что принимает приёмка."""
    code = f"010{ean13}21{serial}{GS}93{serial[-4:]}"
    assert normalize_scanned_code(code) == code
    assert extract_gtin_from_cis(code) == f"0{ean13}"
    return code


def _head(code: str) -> str:
    """Неизменяемая часть КИЗ (`01…21серийный`) для сравнения с сохранённым кодом."""
    return code.split(GS, 1)[0]


def _link_for(links: list[dict[str, Any]], code: str) -> dict[str, Any]:
    found = [link for link in links if str(link["cis_code"]).startswith(_head(code))]
    assert len(found) == 1, (code, links)
    return found[0]


# --------------------------------------------------------------------------- данные


async def _register(client: AsyncClient) -> dict[str, str]:
    """Новая организация ФФ; заголовки её администратора."""
    suffix = uuid.uuid4().hex[:10]
    reg = await client.post(
        "/auth/register",
        json={
            "organization_name": "WMS-686 FBO KIZ",
            "slug": f"wms686-{suffix}",
            "admin_email": f"wms686-{suffix}@example.com",
            "password": "password123",
        },
    )
    assert reg.status_code in (200, 201), reg.text
    return {"Authorization": f"Bearer {reg.json()['access_token']}"}


async def _org(
    client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> tuple[dict[str, str], str, str, int]:
    """Организация ФФ, склад, селлер с подключённым WB и склад WB для отгрузки."""
    h = await _register(client)
    wh = await client.post(
        "/warehouses", headers=h, json={"name": "W", "code": f"w-{uuid.uuid4().hex[:10]}"}
    )
    assert wh.status_code in (200, 201), wh.text
    sid, wb_wid = await _seller_wb_mp_warehouse(client, h, monkeypatch)
    return h, str(wh.json()["id"]), str(sid), int(wb_wid)


async def _cell(client: AsyncClient, h: dict[str, str], wid: str, code: str) -> tuple[str, str]:
    loc = await client.post(f"/warehouses/{wid}/locations", headers=h, json={"code": code})
    assert loc.status_code == 200, loc.text
    return str(loc.json()["id"]), str(loc.json()["barcode"])


async def _product(
    client: AsyncClient,
    h: dict[str, str],
    sid: str,
    *,
    barcode: str,
    name: str = "Футболка 48",
    honest_sign: bool = False,
    sku: str | None = None,
) -> str:
    created = await client.post(
        "/products",
        headers=h,
        json={
            "name": name,
            "sku_code": sku or f"SKU-{uuid.uuid4().hex[:8]}",
            "length_mm": 1,
            "width_mm": 1,
            "height_mm": 1,
            "seller_id": sid,
            "wb_barcode": barcode,
            "requires_honest_sign": honest_sign,
        },
    )
    assert created.status_code in (200, 201), created.text
    pid = str(created.json()["id"])
    await _patch_packaging_instructions(client, h, pid)
    return pid


async def _receive(
    client: AsyncClient,
    h: dict[str, str],
    *,
    wid: str,
    pid: str,
    qty: int,
    loc_id: str,
    kiz: tuple[str, ...] = (),
) -> str:
    """Настоящая приёмка россыпью в ячейку; КИЗ сканируются в приёмке как в вебе.

    Возвращает номер документа приёмки.
    """
    created = await client.post(INBOUND, headers=h, json={"warehouse_id": wid})
    assert created.status_code == 201, created.text
    rid = str(created.json()["id"])
    line = await client.post(
        f"{INBOUND}/{rid}/lines",
        headers=h,
        json={"product_id": pid, "expected_qty": qty, "storage_location_id": loc_id},
    )
    assert line.status_code == 201, line.text
    await set_planned_boxes(client, INBOUND, rid, h)
    submit = await client.post(f"{INBOUND}/{rid}/submit", headers=h)
    assert submit.status_code == 200, submit.text
    await post_primary_accept(client, INBOUND, rid, h)
    await fulfill_inbound_via_box_scans(client, h, rid, line.json()["sku_code"], qty)
    for code in kiz:
        scanned = await client.post(
            f"{INBOUND}/{rid}/marking-codes/scan",
            headers=h,
            json={"line_id": line.json()["id"], "cis_code": code},
        )
        assert scanned.status_code == 200, scanned.text
    verify = await client.post(f"{INBOUND}/{rid}/verify", headers=h)
    assert verify.status_code == 200, verify.text
    post = await client.post(f"{INBOUND}/{rid}/post", headers=h)
    assert post.status_code == 200, post.text
    detail = await client.get(f"{INBOUND}/{rid}", headers=h)
    assert detail.status_code == 200, detail.text
    return str(detail.json()["document_number"])


async def _box_in_cell(
    client: AsyncClient,
    h: dict[str, str],
    *,
    wid: str,
    loc_id: str,
    contents: dict[str, int],
) -> tuple[str, str]:
    """Короб склада в ячейке; товар перекладывается в него из россыпи этой ячейки."""
    box = await client.post(f"/warehouses/{wid}/sorting-objects", headers=h, json={"kind": "box"})
    assert box.status_code == 201, box.text
    box_id, box_barcode = str(box.json()["id"]), str(box.json()["barcode"])
    placed = await client.post(
        f"/warehouses/{wid}/map/move",
        headers=h,
        json={"kind": "box", "id": box_id, "to_kind": "cell", "to_id": loc_id, "qty": 1},
    )
    assert placed.status_code == 200, placed.text
    for pid, qty in contents.items():
        moved = await client.post(
            f"/warehouses/{wid}/map/move",
            headers=h,
            json={
                "kind": "product",
                "id": await _loose_balance_id(loc_id, pid),
                "to_kind": "box",
                "to_id": box_id,
                "qty": qty,
            },
        )
        assert moved.status_code == 200, moved.text
    return box_id, box_barcode


async def _shipment(
    client: AsyncClient,
    h: dict[str, str],
    *,
    wid: str,
    sid: str,
    wb_wid: int,
    lines: dict[str, int],
) -> tuple[str, str]:
    """Утверждённая отгрузка FBO на WB; возвращает id и номер документа."""
    created = await client.post(
        BASE, headers=h, json={"warehouse_id": wid, "seller_id": sid, "wb_mp_warehouse_id": wb_wid}
    )
    assert created.status_code == 201, created.text
    mid = str(created.json()["id"])
    for pid, qty in lines.items():
        added = await client.post(
            f"{BASE}/{mid}/lines", headers=h, json={"product_id": pid, "quantity": qty}
        )
        assert added.status_code in (200, 201), added.text
    await _patch_mp_planned_date(client, h, mid)
    submitted = await client.post(f"{BASE}/{mid}/submit", headers=h)
    assert submitted.status_code == 200, submitted.text
    assert submitted.json()["status"] == "confirmed"
    return mid, str(submitted.json()["document_number"])


async def _seller_headers(
    client: AsyncClient, admin_h: dict[str, str], sid: str
) -> dict[str, str]:
    email = f"wms686-seller-{uuid.uuid4().hex[:10]}@example.com"
    account = await client.post(
        "/auth/seller-accounts",
        headers=admin_h,
        json={"seller_id": sid, "email": email, "password": "password123"},
    )
    assert account.status_code == 201, account.text
    login = await client.post("/auth/login", json={"email": email, "password": "password123"})
    assert login.status_code == 200, login.text
    return {"Authorization": f"Bearer {login.json()['access_token']}"}


# ----------------------------------------------------------------------- действия


async def _pick_scan(
    client: AsyncClient,
    h: dict[str, str],
    mid: str,
    barcode: str,
    *,
    pid: str | None = None,
    loc_id: str | None = None,
    box_id: str | None = None,
    mutation_id: str | None = None,
) -> Response:
    """Скан на вкладке «Подбор» в том виде, в каком его шлёт веб.

    Для ШК и КИЗ передаются текущий источник (ячейка и тара) и целевой товар
    (последний отсканированный) — раздел 9 п. 1.
    """
    body: dict[str, Any] = {"barcode": barcode}
    if pid is not None:
        body["product_id"] = pid
    if loc_id is not None:
        body["storage_location_id"] = loc_id
    if box_id is not None:
        body["container_kind"] = "box"
        body["container_id"] = box_id
    if mutation_id is not None:
        body["mutation_id"] = mutation_id
    return await client.post(f"{BASE}/{mid}/pick/scan", headers=h, json=body)


async def _pick_set(
    client: AsyncClient,
    h: dict[str, str],
    mid: str,
    *,
    pid: str,
    loc_id: str,
    box_id: str | None,
    qty: int,
) -> Response:
    body: dict[str, Any] = {"product_id": pid, "storage_location_id": loc_id, "quantity": qty}
    if box_id is not None:
        body["container_kind"] = "box"
        body["container_id"] = box_id
    return await client.post(f"{BASE}/{mid}/pick/set", headers=h, json=body)


async def _new_box(client: AsyncClient, h: dict[str, str], mid: str) -> tuple[str, str]:
    created = await client.post(f"{BASE}/{mid}/boxes", headers=h, json={"box_preset": "60_40_40"})
    assert created.status_code == 201, created.text
    return str(created.json()["id"]), str(created.json()["internal_barcode"])


async def _pack_scan(
    client: AsyncClient,
    h: dict[str, str],
    mid: str,
    box_id: str,
    barcode: str,
    *,
    pid: str | None = None,
    loc_id: str | None = None,
) -> Response:
    """Скан поля «Упаковки»: только уже подобранные единицы (раздел 9 п. 5)."""
    body: dict[str, Any] = {"barcode": barcode, "pick_from_storage": False}
    if pid is not None:
        body["product_id"] = pid
    if loc_id is not None:
        body["storage_location_id"] = loc_id
    return await client.post(f"{BASE}/{mid}/boxes/{box_id}/scan", headers=h, json=body)


async def _ship(client: AsyncClient, h: dict[str, str], mid: str) -> Response:
    try:
        return await client.post(f"{BASE}/{mid}/ship", headers=h)
    except Exception as exc:  # отказ сервера без HTTP-ответа — тоже провал «Завершить»
        pytest.fail(f"«Завершить» не вернуло ответ: {exc!r}")


async def _links(client: AsyncClient, h: dict[str, str], mid: str) -> list[dict[str, Any]]:
    got = await client.get(f"{BASE}/{mid}/marking-codes", headers=h)
    assert got.status_code == 200, got.text
    return cast(list[dict[str, Any]], got.json())


async def _picked(client: AsyncClient, h: dict[str, str], mid: str, pid: str) -> int:
    """«Снято» по товару на вкладке «Подбор»."""
    options = await client.get(f"{BASE}/{mid}/pick-options", headers=h)
    assert options.status_code == 200, options.text
    return int(next(row for row in options.json() if row["product_id"] == pid)["picked_qty"])


async def _box_lines(
    client: AsyncClient, h: dict[str, str], mid: str, box_id: str
) -> dict[str, int]:
    detail = await client.get(f"{BASE}/{mid}", headers=h)
    assert detail.status_code == 200, detail.text
    box = next(row for row in detail.json()["boxes"] if row["id"] == box_id)
    return {str(line["product_id"]): int(line["quantity"]) for line in box["lines"]}


async def _box_line_id(
    client: AsyncClient, h: dict[str, str], mid: str, box_id: str, pid: str
) -> str:
    detail = await client.get(f"{BASE}/{mid}", headers=h)
    assert detail.status_code == 200, detail.text
    box = next(row for row in detail.json()["boxes"] if row["id"] == box_id)
    return str(next(line["id"] for line in box["lines"] if line["product_id"] == pid))


async def _two_boxes(client: AsyncClient, h: dict[str, str], mid: str) -> tuple[str, str]:
    """Два открытых короба отгрузки (существующее «Создать короба» пачкой)."""
    created = await client.post(
        f"{BASE}/{mid}/boxes/batch", headers=h, json={"count": 2, "box_preset": "60_40_40"}
    )
    assert created.status_code == 201, created.text
    first, second = created.json()
    return str(first["id"]), str(second["id"])


async def _remove_unit(
    client: AsyncClient,
    h: dict[str, str],
    mid: str,
    box_id: str,
    line_id: str,
    *,
    mutation_id: str,
    marking_code_id: str | None = None,
) -> Response:
    """«Убрать» одну единицу из короба отгрузки (раздел 9 п. 10)."""
    body: dict[str, Any] = {"quantity": 1, "mutation_id": mutation_id}
    if marking_code_id is not None:
        body["marking_code_id"] = marking_code_id
    return await client.post(
        f"{BASE}/{mid}/boxes/{box_id}/lines/{line_id}/remove", headers=h, json=body
    )


async def _container_source(
    client: AsyncClient, h: dict[str, str], mid: str, pid: str, container_id: str
) -> dict[str, Any]:
    """Источник-тара товара на вкладке «Подбор» (`pick-options`)."""
    options = await client.get(f"{BASE}/{mid}/pick-options", headers=h)
    assert options.status_code == 200, options.text
    product = next(row for row in options.json() if row["product_id"] == pid)
    sources = [
        source
        for location in product["locations"]
        for source in location["sources"]
        if source["container_path"] and source["container_path"][-1]["id"] == container_id
    ]
    assert len(sources) == 1, product
    return cast(dict[str, Any], sources[0])


def _tara_codes(source: dict[str, Any], key: str) -> list[str]:
    """Коды тары из `pick-options` (раздел 9 п. 12) как неизменяемые части `01…21…`.

    `key` — `known_marking_codes` («точно в таре») или `uncertain_marking_codes`
    («возможно в таре»). Элемент может быть строкой кода или объектом с `cis_code`.
    """
    codes = source.get(key)
    assert isinstance(codes, list), f"нет {key} у источника: {source}"
    return sorted(
        _head(str(item["cis_code"] if isinstance(item, dict) else item)) for item in codes
    )


async def _import_pool(
    client: AsyncClient,
    h: dict[str, str],
    *,
    sid: str,
    pid: str,
    sku: str,
    codes: list[str],
) -> None:
    """Пул кодов товара через существующий импорт «Честного знака»."""
    body = "cis,sku_code\n" + "\n".join(f"{code},{sku}" for code in codes)
    imported = await client.post(
        "/operations/marking-codes/import",
        headers=h,
        data={
            "seller_id": sid,
            "pools_json": json.dumps([{"title": "WMS-686 пул", "product_ids": [pid]}]),
        },
        files=[("files", ("codes.csv", body.encode(), "text/csv"))],
    )
    assert imported.status_code == 200, imported.text
    assert imported.json()["accepted_count"] == len(codes), imported.text


async def _pool_available(client: AsyncClient, h: dict[str, str], sid: str, pid: str) -> int:
    inventory = await client.get(
        "/operations/marking-codes/inventory", headers=h, params={"seller_id": sid}
    )
    assert inventory.status_code == 200, inventory.text
    return int(
        next(row for row in inventory.json()["rows"] if row["product_id"] == pid)[
            "available_count"
        ]
    )


# ------------------------------------------------------------------- остаток в БД


async def _stock(
    pid: str,
    *,
    loc_id: str | None = None,
    container_id: str | None = None,
    loose: bool = False,
) -> int:
    """Остаток товара: весь, по ячейке, по таре или россыпью в ячейке."""
    stmt = select(func.coalesce(func.sum(InventoryBalance.quantity), 0)).where(
        InventoryBalance.product_id == uuid.UUID(pid)
    )
    if loc_id is not None:
        stmt = stmt.where(InventoryBalance.storage_location_id == uuid.UUID(loc_id))
    if container_id is not None:
        stmt = stmt.where(InventoryBalance.container_id == uuid.UUID(container_id))
    if loose:
        stmt = stmt.where(InventoryBalance.container_id.is_(None))
    async with SessionLocal() as session:
        return int(await session.scalar(stmt) or 0)


async def _sorting_stock(wid: str, pid: str) -> int:
    """Сколько штук товара лежит на системной ячейке «Сортировка» склада."""
    stmt = (
        select(func.coalesce(func.sum(InventoryBalance.quantity), 0))
        .join(StorageLocation, StorageLocation.id == InventoryBalance.storage_location_id)
        .where(
            InventoryBalance.product_id == uuid.UUID(pid),
            StorageLocation.warehouse_id == uuid.UUID(wid),
            StorageLocation.code == SORTING_LOCATION_CODE,
        )
    )
    async with SessionLocal() as session:
        return int(await session.scalar(stmt) or 0)


async def _movements(pid: str) -> list[tuple[str, int]]:
    """Все движения остатка товара (тип, изменение) в порядке записи."""
    stmt = (
        select(InventoryMovement.movement_type, InventoryMovement.quantity_delta)
        .where(InventoryMovement.product_id == uuid.UUID(pid))
        .order_by(InventoryMovement.created_at, InventoryMovement.id)
    )
    async with SessionLocal() as session:
        return [(str(kind), int(delta)) for kind, delta in (await session.execute(stmt)).all()]


# ------------------------------------------------------------- типовые сценарии


class Picking:
    """Отгрузка с товаром A в коробе K1 и россыпью в той же ячейке."""

    h: dict[str, str]
    wid: str
    sid: str
    wb_wid: int
    loc_id: str
    pid: str
    k1: str
    k1_barcode: str
    mid: str
    doc: str


async def _picking(
    client: AsyncClient,
    monkeypatch: pytest.MonkeyPatch,
    *,
    in_box: int = 5,
    loose: int = 2,
    plan: int = 5,
) -> Picking:
    ctx = Picking()
    ctx.h, ctx.wid, ctx.sid, ctx.wb_wid = await _org(client, monkeypatch)
    ctx.loc_id, _ = await _cell(client, ctx.h, ctx.wid, "A-1-1")
    ctx.pid = await _product(client, ctx.h, ctx.sid, barcode=EAN_A)
    await _receive(
        client, ctx.h, wid=ctx.wid, pid=ctx.pid, qty=in_box + loose, loc_id=ctx.loc_id
    )
    ctx.k1, ctx.k1_barcode = await _box_in_cell(
        client, ctx.h, wid=ctx.wid, loc_id=ctx.loc_id, contents={ctx.pid: in_box}
    )
    ctx.mid, ctx.doc = await _shipment(
        client, ctx.h, wid=ctx.wid, sid=ctx.sid, wb_wid=ctx.wb_wid, lines={ctx.pid: plan}
    )
    return ctx


async def _select_k1(client: AsyncClient, ctx: Picking) -> None:
    selected = await _pick_scan(client, ctx.h, ctx.mid, ctx.k1_barcode)
    assert selected.status_code == 200, selected.text
    assert selected.json()["kind"] == "container"
    assert selected.json()["container_id"] == ctx.k1


async def _scan_a_from_k1(client: AsyncClient, ctx: Picking) -> Response:
    scanned = await _pick_scan(
        client, ctx.h, ctx.mid, EAN_A, pid=ctx.pid, loc_id=ctx.loc_id, box_id=ctx.k1
    )
    assert scanned.status_code == 200, scanned.text
    assert scanned.json()["kind"] == "product"
    return scanned


async def _scan_kiz_from_k1(client: AsyncClient, ctx: Picking, code: str) -> Response:
    return await _pick_scan(
        client, ctx.h, ctx.mid, code, pid=ctx.pid, loc_id=ctx.loc_id, box_id=ctx.k1
    )


def _assert_marking(resp: Response, *, pid: str, code: str, already_linked: bool) -> None:
    """Ответ `pick/scan` / `boxes/{id}/scan` на КИЗ по разделу 9 п. 1 и 5."""
    assert resp.status_code == 200, f"КИЗ не принят: {resp.status_code} {resp.text}"
    body = resp.json()
    assert body["kind"] == "marking", f"КИЗ обработан как {body.get('kind')!r}: {body}"
    assert body["product_id"] == pid
    assert body["marking_code_id"]
    assert str(body["cis_code"]).startswith(_head(code))
    assert body["already_linked"] is already_linked


# ===================================================================== GREEN


@pytest.mark.asyncio
async def test_c03_each_product_scan_takes_one_unit_from_selected_box_to_sorting(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C3 · R2, R19: ячейка → короб K1 → ШК → ШК: две единицы из K1 на «Сортировку»."""
    ctx = await _picking(async_client, monkeypatch, in_box=5, loose=2, plan=5)
    assert await _stock(ctx.pid) == 7
    await _select_k1(async_client, ctx)

    first = await _scan_a_from_k1(async_client, ctx)
    assert first.json()["allocation_quantity"] == 1
    assert first.json()["picked_qty"] == 1
    # Тот же ШК повторно — следующая физическая единица, а не повтор первой.
    second = await _scan_a_from_k1(async_client, ctx)
    assert second.json()["allocation_quantity"] == 2
    assert second.json()["picked_qty"] == 2

    assert await _picked(async_client, ctx.h, ctx.mid, ctx.pid) == 2
    assert await _stock(ctx.pid, container_id=ctx.k1) == 3
    assert await _stock(ctx.pid, loc_id=ctx.loc_id, loose=True) == 2
    assert await _sorting_stock(ctx.wid, ctx.pid) == 2
    # Подбор — расположение: общий остаток и списания не меняются.
    assert await _stock(ctx.pid) == 7
    assert [kind for kind, _ in await _movements(ctx.pid) if kind == "marketplace_unload"] == []


@pytest.mark.asyncio
async def test_c06_product_without_honest_sign_ships_without_any_kiz(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C6 · R4: подбор → короб → «Завершить» без единого скана КИЗ."""
    h, wid, sid, wb_wid = await _org(async_client, monkeypatch)
    loc_id, _ = await _cell(async_client, h, wid, "A-1-1")
    pid = await _product(async_client, h, sid, barcode=EAN_A)
    await _receive(async_client, h, wid=wid, pid=pid, qty=4, loc_id=loc_id)
    mid, _ = await _shipment(async_client, h, wid=wid, sid=sid, wb_wid=wb_wid, lines={pid: 2})

    for _ in range(2):
        picked = await _pick_scan(async_client, h, mid, EAN_A, pid=pid, loc_id=loc_id)
        assert picked.status_code == 200, picked.text
    box_id, _ = await _new_box(async_client, h, mid)
    for _ in range(2):
        packed = await async_client.post(
            f"{BASE}/{mid}/boxes/{box_id}/scan", headers=h, json={"barcode": EAN_A}
        )
        assert packed.status_code == 200, packed.text
    assert await _box_lines(async_client, h, mid, box_id) == {pid: 2}

    shipped = await _ship(async_client, h, mid)
    assert shipped.status_code == 200, shipped.text
    assert shipped.json()["status"] == "shipped"
    assert await _stock(pid) == 2


@pytest.mark.asyncio
async def test_c11_whole_mixed_warehouse_box_attaches_once_without_opening(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C11 · R7: смешанный короб A x3 + B x2 уходит целиком; повтор не удваивает."""
    h, wid, sid, wb_wid = await _org(async_client, monkeypatch)
    loc_id, _ = await _cell(async_client, h, wid, "A-1-1")
    pid_a = await _product(async_client, h, sid, barcode=EAN_A)
    pid_b = await _product(async_client, h, sid, barcode=EAN_B, name="Футболка 50")
    await _receive(async_client, h, wid=wid, pid=pid_a, qty=4, loc_id=loc_id)
    await _receive(async_client, h, wid=wid, pid=pid_b, qty=3, loc_id=loc_id)
    box_id, box_barcode = await _box_in_cell(
        async_client, h, wid=wid, loc_id=loc_id, contents={pid_a: 3, pid_b: 2}
    )
    mid, _ = await _shipment(
        async_client, h, wid=wid, sid=sid, wb_wid=wb_wid, lines={pid_a: 3, pid_b: 2}
    )

    attached = await async_client.post(
        f"{BASE}/{mid}/boxes/attach",
        headers=h,
        json={"barcode": box_barcode, "box_preset": "60_40_40"},
    )
    assert attached.status_code == 201, attached.text
    shipment_box = attached.json()
    composition = {str(line["product_id"]): int(line["quantity"]) for line in shipment_box["lines"]}
    assert composition == {pid_a: 3, pid_b: 2}
    assert shipment_box["internal_barcode"] == box_barcode
    assert await _stock(pid_a, container_id=box_id) == 0
    assert await _stock(pid_b, container_id=box_id) == 0
    assert (await _stock(pid_a), await _stock(pid_b)) == (4, 3)

    again = await async_client.post(
        f"{BASE}/{mid}/boxes/attach",
        headers=h,
        json={"barcode": box_barcode, "box_preset": "60_40_40"},
    )
    assert again.status_code == 422, again.text
    assert again.json()["detail"] == "box_empty"
    detail = await async_client.get(f"{BASE}/{mid}", headers=h)
    assert detail.status_code == 200, detail.text
    boxes = detail.json()["boxes"]
    assert len(boxes) == 1
    assert {str(line["product_id"]): int(line["quantity"]) for line in boxes[0]["lines"]} == {
        pid_a: 3,
        pid_b: 2,
    }
    assert (await _stock(pid_a), await _stock(pid_b)) == (4, 3)


@pytest.mark.asyncio
async def test_c13_one_box_is_split_between_two_shipments_and_overdraw_is_refused(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C13 · R8: K1 с A x5 делят отгрузки X (2) и Y (3); лишняя штука из K1 — отказ."""
    h, wid, sid, wb_wid = await _org(async_client, monkeypatch)
    loc_id, _ = await _cell(async_client, h, wid, "A-1-1")
    pid = await _product(async_client, h, sid, barcode=EAN_A)
    await _receive(async_client, h, wid=wid, pid=pid, qty=8, loc_id=loc_id)
    k1, _ = await _box_in_cell(async_client, h, wid=wid, loc_id=loc_id, contents={pid: 5})
    # План Y больше, чем осталось в K1: отказ должен дать именно пустой короб, а не план.
    x, _ = await _shipment(async_client, h, wid=wid, sid=sid, wb_wid=wb_wid, lines={pid: 2})
    y, _ = await _shipment(async_client, h, wid=wid, sid=sid, wb_wid=wb_wid, lines={pid: 4})

    taken_x = await _pick_set(async_client, h, x, pid=pid, loc_id=loc_id, box_id=k1, qty=2)
    assert taken_x.status_code == 200, taken_x.text
    assert await _stock(pid, container_id=k1) == 3

    options_y = await async_client.get(f"{BASE}/{y}/pick-options", headers=h)
    assert options_y.status_code == 200, options_y.text
    location_y = next(
        row
        for row in next(o for o in options_y.json() if o["product_id"] == pid)["locations"]
        if row["storage_location_id"] == loc_id
    )
    k1_source = next(s for s in location_y["sources"] if not s["is_loose"])
    assert k1_source["container_path"][-1]["id"] == k1
    assert k1_source["available"] == 3

    taken_y = await _pick_set(async_client, h, y, pid=pid, loc_id=loc_id, box_id=k1, qty=3)
    assert taken_y.status_code == 200, taken_y.text
    assert await _stock(pid, container_id=k1) == 0

    before = await _movements(pid)
    overdraw = await _pick_set(async_client, h, y, pid=pid, loc_id=loc_id, box_id=k1, qty=4)
    assert overdraw.status_code == 422, overdraw.text
    assert overdraw.json()["detail"] == "insufficient_available"
    assert await _movements(pid) == before
    assert (await _picked(async_client, h, x, pid), await _picked(async_client, h, y, pid)) == (
        2,
        3,
    )
    assert await _stock(pid, loc_id=loc_id, loose=True) == 3
    assert await _stock(pid) == 8


@pytest.mark.asyncio
async def test_c32_seller_downloads_wb_xlsx_of_own_shipment(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C32 · R26: селлер своей учётной записью скачивает XLSX для WB своей отгрузки."""
    h, wid, sid, wb_wid = await _org(async_client, monkeypatch)
    loc_id, _ = await _cell(async_client, h, wid, "A-1-1")
    pid = await _product(async_client, h, sid, barcode=EAN_A)
    await _receive(async_client, h, wid=wid, pid=pid, qty=2, loc_id=loc_id)
    mid, _ = await _shipment(async_client, h, wid=wid, sid=sid, wb_wid=wb_wid, lines={pid: 2})
    box_id, box_barcode = await _new_box(async_client, h, mid)
    packed = await async_client.post(
        f"{BASE}/{mid}/boxes/{box_id}/scan",
        headers=h,
        json={"barcode": EAN_A, "storage_location_id": loc_id, "quantity": 2},
    )
    assert packed.status_code == 200, packed.text
    seller_h = await _seller_headers(async_client, h, sid)

    exported = await async_client.get(f"{BASE}/{mid}/wb-fbw-packaging.xlsx", headers=seller_h)
    assert exported.status_code == 200, exported.text
    assert exported.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    sheet = load_workbook(io.BytesIO(exported.content)).active
    assert [tuple(cell.value for cell in row) for row in sheet.iter_rows(min_row=2, max_col=3)] == [
        (EAN_A, 2, box_barcode)
    ]


# ======================================================================= RED


@pytest.mark.asyncio
async def test_c04_kiz_after_product_links_to_source_without_changing_quantity(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C4 · R3: K1 → ШК A → КИЗ1: связь с товаром и K1, «Снято» прежнее."""
    ctx = await _picking(async_client, monkeypatch)
    kiz1 = _kiz(EAN_A, "C4SERIAL00001")
    await _select_k1(async_client, ctx)
    await _scan_a_from_k1(async_client, ctx)
    movements = await _movements(ctx.pid)

    marked = await _scan_kiz_from_k1(async_client, ctx, kiz1)
    _assert_marking(marked, pid=ctx.pid, code=kiz1, already_linked=False)
    assert marked.json()["picked_qty"] == 1
    assert marked.json()["allocation_quantity"] == 1
    assert await _picked(async_client, ctx.h, ctx.mid, ctx.pid) == 1
    assert await _stock(ctx.pid, container_id=ctx.k1) == 4
    assert await _movements(ctx.pid) == movements

    for _ in range(2):  # перечитывание отгрузки: связь сохранена
        links = await _links(async_client, ctx.h, ctx.mid)
        assert len(links) == 1
        link = _link_for(links, kiz1)
        assert link["marking_code_id"] == marked.json()["marking_code_id"]
        assert link["product_id"] == ctx.pid
        assert link["storage_location_id"] == ctx.loc_id
        assert link["container_id"] == ctx.k1
        assert link["box_id"] is None
        assert link["intake_document_number"] is None


@pytest.mark.asyncio
async def test_c07_repeated_kiz_is_already_linked_without_unit_or_second_link(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C7 · R5а: K1 → ШК A → КИЗ1 → КИЗ1: второй ответ already_linked, связь одна."""
    ctx = await _picking(async_client, monkeypatch)
    kiz1 = _kiz(EAN_A, "C7SERIAL00001")
    await _select_k1(async_client, ctx)
    await _scan_a_from_k1(async_client, ctx)

    first = await _scan_kiz_from_k1(async_client, ctx, kiz1)
    _assert_marking(first, pid=ctx.pid, code=kiz1, already_linked=False)
    again = await _scan_kiz_from_k1(async_client, ctx, kiz1)
    _assert_marking(again, pid=ctx.pid, code=kiz1, already_linked=True)
    assert again.json()["marking_code_id"] == first.json()["marking_code_id"]

    assert await _picked(async_client, ctx.h, ctx.mid, ctx.pid) == 1
    assert len(await _links(async_client, ctx.h, ctx.mid)) == 1


@pytest.mark.asyncio
async def test_c08_kiz_beyond_picked_units_is_refused_without_quantity_change(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C8 · R5б: K1 → ШК A → КИЗ1 → КИЗ2: КИЗ2 — marking_quantity_exceeded."""
    ctx = await _picking(async_client, monkeypatch)
    kiz1, kiz2 = _kiz(EAN_A, "C8SERIAL00001"), _kiz(EAN_A, "C8SERIAL00002")
    await _select_k1(async_client, ctx)
    await _scan_a_from_k1(async_client, ctx)
    _assert_marking(
        await _scan_kiz_from_k1(async_client, ctx, kiz1),
        pid=ctx.pid,
        code=kiz1,
        already_linked=False,
    )
    movements = await _movements(ctx.pid)

    extra = await _scan_kiz_from_k1(async_client, ctx, kiz2)
    assert extra.status_code == 422, extra.text
    assert extra.json()["detail"] == "marking_quantity_exceeded"
    assert await _picked(async_client, ctx.h, ctx.mid, ctx.pid) == 1
    assert await _movements(ctx.pid) == movements
    links = await _links(async_client, ctx.h, ctx.mid)
    assert len(links) == 1
    _link_for(links, kiz1)


@pytest.mark.asyncio
async def test_c09_kiz_of_another_shipment_is_refused_and_first_link_kept(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C9 · R5в, R8: КИЗ1 привязан в X; в Y — отказ, единица Y посчитана, связь X цела."""
    ctx = await _picking(async_client, monkeypatch, in_box=5, loose=2, plan=2)
    kiz1 = _kiz(EAN_A, "C9SERIAL00001")
    y, _ = await _shipment(
        async_client, ctx.h, wid=ctx.wid, sid=ctx.sid, wb_wid=ctx.wb_wid, lines={ctx.pid: 2}
    )
    await _select_k1(async_client, ctx)
    await _scan_a_from_k1(async_client, ctx)
    in_x = await _scan_kiz_from_k1(async_client, ctx, kiz1)
    _assert_marking(in_x, pid=ctx.pid, code=kiz1, already_linked=False)

    picked_y = await _pick_scan(
        async_client, ctx.h, y, EAN_A, pid=ctx.pid, loc_id=ctx.loc_id, box_id=ctx.k1
    )
    assert picked_y.status_code == 200, picked_y.text
    in_y = await _pick_scan(
        async_client, ctx.h, y, kiz1, pid=ctx.pid, loc_id=ctx.loc_id, box_id=ctx.k1
    )
    assert in_y.status_code == 422, in_y.text
    assert in_y.json()["detail"] == "marking_code_other_shipment"

    assert await _picked(async_client, ctx.h, y, ctx.pid) == 1
    assert await _links(async_client, ctx.h, y) == []
    links_x = await _links(async_client, ctx.h, ctx.mid)
    assert len(links_x) == 1
    assert _link_for(links_x, kiz1)["marking_code_id"] == in_x.json()["marking_code_id"]


@pytest.mark.asyncio
async def test_c10_kiz_of_other_product_after_product_scan_is_refused(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C10 · R6а: ШК A → КИЗ(B) с product_id = A — marking_code_other_product, связей нет."""
    ctx = await _picking(async_client, monkeypatch)
    await _product(async_client, ctx.h, ctx.sid, barcode=EAN_B, name="Футболка 50")
    kiz_b = _kiz(EAN_B, "C10SERIAL0001")
    await _select_k1(async_client, ctx)
    await _scan_a_from_k1(async_client, ctx)
    movements = await _movements(ctx.pid)

    # Контекст единицы A (`product_id` = A), а GTIN кода — товар B (D16).
    wrong = await _scan_kiz_from_k1(async_client, ctx, kiz_b)
    assert wrong.status_code == 422, f"КИЗ товара B не отклонён: {wrong.text}"
    assert wrong.json()["detail"] == "marking_code_other_product"
    assert await _picked(async_client, ctx.h, ctx.mid, ctx.pid) == 1
    assert await _movements(ctx.pid) == movements
    assert await _links(async_client, ctx.h, ctx.mid) == []


@pytest.mark.asyncio
async def test_c14_decrease_returns_unmarked_unit_first_and_refuses_marked_one(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C14 · R9а-в: из K1 снято 2, одна с КИЗ1; уменьшение — сначала без КИЗ."""
    ctx = await _picking(async_client, monkeypatch, in_box=5, loose=0, plan=5)
    kiz1 = _kiz(EAN_A, "C14SERIAL0001")
    await _select_k1(async_client, ctx)
    await _scan_a_from_k1(async_client, ctx)
    marked = await _scan_kiz_from_k1(async_client, ctx, kiz1)
    _assert_marking(marked, pid=ctx.pid, code=kiz1, already_linked=False)
    await _scan_a_from_k1(async_client, ctx)
    assert await _stock(ctx.pid, container_id=ctx.k1) == 3

    to_one = await _pick_set(
        async_client, ctx.h, ctx.mid, pid=ctx.pid, loc_id=ctx.loc_id, box_id=ctx.k1, qty=1
    )
    assert to_one.status_code == 200, to_one.text
    assert await _stock(ctx.pid, container_id=ctx.k1) == 4
    assert _link_for(await _links(async_client, ctx.h, ctx.mid), kiz1)["container_id"] == ctx.k1

    movements = await _movements(ctx.pid)
    ambiguous = await _pick_set(
        async_client, ctx.h, ctx.mid, pid=ctx.pid, loc_id=ctx.loc_id, box_id=ctx.k1, qty=0
    )
    assert ambiguous.status_code == 422, ambiguous.text
    assert ambiguous.json()["detail"] == "marking_unit_ambiguous"
    assert await _picked(async_client, ctx.h, ctx.mid, ctx.pid) == 1
    assert await _stock(ctx.pid, container_id=ctx.k1) == 4
    assert await _movements(ctx.pid) == movements
    _link_for(await _links(async_client, ctx.h, ctx.mid), kiz1)

    removed = await async_client.delete(
        f"{BASE}/{ctx.mid}/marking-codes/{marked.json()['marking_code_id']}", headers=ctx.h
    )
    assert removed.status_code == 200, removed.text
    to_zero = await _pick_set(
        async_client, ctx.h, ctx.mid, pid=ctx.pid, loc_id=ctx.loc_id, box_id=ctx.k1, qty=0
    )
    assert to_zero.status_code == 200, to_zero.text
    assert await _picked(async_client, ctx.h, ctx.mid, ctx.pid) == 0
    assert await _stock(ctx.pid, container_id=ctx.k1) == 5


@pytest.mark.asyncio
async def test_c15_removing_kiz_keeps_quantity_and_repeat_is_noop(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C15 · R9в: ✕ КИЗ1 дважды: 200 removed=true, затем removed=false; «Снято» прежнее."""
    ctx = await _picking(async_client, monkeypatch)
    kiz1 = _kiz(EAN_A, "C15SERIAL0001")
    await _select_k1(async_client, ctx)
    await _scan_a_from_k1(async_client, ctx)
    marked = await _scan_kiz_from_k1(async_client, ctx, kiz1)
    _assert_marking(marked, pid=ctx.pid, code=kiz1, already_linked=False)
    movements = await _movements(ctx.pid)
    url = f"{BASE}/{ctx.mid}/marking-codes/{marked.json()['marking_code_id']}"

    first = await async_client.delete(url, headers=ctx.h)
    assert first.status_code == 200, first.text
    assert first.json() == {"removed": True}
    second = await async_client.delete(url, headers=ctx.h)
    assert second.status_code == 200, second.text
    assert second.json() == {"removed": False}

    assert await _links(async_client, ctx.h, ctx.mid) == []
    assert await _picked(async_client, ctx.h, ctx.mid, ctx.pid) == 1
    assert await _movements(ctx.pid) == movements


@pytest.mark.asyncio
async def test_c16_pick_scan_replay_with_same_key_takes_one_unit(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C16 · R10: ШК A с ключом M1 дважды — одна единица; новый ключ M2 — вторая."""
    ctx = await _picking(async_client, monkeypatch)
    await _select_k1(async_client, ctx)
    m1, m2 = str(uuid.uuid4()), str(uuid.uuid4())

    async def scan(key: str) -> Response:
        resp = await _pick_scan(
            async_client,
            ctx.h,
            ctx.mid,
            EAN_A,
            pid=ctx.pid,
            loc_id=ctx.loc_id,
            box_id=ctx.k1,
            mutation_id=key,
        )
        assert resp.status_code == 200, resp.text
        return resp

    original = await scan(m1)
    replay = await scan(m1)
    assert await _picked(async_client, ctx.h, ctx.mid, ctx.pid) == 1, "повтор M1 снял вторую"
    assert replay.json() == original.json()
    assert await _sorting_stock(ctx.wid, ctx.pid) == 1
    assert await _stock(ctx.pid, container_id=ctx.k1) == 4

    await scan(m2)
    assert await _picked(async_client, ctx.h, ctx.mid, ctx.pid) == 2
    assert await _sorting_stock(ctx.wid, ctx.pid) == 2


@pytest.mark.asyncio
async def test_c19_packing_scan_boxes_only_picked_units_without_stock_moves(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C19 · R13, R19: снято 2; три скана ШК в B1 — два уложено, третий отказ."""
    h, wid, sid, wb_wid = await _org(async_client, monkeypatch)
    loc_id, _ = await _cell(async_client, h, wid, "A-1-1")
    pid = await _product(async_client, h, sid, barcode=EAN_A)
    await _receive(async_client, h, wid=wid, pid=pid, qty=5, loc_id=loc_id)
    # План больше подобранного: третий скан не упирается в план, только в «нечего класть».
    mid, _ = await _shipment(async_client, h, wid=wid, sid=sid, wb_wid=wb_wid, lines={pid: 5})
    for _ in range(2):
        picked = await _pick_scan(async_client, h, mid, EAN_A, pid=pid, loc_id=loc_id)
        assert picked.status_code == 200, picked.text
    b1, _ = await _new_box(async_client, h, mid)
    movements = await _movements(pid)

    for expected in (1, 2):
        packed = await _pack_scan(async_client, h, mid, b1, EAN_A, loc_id=loc_id)
        assert packed.status_code == 200, packed.text
        assert packed.json()["kind"] == "product"
        assert packed.json()["quantity"] == expected
    nothing = await _pack_scan(async_client, h, mid, b1, EAN_A, loc_id=loc_id)
    assert nothing.status_code == 422, f"упаковка сняла товар со склада: {nothing.text}"
    assert nothing.json()["detail"] == "nothing_picked_to_pack"

    assert await _box_lines(async_client, h, mid, b1) == {pid: 2}
    assert await _picked(async_client, h, mid, pid) == 2
    assert await _movements(pid) == movements
    assert await _stock(pid) == 5


@pytest.mark.asyncio
async def test_c20_packing_kiz_links_to_current_box_and_repeat_is_idempotent(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C20 · R14: уложить A в B1, КИЗ1 в B1 дважды: связь с B1, повтор already_linked."""
    h, wid, sid, wb_wid = await _org(async_client, monkeypatch)
    loc_id, _ = await _cell(async_client, h, wid, "A-1-1")
    pid = await _product(async_client, h, sid, barcode=EAN_A)
    await _receive(async_client, h, wid=wid, pid=pid, qty=3, loc_id=loc_id)
    mid, _ = await _shipment(async_client, h, wid=wid, sid=sid, wb_wid=wb_wid, lines={pid: 3})
    picked = await _pick_scan(async_client, h, mid, EAN_A, pid=pid, loc_id=loc_id)
    assert picked.status_code == 200, picked.text
    b1, _ = await _new_box(async_client, h, mid)
    packed = await _pack_scan(async_client, h, mid, b1, EAN_A)
    assert packed.status_code == 200, packed.text
    kiz1 = _kiz(EAN_A, "C20SERIAL0001")
    movements = await _movements(pid)

    first = await _pack_scan(async_client, h, mid, b1, kiz1, pid=pid)
    _assert_marking(first, pid=pid, code=kiz1, already_linked=False)
    link = _link_for(await _links(async_client, h, mid), kiz1)
    assert link["box_id"] == b1
    assert link["product_id"] == pid

    again = await _pack_scan(async_client, h, mid, b1, kiz1, pid=pid)
    _assert_marking(again, pid=pid, code=kiz1, already_linked=True)
    assert len(await _links(async_client, h, mid)) == 1
    assert await _box_lines(async_client, h, mid, b1) == {pid: 1}
    assert await _movements(pid) == movements


@pytest.mark.asyncio
async def test_c24_scanned_kiz_counts_as_marking_so_ship_passes_without_printing(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C24 · R18: товар с ЧЗ, план 2, печати не было: ШК→КИЗ→ШК→КИЗ, уложить, «Завершить»."""
    h, wid, sid, wb_wid = await _org(async_client, monkeypatch)
    loc_id, _ = await _cell(async_client, h, wid, "A-1-1")
    pid = await _product(async_client, h, sid, barcode=EAN_A, honest_sign=True)
    await _receive(async_client, h, wid=wid, pid=pid, qty=3, loc_id=loc_id)
    mid, _ = await _shipment(async_client, h, wid=wid, sid=sid, wb_wid=wb_wid, lines={pid: 2})

    for serial in ("C24SERIAL0001", "C24SERIAL0002"):
        unit = await _pick_scan(async_client, h, mid, EAN_A, pid=pid, loc_id=loc_id)
        assert unit.status_code == 200, unit.text
        code = _kiz(EAN_A, serial)
        marked = await _pick_scan(async_client, h, mid, code, pid=pid, loc_id=loc_id)
        _assert_marking(marked, pid=pid, code=code, already_linked=False)
    b1, _ = await _new_box(async_client, h, mid)
    for _ in range(2):
        packed = await _pack_scan(async_client, h, mid, b1, EAN_A)
        assert packed.status_code == 200, packed.text

    shipped = await _ship(async_client, h, mid)
    assert shipped.status_code == 200, shipped.text
    assert shipped.json()["status"] == "shipped"


@pytest.mark.asyncio
async def test_c26_intake_kiz_shows_intake_number_external_kiz_shows_none(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C26 · R20: КИЗ1 отсканирован в приёмке N, КИЗ2 нигде — номер N и null."""

    async def no_honest_sign_check(job_id: uuid.UUID) -> None:
        # Внешняя граница: проверка кода в Честном знаке после проведения приёмки.
        del job_id

    monkeypatch.setattr(
        "app.services.inbound_marking_service.run_check_job", no_honest_sign_check
    )
    h, wid, sid, wb_wid = await _org(async_client, monkeypatch)
    loc_id, _ = await _cell(async_client, h, wid, "A-1-1")
    pid = await _product(async_client, h, sid, barcode=EAN_A)
    kiz1, kiz2 = _kiz(EAN_A, "C26SERIAL0001"), _kiz(EAN_A, "C26SERIAL0002")
    intake_number = await _receive(
        async_client, h, wid=wid, pid=pid, qty=3, loc_id=loc_id, kiz=(kiz1,)
    )
    mid, _ = await _shipment(async_client, h, wid=wid, sid=sid, wb_wid=wb_wid, lines={pid: 2})

    for code in (kiz1, kiz2):
        unit = await _pick_scan(async_client, h, mid, EAN_A, pid=pid, loc_id=loc_id)
        assert unit.status_code == 200, unit.text
        marked = await _pick_scan(async_client, h, mid, code, pid=pid, loc_id=loc_id)
        _assert_marking(marked, pid=pid, code=code, already_linked=False)

    links = await _links(async_client, h, mid)
    assert _link_for(links, kiz1)["intake_document_number"] == intake_number
    assert _link_for(links, kiz2)["intake_document_number"] is None


async def _shipped_with_kiz(
    client: AsyncClient,
    monkeypatch: pytest.MonkeyPatch,
    serial: str,
) -> tuple[dict[str, str], str, str]:
    """Отгрузка с одной единицей и КИЗ1, прошедшая «Завершить».

    Возвращает заголовки ФФ, номер документа отгрузки и КИЗ1.
    """
    h, wid, sid, wb_wid = await _org(client, monkeypatch)
    loc_id, _ = await _cell(client, h, wid, "A-1-1")
    pid = await _product(client, h, sid, barcode=EAN_A)
    await _receive(client, h, wid=wid, pid=pid, qty=3, loc_id=loc_id)
    mid, doc = await _shipment(client, h, wid=wid, sid=sid, wb_wid=wb_wid, lines={pid: 1})
    unit = await _pick_scan(client, h, mid, EAN_A, pid=pid, loc_id=loc_id)
    assert unit.status_code == 200, unit.text
    kiz1 = _kiz(EAN_A, serial)
    marked = await _pick_scan(client, h, mid, kiz1, pid=pid, loc_id=loc_id)
    _assert_marking(marked, pid=pid, code=kiz1, already_linked=False)
    b1, _ = await _new_box(client, h, mid)
    packed = await _pack_scan(client, h, mid, b1, EAN_A)
    assert packed.status_code == 200, packed.text
    shipped = await _ship(client, h, mid)
    assert shipped.status_code == 200, shipped.text
    return h, doc, kiz1


@pytest.mark.asyncio
async def test_c27_ship_writes_shipped_event_and_cancel_releases_codes(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C27 · R21: «Завершить» пишет shipped с номером; отмена освобождает КИЗ2."""
    h, wid, sid, wb_wid = await _org(async_client, monkeypatch)
    loc_id, _ = await _cell(async_client, h, wid, "A-1-1")
    pid = await _product(async_client, h, sid, barcode=EAN_A)
    await _receive(async_client, h, wid=wid, pid=pid, qty=4, loc_id=loc_id)
    x, x_doc = await _shipment(async_client, h, wid=wid, sid=sid, wb_wid=wb_wid, lines={pid: 2})
    y, _ = await _shipment(async_client, h, wid=wid, sid=sid, wb_wid=wb_wid, lines={pid: 1})
    kiz1, kiz2 = _kiz(EAN_A, "C27SERIAL0001"), _kiz(EAN_A, "C27SERIAL0002")

    unit = await _pick_scan(async_client, h, x, EAN_A, pid=pid, loc_id=loc_id)
    assert unit.status_code == 200, unit.text
    first = await _pick_scan(async_client, h, x, kiz1, pid=pid, loc_id=loc_id)
    _assert_marking(first, pid=pid, code=kiz1, already_linked=False)

    # Y ещё не отгружена: КИЗ2 привязан, отмена Y снимает связь.
    unit_y = await _pick_scan(async_client, h, y, EAN_A, pid=pid, loc_id=loc_id)
    assert unit_y.status_code == 200, unit_y.text
    in_y = await _pick_scan(async_client, h, y, kiz2, pid=pid, loc_id=loc_id)
    _assert_marking(in_y, pid=pid, code=kiz2, already_linked=False)
    cancelled = await async_client.post(f"{BASE}/{y}/cancel", headers=h)
    assert cancelled.status_code == 200, cancelled.text
    assert cancelled.json()["status"] == "cancelled"
    assert await _links(async_client, h, y) == []

    # Освобождённый КИЗ2 привязывается в другой отгрузке.
    unit2 = await _pick_scan(async_client, h, x, EAN_A, pid=pid, loc_id=loc_id)
    assert unit2.status_code == 200, unit2.text
    second = await _pick_scan(async_client, h, x, kiz2, pid=pid, loc_id=loc_id)
    _assert_marking(second, pid=pid, code=kiz2, already_linked=False)
    b1, _ = await _new_box(async_client, h, x)
    for _ in range(2):
        packed = await _pack_scan(async_client, h, x, b1, EAN_A)
        assert packed.status_code == 200, packed.text
    shipped = await _ship(async_client, h, x)
    assert shipped.status_code == 200, shipped.text

    for marked in (first, second):
        history = await async_client.get(
            f"/operations/marking-codes/codes/{marked.json()['marking_code_id']}/history",
            headers=h,
        )
        assert history.status_code == 200, history.text
        shipped_events = [row for row in history.json() if row["event_type"] == "shipped"]
        assert [row["document_number"] for row in shipped_events] == [x_doc]


@pytest.mark.asyncio
async def test_c28_honest_sign_history_csv_has_kiz_with_fbo_shipment_number(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C28 · R22: после «Завершить» CSV «Истории ЧЗ» содержит КИЗ1 и номер отгрузки."""
    h, doc, kiz1 = await _shipped_with_kiz(async_client, monkeypatch, "C28SERIAL0001")

    exported = await async_client.get("/operations/marking-codes/ledger/export", headers=h)
    assert exported.status_code == 200, exported.text
    rows = list(csv.DictReader(io.StringIO(exported.content.decode("utf-8-sig"))))
    shipped_rows = [
        row
        for row in rows
        if row["event_type"] == "shipped" and row["cis_code"].startswith(_head(kiz1))
    ]
    assert [row["document_number"] for row in shipped_rows] == [doc], rows


@pytest.mark.asyncio
async def test_c29_attached_intake_box_keeps_its_barcode_in_shipment_and_xlsx(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C29 · R23: короб приёмки INB целиком в отгрузку — его ШК в отгрузке и в XLSX."""
    h, wid, sid, wb_wid = await _org(async_client, monkeypatch)
    loc_id, _ = await _cell(async_client, h, wid, "A-1-1")
    pid = await _product(async_client, h, sid, barcode=EAN_A)

    created = await async_client.post(INBOUND, headers=h, json={"warehouse_id": wid})
    assert created.status_code == 201, created.text
    rid = str(created.json()["id"])
    line = await async_client.post(
        f"{INBOUND}/{rid}/lines", headers=h, json={"product_id": pid, "expected_qty": 4}
    )
    assert line.status_code == 201, line.text
    await set_planned_boxes(async_client, INBOUND, rid, h)
    submitted = await async_client.post(f"{INBOUND}/{rid}/submit", headers=h)
    assert submitted.status_code == 200, submitted.text
    await post_primary_accept(async_client, INBOUND, rid, h)
    intake = (await async_client.get(f"{INBOUND}/{rid}", headers=h)).json()
    inb_id = str(intake["boxes"][0]["id"])
    inb_barcode = str(intake["boxes"][0]["internal_barcode"])
    assert inb_barcode.startswith("INB-")
    await fulfill_inbound_via_box_scans(async_client, h, rid, line.json()["sku_code"], 4)
    verified = await async_client.post(f"{INBOUND}/{rid}/verify", headers=h)
    assert verified.status_code == 200, verified.text
    putaway = await async_client.post(
        f"{INBOUND}/{rid}/boxes/{inb_id}/putaway", headers=h, json={"storage_location_id": loc_id}
    )
    assert putaway.status_code == 200, putaway.text
    assert await _stock(pid, container_id=inb_id) == 4
    mid, _ = await _shipment(async_client, h, wid=wid, sid=sid, wb_wid=wb_wid, lines={pid: 4})

    attached = await async_client.post(
        f"{BASE}/{mid}/boxes/attach", headers=h, json={"barcode": inb_barcode}
    )
    assert attached.status_code == 201, attached.text
    assert sum(int(row["quantity"]) for row in attached.json()["lines"]) == 4
    assert attached.json()["internal_barcode"] == inb_barcode, "ШК короба приёмки потерян"
    detail = await async_client.get(f"{BASE}/{mid}", headers=h)
    assert [box["internal_barcode"] for box in detail.json()["boxes"]] == [inb_barcode]

    exported = await async_client.get(f"{BASE}/{mid}/wb-fbw-packaging.xlsx", headers=h)
    assert exported.status_code == 200, exported.text
    sheet = load_workbook(io.BytesIO(exported.content)).active
    assert [tuple(cell.value for cell in row) for row in sheet.iter_rows(min_row=2, max_col=3)] == [
        (EAN_A, 4, inb_barcode)
    ]


@pytest.mark.asyncio
async def test_c31_pass_attach_download_replace_by_ff_and_seller_ship_without_pass(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C31 · R25: пропуск — прикрепить, скачать, заменить; «Завершить» без пропуска."""
    h, wid, sid, wb_wid = await _org(async_client, monkeypatch)
    loc_id, _ = await _cell(async_client, h, wid, "A-1-1")
    pid = await _product(async_client, h, sid, barcode=EAN_A)
    await _receive(async_client, h, wid=wid, pid=pid, qty=3, loc_id=loc_id)
    x, _ = await _shipment(async_client, h, wid=wid, sid=sid, wb_wid=wb_wid, lines={pid: 1})
    other, _ = await _shipment(async_client, h, wid=wid, sid=sid, wb_wid=wb_wid, lines={pid: 1})
    seller_h = await _seller_headers(async_client, h, sid)
    pdf = b"%PDF-1.4\n% WMS-686 pass\n%%EOF\n"
    png = b"\x89PNG\r\n\x1a\nWMS-686-replaced-pass"
    jpeg = b"\xff\xd8\xff\xe0WMS-686-seller-pass\xff\xd9"

    attached = await async_client.put(
        f"{BASE}/{x}/pass", headers=h, files={"file": ("pass.pdf", pdf, "application/pdf")}
    )
    assert attached.status_code == 200, attached.text
    assert attached.json() == {
        "filename": "pass.pdf",
        "content_type": "application/pdf",
        "size": len(pdf),
    }
    seller_copy = await async_client.get(f"{BASE}/{x}/pass", headers=seller_h)
    assert seller_copy.status_code == 200, seller_copy.text
    assert seller_copy.content == pdf

    replaced = await async_client.put(
        f"{BASE}/{x}/pass", headers=h, files={"file": ("pass.png", png, "image/png")}
    )
    assert replaced.status_code == 200, replaced.text
    assert replaced.json()["filename"] == "pass.png"
    assert (await async_client.get(f"{BASE}/{x}/pass", headers=seller_h)).content == png

    by_seller = await async_client.put(
        f"{BASE}/{x}/pass", headers=seller_h, files={"file": ("pass.jpg", jpeg, "image/jpeg")}
    )
    assert by_seller.status_code == 200, by_seller.text
    ff_copy = await async_client.get(f"{BASE}/{x}/pass", headers=h)
    assert ff_copy.status_code == 200, ff_copy.text
    assert ff_copy.content == jpeg

    # Пропуск не блокирует работу: другая отгрузка без пропуска завершается.
    unit = await _pick_scan(async_client, h, other, EAN_A, pid=pid, loc_id=loc_id)
    assert unit.status_code == 200, unit.text
    b1, _ = await _new_box(async_client, h, other)
    packed = await async_client.post(
        f"{BASE}/{other}/boxes/{b1}/scan", headers=h, json={"barcode": EAN_A}
    )
    assert packed.status_code == 200, packed.text
    shipped = await _ship(async_client, h, other)
    assert shipped.status_code == 200, shipped.text
    assert shipped.json()["status"] == "shipped"



# ===================================== редакции 2-3: R28-R33, D16, D17 (C37-C57)
#
# Контекст КИЗ (D16, раздел 9 п. 1 и 9): запрос с `product_id` — контекст единицы
# последнего ШК; без `product_id` — контекст проверки выбранного источника
# (`pick/scan` с тарой) или короба (`boxes/{box}/scan`).


async def _select(client: AsyncClient, h: dict[str, str], mid: str, barcode: str) -> None:
    """Скан тары на «Подборе»: выбирает источник для следующих сканов."""
    selected = await _pick_scan(client, h, mid, barcode)
    assert selected.status_code == 200, selected.text
    assert selected.json()["kind"] == "container"


async def _pick_unit(
    client: AsyncClient,
    h: dict[str, str],
    mid: str,
    pid: str,
    barcode: str,
    loc_id: str,
    box_id: str | None = None,
    kiz: str | None = None,
) -> str | None:
    """ШК товара из источника и, если дан, его КИЗ в контексте этой единицы.

    Возвращает id привязанного кода.
    """
    unit = await _pick_scan(client, h, mid, barcode, pid=pid, loc_id=loc_id, box_id=box_id)
    assert unit.status_code == 200, unit.text
    assert unit.json()["kind"] == "product"
    if kiz is None:
        return None
    marked = await _pick_scan(client, h, mid, kiz, pid=pid, loc_id=loc_id, box_id=box_id)
    _assert_marking(marked, pid=pid, code=kiz, already_linked=False)
    return str(marked.json()["marking_code_id"])


async def _pack_unit(
    client: AsyncClient,
    h: dict[str, str],
    mid: str,
    box_id: str,
    pid: str,
    barcode: str,
    kiz: str | None = None,
    *,
    already_linked: bool = False,
) -> str | None:
    """ШК товара на «Упаковке» и, если дан, КИЗ этой единицы (контекст единицы)."""
    packed = await _pack_scan(client, h, mid, box_id, barcode)
    assert packed.status_code == 200, packed.text
    assert packed.json()["kind"] == "product"
    if kiz is None:
        return None
    marked = await _pack_scan(client, h, mid, box_id, kiz, pid=pid)
    _assert_marking(marked, pid=pid, code=kiz, already_linked=already_linked)
    return str(marked.json()["marking_code_id"])


@pytest.mark.asyncio
async def test_c37_remove_marked_unit_from_box_returns_it_to_its_source_once(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C37 · R28: «Убрать» → «КИЗ k1»: в B1 и в подборе на 1 меньше, k1 свободен, в K1."""
    ctx = await _picking(async_client, monkeypatch, in_box=5, loose=2, plan=5)
    k1, k2 = _kiz(EAN_A, "C37SERIAL0001"), _kiz(EAN_A, "C37SERIAL0002")
    await _select_k1(async_client, ctx)
    k1_id = await _pick_unit(
        async_client, ctx.h, ctx.mid, ctx.pid, EAN_A, ctx.loc_id, ctx.k1, kiz=k1
    )
    await _pick_unit(async_client, ctx.h, ctx.mid, ctx.pid, EAN_A, ctx.loc_id, ctx.k1, kiz=k2)
    b1, _ = await _new_box(async_client, ctx.h, ctx.mid)
    # Коды подбора подтверждаются в коробе после ШК своей единицы (R29б).
    for code in (k1, k2):
        await _pack_unit(
            async_client, ctx.h, ctx.mid, b1, ctx.pid, EAN_A, code, already_linked=True
        )
    assert k1_id is not None
    assert await _stock(ctx.pid, container_id=ctx.k1) == 3
    line_id = await _box_line_id(async_client, ctx.h, ctx.mid, b1, ctx.pid)
    key = str(uuid.uuid4())

    removed = await _remove_unit(
        async_client, ctx.h, ctx.mid, b1, line_id, mutation_id=key, marking_code_id=k1_id
    )
    assert removed.status_code == 200, removed.text
    assert removed.json()["returned_to"] == {
        "storage_location_id": ctx.loc_id,
        "container_id": ctx.k1,
    }
    assert removed.json()["source_known"] is True
    assert await _box_lines(async_client, ctx.h, ctx.mid, b1) == {ctx.pid: 1}
    assert await _picked(async_client, ctx.h, ctx.mid, ctx.pid) == 1
    assert await _stock(ctx.pid, container_id=ctx.k1) == 4
    links = await _links(async_client, ctx.h, ctx.mid)
    assert len(links) == 1
    assert _link_for(links, k2)["box_id"] == b1

    replay = await _remove_unit(
        async_client, ctx.h, ctx.mid, b1, line_id, mutation_id=key, marking_code_id=k1_id
    )
    assert replay.status_code == 200, replay.text
    assert await _box_lines(async_client, ctx.h, ctx.mid, b1) == {ctx.pid: 1}
    assert await _picked(async_client, ctx.h, ctx.mid, ctx.pid) == 1
    assert await _stock(ctx.pid, container_id=ctx.k1) == 4
    assert await _stock(ctx.pid) == 7


@pytest.mark.asyncio
async def test_c38_remove_unmarked_unit_goes_to_its_source_then_all_marked_is_ambiguous(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C38 · R28, D14: B1 = A с КИЗ (из K1) + A без КИЗ (из K2); «Убрать» без кода."""
    h, wid, sid, wb_wid = await _org(async_client, monkeypatch)
    loc_id, _ = await _cell(async_client, h, wid, "A-1-1")
    pid = await _product(async_client, h, sid, barcode=EAN_A)
    await _receive(async_client, h, wid=wid, pid=pid, qty=4, loc_id=loc_id)
    k1_box, k1_barcode = await _box_in_cell(
        async_client, h, wid=wid, loc_id=loc_id, contents={pid: 2}
    )
    k2_box, k2_barcode = await _box_in_cell(
        async_client, h, wid=wid, loc_id=loc_id, contents={pid: 2}
    )
    mid, _ = await _shipment(async_client, h, wid=wid, sid=sid, wb_wid=wb_wid, lines={pid: 2})
    kiz1 = _kiz(EAN_A, "C38SERIAL0001")
    await _select(async_client, h, mid, k1_barcode)
    await _pick_unit(async_client, h, mid, pid, EAN_A, loc_id, k1_box, kiz=kiz1)
    await _select(async_client, h, mid, k2_barcode)
    await _pick_unit(async_client, h, mid, pid, EAN_A, loc_id, k2_box)
    b1, _ = await _new_box(async_client, h, mid)
    await _pack_unit(async_client, h, mid, b1, pid, EAN_A, kiz1, already_linked=True)
    await _pack_unit(async_client, h, mid, b1, pid, EAN_A)
    line_id = await _box_line_id(async_client, h, mid, b1, pid)

    unmarked = await _remove_unit(async_client, h, mid, b1, line_id, mutation_id=str(uuid.uuid4()))
    assert unmarked.status_code == 200, unmarked.text
    assert unmarked.json()["returned_to"] == {"storage_location_id": loc_id, "container_id": k2_box}
    assert unmarked.json()["source_known"] is False
    assert await _stock(pid, container_id=k2_box) == 2
    assert await _stock(pid, container_id=k1_box) == 1
    assert await _box_lines(async_client, h, mid, b1) == {pid: 1}
    assert _link_for(await _links(async_client, h, mid), kiz1)["box_id"] == b1

    ambiguous = await _remove_unit(
        async_client, h, mid, b1, line_id, mutation_id=str(uuid.uuid4())
    )
    assert ambiguous.status_code == 422, ambiguous.text
    assert ambiguous.json()["detail"] == "marking_unit_ambiguous"
    assert await _box_lines(async_client, h, mid, b1) == {pid: 1}
    assert await _picked(async_client, h, mid, pid) == 1
    assert _link_for(await _links(async_client, h, mid), kiz1)["box_id"] == b1


@pytest.mark.asyncio
async def test_c40_kiz_goes_only_to_the_box_in_url_and_not_to_a_neighbour_box(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C40 · R14, R15: B1 и B2 по A без КИЗ; КИЗ в контексте проверки только своего короба."""
    h, wid, sid, wb_wid = await _org(async_client, monkeypatch)
    loc_id, _ = await _cell(async_client, h, wid, "A-1-1")
    pid = await _product(async_client, h, sid, barcode=EAN_A)
    await _receive(async_client, h, wid=wid, pid=pid, qty=3, loc_id=loc_id)
    mid, _ = await _shipment(async_client, h, wid=wid, sid=sid, wb_wid=wb_wid, lines={pid: 2})
    for _ in range(2):
        await _pick_unit(async_client, h, mid, pid, EAN_A, loc_id)
    b1, b2 = await _two_boxes(async_client, h, mid)
    for box in (b1, b2):
        await _pack_unit(async_client, h, mid, box, pid, EAN_A)
    kiz1, kiz2, kiz3 = (_kiz(EAN_A, f"C40SERIAL000{n}") for n in (1, 2, 3))

    # Текущий короб задан URL; КИЗ без ШК — контекст проверки этого короба (R6б).
    in_b2 = await _pack_scan(async_client, h, mid, b2, kiz1)
    _assert_marking(in_b2, pid=pid, code=kiz1, already_linked=False)
    in_b1 = await _pack_scan(async_client, h, mid, b1, kiz2)
    _assert_marking(in_b1, pid=pid, code=kiz2, already_linked=False)
    no_unit = await _pack_scan(async_client, h, mid, b2, kiz3)
    assert no_unit.status_code == 422, no_unit.text
    assert no_unit.json()["detail"] == "marking_unit_not_in_box"

    links = await _links(async_client, h, mid)
    assert len(links) == 2
    assert _link_for(links, kiz1)["box_id"] == b2
    assert _link_for(links, kiz2)["box_id"] == b1
    assert await _box_lines(async_client, h, mid, b1) == {pid: 1}
    assert await _box_lines(async_client, h, mid, b2) == {pid: 1}


@pytest.mark.asyncio
async def test_c41_kiz_of_earlier_product_after_other_product_scan_is_refused(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C41 · R6, D16: K1 → ШК A → ШК B → КИЗ(A) — отказ; затем K1 → КИЗ(A) — к A без +1."""
    h, wid, sid, wb_wid = await _org(async_client, monkeypatch)
    loc_id, _ = await _cell(async_client, h, wid, "A-1-1")
    pid_a = await _product(async_client, h, sid, barcode=EAN_A)
    pid_b = await _product(async_client, h, sid, barcode=EAN_B, name="Футболка 50")
    await _receive(async_client, h, wid=wid, pid=pid_a, qty=2, loc_id=loc_id)
    await _receive(async_client, h, wid=wid, pid=pid_b, qty=2, loc_id=loc_id)
    k1, k1_barcode = await _box_in_cell(
        async_client, h, wid=wid, loc_id=loc_id, contents={pid_a: 2, pid_b: 2}
    )
    mid, _ = await _shipment(
        async_client, h, wid=wid, sid=sid, wb_wid=wb_wid, lines={pid_a: 1, pid_b: 1}
    )
    await _select(async_client, h, mid, k1_barcode)
    await _pick_unit(async_client, h, mid, pid_a, EAN_A, loc_id, k1)
    await _pick_unit(async_client, h, mid, pid_b, EAN_B, loc_id, k1)
    kiz_a = _kiz(EAN_A, "C41SERIAL0001")
    movements = await _movements(pid_a)

    # Контекст единицы B: код товара A не привязывается ни к B, ни к ранее снятой A.
    refused = await _pick_scan(
        async_client, h, mid, kiz_a, pid=pid_b, loc_id=loc_id, box_id=k1
    )
    assert refused.status_code == 422, f"КИЗ(A) после ШК B не отклонён: {refused.text}"
    assert refused.json()["detail"] == "marking_code_other_product"
    assert await _links(async_client, h, mid) == []

    # Контекст проверки источника K1: тот же код — к уже посчитанной единице A.
    await _select(async_client, h, mid, k1_barcode)
    checked = await _pick_scan(async_client, h, mid, kiz_a, loc_id=loc_id, box_id=k1)
    _assert_marking(checked, pid=pid_a, code=kiz_a, already_linked=False)
    link = _link_for(await _links(async_client, h, mid), kiz_a)
    assert link["product_id"] == pid_a
    assert link["container_id"] == k1
    assert await _picked(async_client, h, mid, pid_a) == 1
    assert await _picked(async_client, h, mid, pid_b) == 1
    assert await _movements(pid_a) == movements


@pytest.mark.asyncio
async def test_c42_picked_kiz_is_confirmed_in_box_without_new_unit_or_link(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C42 · R29: k1 подбора → ШК в B1 (короб не указан) → k1 в B1 → повтор → k1 в B2."""
    ctx = await _picking(async_client, monkeypatch, in_box=5, loose=2, plan=5)
    k1 = _kiz(EAN_A, "C42SERIAL0001")
    await _select_k1(async_client, ctx)
    await _pick_unit(async_client, ctx.h, ctx.mid, ctx.pid, EAN_A, ctx.loc_id, ctx.k1, kiz=k1)
    await _pick_unit(async_client, ctx.h, ctx.mid, ctx.pid, EAN_A, ctx.loc_id, ctx.k1)
    b1, b2 = await _two_boxes(async_client, ctx.h, ctx.mid)
    await _pack_unit(async_client, ctx.h, ctx.mid, b1, ctx.pid, EAN_A)
    assert _link_for(await _links(async_client, ctx.h, ctx.mid), k1)["box_id"] is None
    movements = await _movements(ctx.pid)

    confirmed = await _pack_scan(async_client, ctx.h, ctx.mid, b1, k1, pid=ctx.pid)
    _assert_marking(confirmed, pid=ctx.pid, code=k1, already_linked=True)
    links = await _links(async_client, ctx.h, ctx.mid)
    assert len(links) == 1
    assert _link_for(links, k1)["box_id"] == b1
    assert _link_for(links, k1)["container_id"] == ctx.k1
    assert await _box_lines(async_client, ctx.h, ctx.mid, b1) == {ctx.pid: 1}

    again = await _pack_scan(async_client, ctx.h, ctx.mid, b1, k1, pid=ctx.pid)
    _assert_marking(again, pid=ctx.pid, code=k1, already_linked=True)
    assert len(await _links(async_client, ctx.h, ctx.mid)) == 1

    # В B2 есть своя непомеченная единица A — но k1 числится в B1 и молча не переносится.
    await _pack_unit(async_client, ctx.h, ctx.mid, b2, ctx.pid, EAN_A)
    other_box = await _pack_scan(async_client, ctx.h, ctx.mid, b2, k1, pid=ctx.pid)
    assert other_box.status_code == 422, other_box.text
    assert other_box.json()["detail"] == "marking_code_other_box"
    assert _link_for(await _links(async_client, ctx.h, ctx.mid), k1)["box_id"] == b1
    assert await _box_lines(async_client, ctx.h, ctx.mid, b1) == {ctx.pid: 1}
    assert await _box_lines(async_client, ctx.h, ctx.mid, b2) == {ctx.pid: 1}
    assert await _picked(async_client, ctx.h, ctx.mid, ctx.pid) == 2
    assert await _movements(ctx.pid) == movements


@pytest.mark.asyncio
async def test_c43_print_marking_takes_one_pool_code_per_unit_and_respects_picked_codes(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C43 · R16б, R29г: печать ЧЗ из пула на единицу в коробе; повтор ключа; отказы."""
    h, wid, sid, wb_wid = await _org(async_client, monkeypatch)
    loc_id, _ = await _cell(async_client, h, wid, "A-1-1")
    sku = f"SKU-{uuid.uuid4().hex[:8]}"
    pid = await _product(async_client, h, sid, barcode=EAN_A, honest_sign=True, sku=sku)
    await _receive(async_client, h, wid=wid, pid=pid, qty=4, loc_id=loc_id)
    pool_codes = [_kiz(EAN_A, "C43POOL00001"), _kiz(EAN_A, "C43POOL00002")]
    await _import_pool(async_client, h, sid=sid, pid=pid, sku=sku, codes=pool_codes)
    assert await _pool_available(async_client, h, sid, pid) == 2
    mid, _ = await _shipment(async_client, h, wid=wid, sid=sid, wb_wid=wb_wid, lines={pid: 4})
    b1, _ = await _new_box(async_client, h, mid)
    pool_heads = {_head(code) for code in pool_codes}

    async def print_marking(key: str) -> Response:
        return await async_client.post(
            f"{BASE}/{mid}/boxes/{b1}/print-marking",
            headers=h,
            json={"product_id": pid, "mutation_id": key},
        )

    # Единица без кодов подбора: один код из пула, повтор ключа — тот же код.
    await _pick_unit(async_client, h, mid, pid, EAN_A, loc_id)
    await _pack_unit(async_client, h, mid, b1, pid, EAN_A)
    key = str(uuid.uuid4())
    printed = await print_marking(key)
    assert printed.status_code == 200, printed.text
    code_id, cis = printed.json()["marking_code_id"], str(printed.json()["cis_code"])
    assert _head(cis) in pool_heads
    replay = await print_marking(key)
    assert replay.status_code == 200, replay.text
    assert (replay.json()["marking_code_id"], replay.json()["cis_code"]) == (code_id, cis)
    assert await _pool_available(async_client, h, sid, pid) == 1
    link = _link_for(await _links(async_client, h, mid), cis)
    assert link["box_id"] == b1
    assert link["marking_code_id"] == code_id

    # Код подбора k1 ещё не подтверждён в коробе: новый код не выдаётся.
    k1 = _kiz(EAN_A, "C43SERIAL0001")
    await _pick_unit(async_client, h, mid, pid, EAN_A, loc_id, kiz=k1)
    await _pack_unit(async_client, h, mid, b1, pid, EAN_A)
    may_be_marked = await print_marking(str(uuid.uuid4()))
    assert may_be_marked.status_code == 422, may_be_marked.text
    assert may_be_marked.json()["detail"] == "unit_may_be_marked"
    assert await _pool_available(async_client, h, sid, pid) == 1
    confirmed = await _pack_scan(async_client, h, mid, b1, k1, pid=pid)
    _assert_marking(confirmed, pid=pid, code=k1, already_linked=True)

    # Последний код пула уходит третьей единице; четвёртой — пул пуст, единица упакована.
    for expected in ("printed", "empty"):
        await _pick_unit(async_client, h, mid, pid, EAN_A, loc_id)
        await _pack_unit(async_client, h, mid, b1, pid, EAN_A)
        result = await print_marking(str(uuid.uuid4()))
        if expected == "printed":
            assert result.status_code == 200, result.text
            assert await _pool_available(async_client, h, sid, pid) == 0
        else:
            assert result.status_code == 422, result.text
            assert result.json()["detail"] == "marking_pool_empty"
    assert await _box_lines(async_client, h, mid, b1) == {pid: 4}
    assert len(await _links(async_client, h, mid)) == 3


async def _return_k1_via_box(
    client: AsyncClient, ctx: Picking, serial: str
) -> tuple[str, str]:
    """k1 снят из K1 с КИЗ, уложен в B1 и возвращён в K1 по R28 (C37).

    Возвращает КИЗ k1 и его id.
    """
    k1 = _kiz(EAN_A, serial)
    await _select_k1(client, ctx)
    k1_id = await _pick_unit(client, ctx.h, ctx.mid, ctx.pid, EAN_A, ctx.loc_id, ctx.k1, kiz=k1)
    assert k1_id is not None
    b1, _ = await _new_box(client, ctx.h, ctx.mid)
    await _pack_unit(client, ctx.h, ctx.mid, b1, ctx.pid, EAN_A, k1, already_linked=True)
    line_id = await _box_line_id(client, ctx.h, ctx.mid, b1, ctx.pid)
    returned = await _remove_unit(
        client, ctx.h, ctx.mid, b1, line_id, mutation_id=str(uuid.uuid4()), marking_code_id=k1_id
    )
    assert returned.status_code == 200, returned.text
    assert returned.json()["source_known"] is True
    assert returned.json()["returned_to"]["container_id"] == ctx.k1
    return k1, k1_id


@pytest.mark.asyncio
async def test_c46_exactly_known_code_is_shown_and_moves_with_whole_box(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C46 · R30а, R30г: k1 возвращён в K1 — «точно в таре»; «целиком» K1 переносит k1."""
    ctx = await _picking(async_client, monkeypatch, in_box=5, loose=2, plan=2)
    k1, _ = await _return_k1_via_box(async_client, ctx, "C46SERIAL0001")
    assert await _stock(ctx.pid, container_id=ctx.k1) == 5
    z, _ = await _shipment(
        async_client, ctx.h, wid=ctx.wid, sid=ctx.sid, wb_wid=ctx.wb_wid, lines={ctx.pid: 5}
    )

    source = await _container_source(async_client, ctx.h, z, ctx.pid, ctx.k1)
    assert _tara_codes(source, "known_marking_codes") == [_head(k1)]
    assert _tara_codes(source, "uncertain_marking_codes") == []

    attached = await async_client.post(
        f"{BASE}/{z}/boxes/attach", headers=ctx.h, json={"barcode": ctx.k1_barcode}
    )
    assert attached.status_code == 201, attached.text
    z_box = str(attached.json()["id"])
    link = _link_for(await _links(async_client, ctx.h, z), k1)
    assert link["box_id"] == z_box
    assert link["product_id"] == ctx.pid
    assert await _links(async_client, ctx.h, ctx.mid) == []


@pytest.mark.asyncio
async def test_c47_one_box_split_by_exact_codes_and_unknown_rest_is_not_invented(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C47 · R30, R5в: из K1 X берёт {k1,k2}, Y — {k3,k4,k5}; k1 в Y отклонён."""
    ctx = await _picking(async_client, monkeypatch, in_box=6, loose=0, plan=2)
    y, _ = await _shipment(
        async_client, ctx.h, wid=ctx.wid, sid=ctx.sid, wb_wid=ctx.wb_wid, lines={ctx.pid: 3}
    )
    k1, k2, k3, k4, k5 = (_kiz(EAN_A, f"C47SERIAL000{n}") for n in range(1, 6))
    await _select_k1(async_client, ctx)
    for code in (k1, k2):
        await _pick_unit(
            async_client, ctx.h, ctx.mid, ctx.pid, EAN_A, ctx.loc_id, ctx.k1, kiz=code
        )
    await _select(async_client, ctx.h, y, ctx.k1_barcode)
    for code in (k3, k4):
        await _pick_unit(async_client, ctx.h, y, ctx.pid, EAN_A, ctx.loc_id, ctx.k1, kiz=code)
    await _pick_unit(async_client, ctx.h, y, ctx.pid, EAN_A, ctx.loc_id, ctx.k1)
    # Проверка источника K1 в Y: есть непомеченная единица, но k1 уже уехал в X.
    await _select(async_client, ctx.h, y, ctx.k1_barcode)
    crossed = await _pick_scan(async_client, ctx.h, y, k1, loc_id=ctx.loc_id, box_id=ctx.k1)
    assert crossed.status_code == 422, crossed.text
    assert crossed.json()["detail"] == "marking_code_other_shipment"
    last = await _pick_scan(async_client, ctx.h, y, k5, loc_id=ctx.loc_id, box_id=ctx.k1)
    _assert_marking(last, pid=ctx.pid, code=k5, already_linked=False)

    def heads(links: list[dict[str, Any]]) -> set[str]:
        return {_head(str(link["cis_code"])) for link in links}

    assert heads(await _links(async_client, ctx.h, ctx.mid)) == {_head(k1), _head(k2)}
    assert heads(await _links(async_client, ctx.h, y)) == {_head(k3), _head(k4), _head(k5)}
    assert await _stock(ctx.pid, container_id=ctx.k1) == 1
    source = await _container_source(async_client, ctx.h, y, ctx.pid, ctx.k1)
    assert _tara_codes(source, "known_marking_codes") == []
    assert _tara_codes(source, "uncertain_marking_codes") == []


@pytest.mark.asyncio
async def test_c49_other_organization_gets_404_on_shipment_kiz_links(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C49 · R33а: своя организация видит связи КИЗ (200), чужая — 404 на чтение и удаление."""
    ctx = await _picking(async_client, monkeypatch)
    await _select_k1(async_client, ctx)
    await _scan_a_from_k1(async_client, ctx)
    assert await _links(async_client, ctx.h, ctx.mid) == []

    kiz1 = _kiz(EAN_A, "C49SERIAL0001")
    marked = await _scan_kiz_from_k1(async_client, ctx, kiz1)
    _assert_marking(marked, pid=ctx.pid, code=kiz1, already_linked=False)
    foreign_h = await _register(async_client)
    url = f"{BASE}/{ctx.mid}/marking-codes"

    foreign_read = await async_client.get(url, headers=foreign_h)
    assert foreign_read.status_code == 404, foreign_read.text
    foreign_delete = await async_client.delete(
        f"{url}/{marked.json()['marking_code_id']}", headers=foreign_h
    )
    assert foreign_delete.status_code == 404, foreign_delete.text
    assert _link_for(await _links(async_client, ctx.h, ctx.mid), kiz1)["marking_code_id"] == (
        marked.json()["marking_code_id"]
    )


@pytest.mark.asyncio
@pytest.mark.postgresql_concurrency
async def test_c50_same_kiz_in_two_shipments_at_once_has_one_winner(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C50 · R33б: два одновременных `pick/scan` одного КИЗ в X и Y — ровно один 200."""
    h, wid, sid, wb_wid = await _org(async_client, monkeypatch)
    loc_id, _ = await _cell(async_client, h, wid, "A-1-1")
    pid = await _product(async_client, h, sid, barcode=EAN_A)
    await _receive(async_client, h, wid=wid, pid=pid, qty=4, loc_id=loc_id)
    x, _ = await _shipment(async_client, h, wid=wid, sid=sid, wb_wid=wb_wid, lines={pid: 2})
    y, _ = await _shipment(async_client, h, wid=wid, sid=sid, wb_wid=wb_wid, lines={pid: 2})
    for mid in (x, y):
        await _pick_unit(async_client, h, mid, pid, EAN_A, loc_id)
    kiz1 = _kiz(EAN_A, "C50SERIAL0001")

    results = await asyncio.wait_for(
        asyncio.gather(
            _pick_scan(async_client, h, x, kiz1, pid=pid, loc_id=loc_id),
            _pick_scan(async_client, h, y, kiz1, pid=pid, loc_id=loc_id),
            return_exceptions=True,
        ),
        timeout=60,
    )
    responses = [r for r in results if isinstance(r, Response)]
    assert len(responses) == 2, f"сервер не ответил: {results!r}"
    by_status = sorted(responses, key=lambda r: r.status_code)
    assert [r.status_code for r in by_status] == [200, 422], [r.text for r in by_status]
    winner, loser = by_status
    assert winner.json()["kind"] == "marking", winner.text
    assert loser.json()["detail"] == "marking_code_other_shipment"
    all_links = await _links(async_client, h, x) + await _links(async_client, h, y)
    assert len(all_links) == 1
    assert (await _picked(async_client, h, x, pid), await _picked(async_client, h, y, pid)) == (
        1,
        1,
    )


@pytest.mark.asyncio
@pytest.mark.postgresql_concurrency
async def test_c51_last_unit_of_source_taken_at_once_by_two_shipments_goes_to_one(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C51 · R33в: в K1 одна штука A; X и Y одновременно снимают её — успешна одна."""
    h, wid, sid, wb_wid = await _org(async_client, monkeypatch)
    loc_id, _ = await _cell(async_client, h, wid, "A-1-1")
    pid = await _product(async_client, h, sid, barcode=EAN_A)
    # Россыпь в ячейке нужна только для резерва второй отгрузки; источник — K1 с одной штукой.
    await _receive(async_client, h, wid=wid, pid=pid, qty=2, loc_id=loc_id)
    k1, _ = await _box_in_cell(async_client, h, wid=wid, loc_id=loc_id, contents={pid: 1})
    x, _ = await _shipment(async_client, h, wid=wid, sid=sid, wb_wid=wb_wid, lines={pid: 1})
    y, _ = await _shipment(async_client, h, wid=wid, sid=sid, wb_wid=wb_wid, lines={pid: 1})

    results = await asyncio.wait_for(
        asyncio.gather(
            _pick_scan(async_client, h, x, EAN_A, pid=pid, loc_id=loc_id, box_id=k1),
            _pick_scan(async_client, h, y, EAN_A, pid=pid, loc_id=loc_id, box_id=k1),
            return_exceptions=True,
        ),
        timeout=60,
    )
    responses = [r for r in results if isinstance(r, Response)]
    assert len(responses) == 2, f"сервер не ответил: {results!r}"
    by_status = sorted(responses, key=lambda r: r.status_code)
    assert [r.status_code for r in by_status] == [200, 422], [r.text for r in by_status]
    assert by_status[1].json()["detail"] == "insufficient_available"
    assert await _stock(pid, container_id=k1) == 0
    async with SessionLocal() as session:
        negative = await session.scalar(
            select(func.count(InventoryBalance.id)).where(
                InventoryBalance.product_id == uuid.UUID(pid), InventoryBalance.quantity < 0
            )
        )
    assert negative == 0
    assert await _picked(async_client, h, x, pid) + await _picked(async_client, h, y, pid) == 1
    assert await _sorting_stock(wid, pid) == 1
    assert await _stock(pid) == 2


@pytest.mark.asyncio
async def test_c54_check_context_links_kiz_to_counted_unit_without_new_unit(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C54 · R6б, R14: скан K1 → КИЗ без ШК; скан короба B1 → КИЗ без ШК; без +1."""
    ctx = await _picking(async_client, monkeypatch, in_box=5, loose=2, plan=5)
    await _select_k1(async_client, ctx)
    for _ in range(2):
        await _pick_unit(async_client, ctx.h, ctx.mid, ctx.pid, EAN_A, ctx.loc_id, ctx.k1)
    b1, _ = await _new_box(async_client, ctx.h, ctx.mid)
    await _pack_unit(async_client, ctx.h, ctx.mid, b1, ctx.pid, EAN_A)
    kiz1, kiz2 = _kiz(EAN_A, "C54SERIAL0001"), _kiz(EAN_A, "C54SERIAL0002")
    movements = await _movements(ctx.pid)

    # «Подбор»: контекст проверки источника K1.
    await _select_k1(async_client, ctx)
    in_source = await _pick_scan(
        async_client, ctx.h, ctx.mid, kiz1, loc_id=ctx.loc_id, box_id=ctx.k1
    )
    _assert_marking(in_source, pid=ctx.pid, code=kiz1, already_linked=False)
    # «Упаковка»: контекст проверки короба B1.
    in_box = await _pack_scan(async_client, ctx.h, ctx.mid, b1, kiz2)
    _assert_marking(in_box, pid=ctx.pid, code=kiz2, already_linked=False)

    links = await _links(async_client, ctx.h, ctx.mid)
    assert _link_for(links, kiz1)["container_id"] == ctx.k1
    assert _link_for(links, kiz1)["box_id"] is None
    assert _link_for(links, kiz2)["box_id"] == b1
    assert await _picked(async_client, ctx.h, ctx.mid, ctx.pid) == 2
    assert await _box_lines(async_client, ctx.h, ctx.mid, b1) == {ctx.pid: 1}
    assert await _movements(ctx.pid) == movements

    # Код с GTIN не из карточек в контексте проверки: «Сначала отсканируйте ШК товара».
    stranger = _kiz(_ean13("460686000099"), "C54SERIAL0003")
    unknown = await _pick_scan(
        async_client, ctx.h, ctx.mid, stranger, loc_id=ctx.loc_id, box_id=ctx.k1
    )
    assert unknown.status_code == 422, unknown.text
    assert unknown.json()["detail"] == "marking_product_unknown"
    assert len(await _links(async_client, ctx.h, ctx.mid)) == 2


@pytest.mark.asyncio
async def test_c55_known_code_becomes_uncertain_after_unmarked_pick_and_stays_behind(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C55 · R30: k1 «точно в K1»; снята A без КИЗ — k1 «возможно»; «целиком» его не берёт."""
    ctx = await _picking(async_client, monkeypatch, in_box=5, loose=2, plan=2)
    k1, _ = await _return_k1_via_box(async_client, ctx, "C55SERIAL0001")
    exact = await _container_source(async_client, ctx.h, ctx.mid, ctx.pid, ctx.k1)
    assert _tara_codes(exact, "known_marking_codes") == [_head(k1)]

    await _select_k1(async_client, ctx)
    await _pick_unit(async_client, ctx.h, ctx.mid, ctx.pid, EAN_A, ctx.loc_id, ctx.k1)
    assert await _stock(ctx.pid, container_id=ctx.k1) == 4
    source = await _container_source(async_client, ctx.h, ctx.mid, ctx.pid, ctx.k1)
    assert _tara_codes(source, "known_marking_codes") == []
    assert _tara_codes(source, "uncertain_marking_codes") == [_head(k1)]
    assert await _links(async_client, ctx.h, ctx.mid) == []

    z, _ = await _shipment(
        async_client, ctx.h, wid=ctx.wid, sid=ctx.sid, wb_wid=ctx.wb_wid, lines={ctx.pid: 4}
    )
    attached = await async_client.post(
        f"{BASE}/{z}/boxes/attach", headers=ctx.h, json={"barcode": ctx.k1_barcode}
    )
    assert attached.status_code == 201, attached.text
    assert sum(int(line["quantity"]) for line in attached.json()["lines"]) == 4
    assert await _links(async_client, ctx.h, z) == []
    assert await _links(async_client, ctx.h, ctx.mid) == []


@pytest.mark.asyncio
async def test_c56_box_code_with_unknown_source_needs_chosen_return_place(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C56 · R28, R29: k9 впервые на упаковке; «Убрать» без места — отказ; с K2 — в K2."""
    h, wid, sid, wb_wid = await _org(async_client, monkeypatch)
    loc_id, _ = await _cell(async_client, h, wid, "A-1-1")
    pid = await _product(async_client, h, sid, barcode=EAN_A)
    await _receive(async_client, h, wid=wid, pid=pid, qty=4, loc_id=loc_id)
    k1_box, _ = await _box_in_cell(async_client, h, wid=wid, loc_id=loc_id, contents={pid: 2})
    k2_box, _ = await _box_in_cell(async_client, h, wid=wid, loc_id=loc_id, contents={pid: 2})
    mid, _ = await _shipment(async_client, h, wid=wid, sid=sid, wb_wid=wb_wid, lines={pid: 2})
    # Подбор количеством («Снять»): по одной штуке из K1 и из K2, без КИЗ.
    for box in (k1_box, k2_box):
        taken = await _pick_set(async_client, h, mid, pid=pid, loc_id=loc_id, box_id=box, qty=1)
        assert taken.status_code == 200, taken.text
    b1, _ = await _new_box(async_client, h, mid)
    await _pack_unit(async_client, h, mid, b1, pid, EAN_A)
    k9 = _kiz(EAN_A, "C56SERIAL0009")
    k9_id = await _pack_unit(async_client, h, mid, b1, pid, EAN_A, k9)
    assert k9_id is not None
    line_id = await _box_line_id(async_client, h, mid, b1, pid)

    no_place = await _remove_unit(
        async_client, h, mid, b1, line_id, mutation_id=str(uuid.uuid4()), marking_code_id=k9_id
    )
    assert no_place.status_code == 422, no_place.text
    assert no_place.json()["detail"] == "return_place_required"
    assert await _box_lines(async_client, h, mid, b1) == {pid: 2}

    place = {"storage_location_id": loc_id, "container_id": k2_box}
    returned = await async_client.post(
        f"{BASE}/{mid}/boxes/{b1}/lines/{line_id}/remove",
        headers=h,
        json={
            "quantity": 1,
            "mutation_id": str(uuid.uuid4()),
            "marking_code_id": k9_id,
            "return_to": place,
        },
    )
    assert returned.status_code == 200, returned.text
    assert returned.json()["returned_to"] == place
    assert returned.json()["source_known"] is False
    assert await _box_lines(async_client, h, mid, b1) == {pid: 1}
    assert await _stock(pid, container_id=k2_box) == 2
    assert await _stock(pid, container_id=k1_box) == 1
    assert await _links(async_client, h, mid) == []
    k2_source = await _container_source(async_client, h, mid, pid, k2_box)
    assert _tara_codes(k2_source, "known_marking_codes") == [_head(k9)]


@pytest.mark.asyncio
async def test_c57_atomic_return_makes_code_known_but_unlink_and_decrease_does_not(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C57 · R9в, R30: «Вернуть эту единицу» k1 — «точно в K1»; «Отвязать» k2 + «Снять» 0 — нет."""
    ctx = await _picking(async_client, monkeypatch, in_box=5, loose=2, plan=5)
    k1, k2 = _kiz(EAN_A, "C57SERIAL0001"), _kiz(EAN_A, "C57SERIAL0002")
    await _select_k1(async_client, ctx)
    k1_id = await _pick_unit(
        async_client, ctx.h, ctx.mid, ctx.pid, EAN_A, ctx.loc_id, ctx.k1, kiz=k1
    )
    k2_id = await _pick_unit(
        async_client, ctx.h, ctx.mid, ctx.pid, EAN_A, ctx.loc_id, ctx.k1, kiz=k2
    )
    assert await _stock(ctx.pid, container_id=ctx.k1) == 3
    key = str(uuid.uuid4())

    returned = await async_client.post(
        f"{BASE}/{ctx.mid}/marking-codes/{k1_id}/return", headers=ctx.h, json={"mutation_id": key}
    )
    assert returned.status_code == 200, returned.text
    assert returned.json()["returned_to"] == {
        "storage_location_id": ctx.loc_id,
        "container_id": ctx.k1,
    }
    replay = await async_client.post(
        f"{BASE}/{ctx.mid}/marking-codes/{k1_id}/return", headers=ctx.h, json={"mutation_id": key}
    )
    assert replay.status_code == 200, replay.text
    assert await _picked(async_client, ctx.h, ctx.mid, ctx.pid) == 1
    assert await _stock(ctx.pid, container_id=ctx.k1) == 4
    links = await _links(async_client, ctx.h, ctx.mid)
    assert len(links) == 1
    _link_for(links, k2)
    source = await _container_source(async_client, ctx.h, ctx.mid, ctx.pid, ctx.k1)
    assert _tara_codes(source, "known_marking_codes") == [_head(k1)]

    unlinked = await async_client.delete(f"{BASE}/{ctx.mid}/marking-codes/{k2_id}", headers=ctx.h)
    assert unlinked.status_code == 200, unlinked.text
    to_zero = await _pick_set(
        async_client, ctx.h, ctx.mid, pid=ctx.pid, loc_id=ctx.loc_id, box_id=ctx.k1, qty=0
    )
    assert to_zero.status_code == 200, to_zero.text
    assert await _picked(async_client, ctx.h, ctx.mid, ctx.pid) == 0
    assert await _stock(ctx.pid, container_id=ctx.k1) == 5
    source = await _container_source(async_client, ctx.h, ctx.mid, ctx.pid, ctx.k1)
    assert _tara_codes(source, "known_marking_codes") == [_head(k1)]
    assert _tara_codes(source, "uncertain_marking_codes") == []


@pytest.mark.asyncio
async def test_c58_scan_of_another_box_while_box_selected_switches_source(
    async_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """C58 · R1, K23: выбран K1; скан ШК K2 (запрос несёт K1, как шлёт веб) выбирает K2."""
    h, wid, sid, wb_wid = await _org(async_client, monkeypatch)
    loc_id, _ = await _cell(async_client, h, wid, "A-1-1")
    pid = await _product(async_client, h, sid, barcode=EAN_A)
    await _receive(async_client, h, wid=wid, pid=pid, qty=4, loc_id=loc_id)
    k1, k1_barcode = await _box_in_cell(async_client, h, wid=wid, loc_id=loc_id, contents={pid: 2})
    k2, k2_barcode = await _box_in_cell(async_client, h, wid=wid, loc_id=loc_id, contents={pid: 2})
    mid, _ = await _shipment(async_client, h, wid=wid, sid=sid, wb_wid=wb_wid, lines={pid: 2})
    await _select(async_client, h, mid, k1_barcode)

    # Веб при выбранной таре шлёт её с каждым сканом (FfUnloadPickPage.tsx:486-489).
    switched = await _pick_scan(async_client, h, mid, k2_barcode, loc_id=loc_id, box_id=k1)
    assert switched.status_code == 200, f"ШК короба K2 не распознан как тара: {switched.text}"
    assert switched.json()["kind"] == "container"
    assert switched.json()["container_id"] == k2
    assert switched.json()["storage_location_id"] == loc_id

    unit = await _pick_scan(async_client, h, mid, EAN_A, pid=pid, loc_id=loc_id, box_id=k2)
    assert unit.status_code == 200, unit.text
    assert unit.json()["kind"] == "product"
    assert await _stock(pid, container_id=k2) == 1
    assert await _stock(pid, container_id=k1) == 2
    assert await _picked(async_client, h, mid, pid) == 1
