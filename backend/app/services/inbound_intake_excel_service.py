"""Atomic import of planned seller intake quantities from an XLSX template."""

from __future__ import annotations

import io
import uuid
from collections import defaultdict
from dataclasses import dataclass
from typing import Any

from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.inbound_intake import InboundIntakeLine
from app.models.product import Product
from app.models.product_barcode import ProductBarcode
from app.models.product_marketplace_link import ProductMarketplaceLink
from app.services import inbound_intake_service as intake

HEADERS = ("Товар", "Артикул", "ШК", "Размер", "Количество (штук)")


@dataclass(frozen=True)
class ImportIssue:
    row: int
    article: str
    message: str


class IntakeExcelError(Exception):
    def __init__(self, message: str, issues: list[ImportIssue] | None = None) -> None:
        super().__init__(message)
        self.message = message
        self.issues = issues or []


def template_xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Товары"
    sheet.append(HEADERS)
    for column, width in zip("ABCDE", (36, 26, 26, 18, 24), strict=True):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    for column in "ABCD":
        sheet.column_dimensions[column].number_format = "@"
        sheet[f"{column}2"].number_format = "@"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _identifier(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _quantity(value: Any) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value if 0 < value <= 1_000_000_000 else None
    if isinstance(value, float):
        return int(value) if value.is_integer() and 0 < value <= 1_000_000_000 else None
    if isinstance(value, str):
        text = value.strip()
        if text.isdecimal():
            number = int(text)
            return number if 0 < number <= 1_000_000_000 else None
    return None


def parse_xlsx(content: bytes) -> tuple[list[tuple[int, str, str, str, int]], list[ImportIssue]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise IntakeExcelError("Excel не в том формате. Заполните шаблон ещё раз.") from exc
    try:
        sheet = workbook.active
        if sheet is None:
            raise IntakeExcelError("Excel не в том формате. Заполните шаблон ещё раз.")
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if (
            header is None
            or tuple(_identifier(v) for v in header[:5]) != HEADERS
            or any(_identifier(v) for v in header[5:])
        ):
            raise IntakeExcelError("Excel не в том формате. Заполните шаблон ещё раз.")
        parsed: list[tuple[int, str, str, str, int]] = []
        issues: list[ImportIssue] = []
        for row_number, row in enumerate(rows, start=2):
            values = tuple(row[:5]) + (None,) * max(0, 5 - len(row))
            if not any(_identifier(value) for value in row):
                continue
            article = _identifier(values[1])
            barcode = _identifier(values[2])
            code = article or barcode
            size = _identifier(values[3])
            quantity = _quantity(values[4])
            if any(_identifier(value) for value in row[5:]):
                issues.append(ImportIssue(row_number, code, "Лишние данные вне шаблона"))
            if not article and not barcode:
                issues.append(ImportIssue(row_number, code, "Укажите артикул или ШК"))
            if quantity is None:
                issues.append(
                    ImportIssue(
                        row_number, code, "Количество должно быть целым положительным числом"
                    )
                )
            if (article or barcode) and quantity is not None:
                parsed.append((row_number, article, barcode, size, quantity))
        if not parsed and not issues:
            raise IntakeExcelError("В Excel нет товаров. Заполните шаблон и загрузите снова.")
        return parsed, issues
    except IntakeExcelError:
        raise
    except Exception as exc:
        raise IntakeExcelError("Excel не в том формате. Заполните шаблон ещё раз.") from exc
    finally:
        workbook.close()


async def import_xlsx(
    session: AsyncSession,
    tenant_id: uuid.UUID,
    request_id: uuid.UUID,
    seller_id: uuid.UUID,
    content: bytes,
) -> int:
    rows, issues = parse_xlsx(content)
    request = await intake.get_request(
        session, tenant_id, request_id, seller_product_owner_id=seller_id, for_update=True
    )
    if request is None:
        raise intake.InboundIntakeError("request_not_found")
    if not intake._request_plan_editable(request, seller_product_owner_id=seller_id):
        raise intake.InboundIntakeError("not_draft")

    products = list(
        (
            await session.scalars(
                select(Product).where(
                    Product.tenant_id == tenant_id, Product.seller_id == seller_id
                )
            )
        ).all()
    )
    by_article: dict[str, list[Product]] = defaultdict(list)
    by_barcode: dict[str, list[Product]] = defaultdict(list)
    for product in products:
        by_article[product.sku_code].append(product)
        if product.wb_barcode:
            by_barcode[product.wb_barcode].append(product)
    product_by_id = {product.id: product for product in products}
    stored_barcodes = await session.execute(
        select(ProductBarcode.product_id, ProductBarcode.barcode).where(
            ProductBarcode.tenant_id == tenant_id,
            ProductBarcode.seller_id == seller_id,
            ProductBarcode.source == "wb",
        )
    )
    for product_id, barcode in stored_barcodes:
        owner = product_by_id.get(product_id)
        if owner is not None and owner not in by_barcode[barcode]:
            by_barcode[barcode].append(owner)
    marketplace_links = await session.execute(
        select(ProductMarketplaceLink.product_id, ProductMarketplaceLink.external_barcodes).where(
            ProductMarketplaceLink.tenant_id == tenant_id,
            ProductMarketplaceLink.seller_id == seller_id,
        )
    )
    for product_id, barcodes in marketplace_links:
        owner = product_by_id.get(product_id)
        if owner is not None:
            for barcode in barcodes or []:
                if isinstance(barcode, str) and barcode and owner not in by_barcode[barcode]:
                    by_barcode[barcode].append(owner)
    quantities: dict[uuid.UUID, int] = defaultdict(int)
    for row_number, article, barcode, size, quantity in rows:
        code = article or barcode
        candidates = by_article.get(article, []) if article else by_barcode.get(barcode, [])
        if article and barcode:
            candidates = [
                product for product in candidates if product in by_barcode.get(barcode, [])
            ]
        if size:
            candidates = [product for product in candidates if (product.wb_size or "") == size]
        if not candidates:
            issues.append(
                ImportIssue(
                    row_number,
                    code,
                    "Товар с указанными артикулом, ШК и размером не найден у селлера",
                )
            )
            continue
        if len(candidates) != 1:
            issues.append(
                ImportIssue(row_number, code, "Неоднозначный товар: уточните ШК и размер")
            )
            continue
        quantities[candidates[0].id] += quantity
        if quantities[candidates[0].id] > 1_000_000_000:
            issues.append(
                ImportIssue(
                    row_number, code, "Суммарное количество превышает допустимое для строки"
                )
            )
    existing_by_product = {line.product_id: line for line in request.lines}
    for product_id in quantities:
        line = existing_by_product.get(product_id)
        if line is not None and line.posted_qty:
            issues.append(ImportIssue(0, line.product.sku_code, "Строка уже проведена"))
    if issues:
        raise IntakeExcelError(
            "Не удалось загрузить Excel. Исправьте строки и повторите загрузку.", issues
        )

    for product_id, quantity in quantities.items():
        line = existing_by_product.get(product_id)
        if line is None:
            session.add(
                InboundIntakeLine(
                    request_id=request_id,
                    product_id=product_id,
                    expected_qty=quantity,
                    actual_qty=None,
                    posted_qty=0,
                )
            )
        else:
            line.expected_qty = quantity
    await session.commit()
    return len(quantities)
