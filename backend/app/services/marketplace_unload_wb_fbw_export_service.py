"""Local XLSX export of an existing WB FBW shipment's box composition.

This module deliberately has no Wildberries client.  The confirmed integration
boundary is a file the operator uploads in WB Partners; WB packaging writes and
WB label data are not an established contract.
"""

from __future__ import annotations

import uuid
import zipfile
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook  # type: ignore[import-untyped]
from sqlalchemy import and_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.marketplace_unload import (
    MarketplaceUnloadBox,
    MarketplaceUnloadBoxLine,
    MarketplaceUnloadRequest,
)
from app.models.product import Product
from app.models.product_barcode import ProductBarcode
from app.models.warehouse_box import WarehouseBox

_HEADERS = (
    "Баркод товара",
    "Кол-во товаров",
    "ШК короба",
    "Срок годности",
    "ШК короба для печати в стороннем сервисе",
)
_XLSX_ZIP_TIME = (2000, 1, 1, 0, 0, 0)


class WbFbwPackagingExportError(Exception):
    """The saved WMS composition cannot truthfully be represented in the WB file."""


@dataclass(frozen=True)
class WbFbwPackagingXlsxExport:
    content: bytes
    warnings: tuple[str, ...]


@dataclass(frozen=True)
class _ExportRow:
    box_id: uuid.UUID
    box_created_at: datetime
    box_barcode: str | None
    request_seller_id: uuid.UUID | None
    line_id: uuid.UUID | None
    line_created_at: datetime | None
    quantity: int | None
    product_id: uuid.UUID | None
    product_name: str | None
    product_tenant_id: uuid.UUID | None
    product_seller_id: uuid.UUID | None
    product_wb_barcode: str | None
    product_wb_shelf_life: str | None
    barcode: str | None


def _normalized(value: str | None) -> str | None:
    if value is None:
        return None
    normalized = value.strip()
    return normalized or None


def _stable_xlsx(workbook: Workbook) -> bytes:
    """Save deterministic bytes so retried downloads of the same snapshot match."""
    raw = BytesIO()
    workbook.save(raw)
    result = BytesIO()
    with zipfile.ZipFile(BytesIO(raw.getvalue()), "r") as source, zipfile.ZipFile(
        result, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as target:
        for name in sorted(source.namelist()):
            entry = zipfile.ZipInfo(name, date_time=_XLSX_ZIP_TIME)
            entry.compress_type = zipfile.ZIP_DEFLATED
            entry.external_attr = 0o600 << 16
            target.writestr(entry, source.read(name))
    return result.getvalue()


def _workbook(rows: list[tuple[str, int, str]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    workbook.properties.created = datetime(2000, 1, 1)
    workbook.properties.modified = datetime(2000, 1, 1)
    sheet.append(_HEADERS)
    for row_number, (product_barcode, quantity, box_barcode) in enumerate(rows, start=2):
        # openpyxl considers a value beginning with '=' to be a formula unless its
        # type is explicitly string.  WB barcodes are data, never spreadsheet code.
        product_cell = sheet.cell(row_number, 1, product_barcode)
        product_cell.data_type = "s"
        sheet.cell(row_number, 2, quantity)
        box_cell = sheet.cell(row_number, 3, box_barcode)
        box_cell.data_type = "s"
        # D is intentionally empty: WMS has shelf-life duration, not a batch expiry.
        # E is intentionally empty: WB has not supplied print data for this box.
    return _stable_xlsx(workbook)


async def export_wb_fbw_packaging_xlsx(
    session: AsyncSession,
    *,
    tenant_id: uuid.UUID,
    request_id: uuid.UUID,
) -> WbFbwPackagingXlsxExport:
    """Build a read-only, one-query snapshot of a WB shipment's existing boxes."""
    rows_result = await session.execute(
        select(
            MarketplaceUnloadBox.id,
            MarketplaceUnloadBox.created_at,
            WarehouseBox.internal_barcode,
            MarketplaceUnloadRequest.seller_id,
            MarketplaceUnloadBoxLine.id,
            MarketplaceUnloadBoxLine.created_at,
            MarketplaceUnloadBoxLine.quantity,
            Product.id,
            Product.name,
            Product.tenant_id,
            Product.seller_id,
            Product.wb_barcode,
            Product.wb_shelf_life,
            ProductBarcode.barcode,
        )
        .select_from(MarketplaceUnloadRequest)
        .join(MarketplaceUnloadBox, MarketplaceUnloadBox.request_id == MarketplaceUnloadRequest.id)
        .outerjoin(WarehouseBox, WarehouseBox.id == MarketplaceUnloadBox.warehouse_box_id)
        .outerjoin(
            MarketplaceUnloadBoxLine,
            MarketplaceUnloadBoxLine.box_id == MarketplaceUnloadBox.id,
        )
        .outerjoin(Product, Product.id == MarketplaceUnloadBoxLine.product_id)
        .outerjoin(
            ProductBarcode,
            and_(
                ProductBarcode.product_id == Product.id,
                ProductBarcode.tenant_id == MarketplaceUnloadRequest.tenant_id,
                ProductBarcode.seller_id == MarketplaceUnloadRequest.seller_id,
                ProductBarcode.source == "wb",
            ),
        )
        .where(
            MarketplaceUnloadRequest.id == request_id,
            MarketplaceUnloadRequest.tenant_id == tenant_id,
        )
        .where(MarketplaceUnloadRequest.marketplace == "wb")
    )
    raw_rows = [
        _ExportRow(*row)
        for row in rows_result.tuples().all()
    ]
    if not raw_rows:
        request = await session.scalar(
            select(MarketplaceUnloadRequest).where(
                MarketplaceUnloadRequest.id == request_id,
                MarketplaceUnloadRequest.tenant_id == tenant_id,
            )
        )
        if request is None:
            raise WbFbwPackagingExportError("Отгрузка не найдена или недоступна.")
        raise WbFbwPackagingExportError("Выгрузка XLSX доступна только для отгрузки WB.")

    wb_barcodes_by_product: dict[uuid.UUID, set[str]] = {}
    for row in raw_rows:
        if row.product_id is not None and (barcode := _normalized(row.barcode)) is not None:
            wb_barcodes_by_product.setdefault(row.product_id, set()).add(barcode)

    export_rows: list[tuple[str, int, str]] = []
    warnings: list[str] = []
    warned_product_ids: set[uuid.UUID] = set()
    handled_line_ids: set[uuid.UUID] = set()
    for row in raw_rows:
        box_barcode = _normalized(row.box_barcode)
        if box_barcode is None:
            raise WbFbwPackagingExportError(
                f"У короба {row.box_id} нет связи с WarehouseBox или внутреннего ШК."
            )
        if row.line_id is None:
            continue
        # One product can have several fallback barcodes, therefore the SQL join
        # repeats its box line.  Validate and emit it exactly once.
        if row.line_id in handled_line_ids:
            continue
        handled_line_ids.add(row.line_id)

        if row.product_id is None or row.product_name is None or row.quantity is None:
            raise WbFbwPackagingExportError(f"У короба {box_barcode} есть неполная строка состава.")
        if row.product_tenant_id != tenant_id:
            raise WbFbwPackagingExportError(
                f"Товар «{row.product_name}» недоступен в текущем тенанте."
            )
        if (
            row.request_seller_id is not None
            and row.product_seller_id != row.request_seller_id
        ):
            raise WbFbwPackagingExportError(
                f"Товар «{row.product_name}» принадлежит другому селлеру."
            )
        if row.quantity < 1:
            raise WbFbwPackagingExportError(
                f"У товара «{row.product_name}» в коробе {box_barcode} некорректное количество."
            )

        product_barcode = _normalized(row.product_wb_barcode)
        if product_barcode is None:
            candidates = wb_barcodes_by_product.get(row.product_id, set())
            if len(candidates) == 1:
                product_barcode = next(iter(candidates))
            elif len(candidates) > 1:
                raise WbFbwPackagingExportError(
                    f"Неоднозначный вариант WB-баркода для товара «{row.product_name}»."
                )
            else:
                raise WbFbwPackagingExportError(
                    f"У товара «{row.product_name}» нет WB-баркода для выгрузки."
                )

        export_rows.append((product_barcode, row.quantity, box_barcode))
        if (
            _normalized(row.product_wb_shelf_life) is not None
            and row.product_id not in warned_product_ids
        ):
            warnings.append(
                f"Для товара «{row.product_name}» перед загрузкой в WB укажите "
                "фактическую дату партии: "
                "в WMS нет даты окончания срока годности."
            )
            warned_product_ids.add(row.product_id)

    return WbFbwPackagingXlsxExport(content=_workbook(export_rows), warnings=tuple(warnings))
