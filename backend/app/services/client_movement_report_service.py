"""WMS-533: read-only client movement feed and its Excel projection."""

from __future__ import annotations

import base64
import io
import json
import uuid
from collections import defaultdict
from datetime import UTC, datetime
from typing import Any

from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.cell import WriteOnlyCell  # type: ignore[import-untyped]
from openpyxl.styles import Font, PatternFill  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]
from sqlalchemy import and_, bindparam, case, or_, select, text
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.fbs_order import FbsOrder, FbsOrderMarking, FbsOrderProduct
from app.models.fbs_shipment_reversal_ledger import FbsShipmentReversalLedger
from app.models.inbound_intake import InboundIntakeLine, InboundIntakeRequest
from app.models.inventory_count import InventoryCount, InventoryCountLine
from app.models.inventory_movement import InventoryMovement
from app.models.marketplace_unload import MarketplaceUnloadRequest
from app.models.marking_code import MarkingCode
from app.models.outbound_shipment import OutboundShipmentLine, OutboundShipmentRequest
from app.models.product import Product
from app.models.product_barcode import ProductBarcode
from app.models.warehouse import Warehouse

MAX_LIMIT = 1000
_BATCH = 500
_STOCK_EVENTS = (
    "inbound_intake",
    "fbs_shipment",
    "marketplace_unload",
    "inventory_count",
)

_EVENT_NAMES = {
    "inbound_intake": "Приёмка",
    "return": "Возврат товара",
    "fbs_shipment": "Отгрузка FBS",
    "marketplace_unload": "Отгрузка FBO",
    "inventory_count": "Инвентаризация",
}


def _period(date_from: datetime, date_to: datetime) -> tuple[datetime, datetime]:
    if date_from.tzinfo is None or date_to.tzinfo is None:
        raise ValueError("date_from and date_to must include a timezone offset")
    start, end = date_from.astimezone(UTC), date_to.astimezone(UTC)
    if end <= start:
        raise ValueError("date_to must be after date_from")
    return start, end


def _decode_cursor(cursor: str | None) -> tuple[datetime, uuid.UUID, int] | None:
    if cursor is None:
        return None
    try:
        values = json.loads(base64.urlsafe_b64decode(cursor + "=" * (-len(cursor) % 4)))
        at = datetime.fromisoformat(values[0])
        if at.tzinfo is None or type(values[2]) is not int or values[2] < 0:
            raise ValueError
        return at.astimezone(UTC), uuid.UUID(values[1]), values[2]
    except (ValueError, TypeError, IndexError, KeyError, UnicodeDecodeError) as exc:
        raise ValueError("invalid cursor") from exc


def _encode_cursor(at: datetime, movement_id: uuid.UUID, unit: int) -> str:
    raw = json.dumps([at.isoformat(), str(movement_id), unit], separators=(",", ":"))
    return base64.urlsafe_b64encode(raw.encode()).decode().rstrip("=")


async def _fbs_index(
    session: AsyncSession,
    tenant_id: uuid.UUID,
    seller_id: uuid.UUID | None,
    movements: list[InventoryMovement],
) -> dict[uuid.UUID, tuple[FbsOrder, int | None]]:
    """Resolve the current batch using explicit movement and group IDs."""
    fbs = [m for m in movements if m.movement_type == "fbs_shipment"]
    if not fbs:
        return {}
    movement_ids = {m.id for m in fbs}
    groups = {m.transfer_group_id for m in fbs if m.transfer_group_id is not None}
    if groups:
        movement_ids.update(
            (
                await session.scalars(
                    select(InventoryMovement.id).where(
                        InventoryMovement.tenant_id == tenant_id,
                        InventoryMovement.transfer_group_id.in_(groups),
                        InventoryMovement.movement_type == "fbs_shipment",
                    )
                )
            ).all()
        )
    if session.get_bind().dialect.name == "postgresql":
        ozon_match = text(
            "EXISTS (SELECT 1 FROM jsonb_array_elements("
            "COALESCE(NULLIF(fbs_shipment_reversal_ledger.ozon_positions_json::jsonb, "
            "'null'::jsonb), '[]'::jsonb)) "
            "AS position WHERE position->>'movement_id' IN :movement_ids)"
        )
    else:
        ozon_match = text(
            "EXISTS (SELECT 1 FROM json_each("
            "fbs_shipment_reversal_ledger.ozon_positions_json) AS position "
            "WHERE json_extract(position.value, '$.movement_id') IN :movement_ids)"
        )
    ozon_match = ozon_match.bindparams(
        bindparam("movement_ids", [str(mid) for mid in movement_ids], expanding=True)
    )
    stmt = (
        select(FbsShipmentReversalLedger, FbsOrder)
        .join(FbsOrder, FbsOrder.id == FbsShipmentReversalLedger.fbs_order_id)
        .where(
            FbsShipmentReversalLedger.tenant_id == tenant_id,
            FbsOrder.tenant_id == tenant_id,
            or_(
                FbsShipmentReversalLedger.shipment_movement_id.in_(movement_ids),
                FbsShipmentReversalLedger.reversal_movement_id.in_(movement_ids),
                ozon_match,
            ),
        )
    )
    if seller_id is not None:
        stmt = stmt.where(FbsOrder.seller_id == seller_id)
    index: dict[uuid.UUID, tuple[FbsOrder, int | None]] = {}
    for ledger, order in (await session.execute(stmt)).all():
        if order.marketplace not in {"wb", "ozon"}:
            continue
        offsets: dict[uuid.UUID, int] = defaultdict(int)
        for position in ledger.ozon_positions_json or []:
            try:
                movement_id = uuid.UUID(str(position["movement_id"]))
                product_id = uuid.UUID(str(position["product_id"]))
                quantity = int(str(position["quantity"]))
            except (ValueError, TypeError, KeyError):
                continue
            index[movement_id] = (order, offsets[product_id])
            offsets[product_id] += max(0, quantity)
        if ledger.shipment_movement_id is not None:
            index.setdefault(ledger.shipment_movement_id, (order, 0))
        if ledger.reversal_movement_id is not None:
            index[ledger.reversal_movement_id] = (order, 0 if order.marketplace == "wb" else None)
    if index:
        anchors = (
            await session.scalars(
                select(InventoryMovement).where(
                    InventoryMovement.tenant_id == tenant_id,
                    InventoryMovement.id.in_(set(index)),
                    InventoryMovement.transfer_group_id.is_not(None),
                    InventoryMovement.movement_type == "fbs_shipment",
                )
            )
        ).all()
        groups = {m.transfer_group_id for m in anchors if m.transfer_group_id is not None}
        if groups:
            siblings = (
                await session.scalars(
                    select(InventoryMovement).where(
                        InventoryMovement.tenant_id == tenant_id,
                        InventoryMovement.transfer_group_id.in_(groups),
                        InventoryMovement.movement_type == "fbs_shipment",
                    )
                )
            ).all()
            by_group: dict[uuid.UUID, list[InventoryMovement]] = defaultdict(list)
            for movement in siblings:
                if movement.transfer_group_id is not None:
                    by_group[movement.transfer_group_id].append(movement)
            for anchor_movement in anchors:
                order, base = index[anchor_movement.id]
                if base is None or anchor_movement.transfer_group_id is None:
                    continue
                offset = base
                members = [
                    anchor_movement,
                    *sorted(
                        (
                            item
                            for item in by_group[anchor_movement.transfer_group_id]
                            if item.id != anchor_movement.id
                        ),
                        key=lambda item: str(item.id),
                    ),
                ]
                for movement in members:
                    if movement.product_id != anchor_movement.product_id:
                        continue
                    if movement.id in index and movement.id != anchor_movement.id:
                        continue
                    index[movement.id] = (order, offset)
                    offset += abs(int(movement.quantity_delta))
    return index


async def _marking_index(
    session: AsyncSession, tenant_id: uuid.UUID, orders: set[uuid.UUID]
) -> dict[tuple[uuid.UUID, uuid.UUID], list[str]]:
    if not orders:
        return {}
    positions = (
        await session.execute(
            select(FbsOrderProduct.order_id, FbsOrderProduct.product_id, FbsOrderProduct.id).where(
                FbsOrderProduct.order_id.in_(orders)
            )
        )
    ).all()
    by_position = {
        position_id: (order_id, product_id) for order_id, product_id, position_id in positions
    }
    counts: dict[tuple[uuid.UUID, uuid.UUID], int] = defaultdict(int)
    positions_by_order: dict[uuid.UUID, list[uuid.UUID | None]] = defaultdict(list)
    for order_id, product_id, _ in positions:
        positions_by_order[order_id].append(product_id)
        if product_id is not None:
            counts[(order_id, product_id)] += 1
    order_products = {
        order_id: product_id
        for order_id, product_id in (
            await session.execute(
                select(FbsOrder.id, FbsOrder.product_id).where(
                    FbsOrder.tenant_id == tenant_id, FbsOrder.id.in_(orders)
                )
            )
        ).all()
    }
    stmt = (
        select(FbsOrderMarking, MarkingCode.cis_code, MarkingCode.product_id)
        .outerjoin(
            MarkingCode,
            and_(
                MarkingCode.id == FbsOrderMarking.marking_code_id,
                MarkingCode.tenant_id == tenant_id,
            ),
        )
        .where(
            FbsOrderMarking.tenant_id == tenant_id,
            FbsOrderMarking.order_id.in_(orders),
            FbsOrderMarking.kind == "sgtin",
            FbsOrderMarking.meta_status.not_in(("rejected", "replacement_required")),
        )
        .order_by(FbsOrderMarking.id)
    )
    codes: dict[tuple[uuid.UUID, uuid.UUID], list[str]] = defaultdict(list)
    for mark, cis_code, code_product_id in (await session.execute(stmt)).all():
        if mark.order_product_id is not None:
            owner = by_position.get(mark.order_product_id)
            if owner is None or owner[0] != mark.order_id or owner[1] is None:
                continue
            product_id = owner[1]
        else:
            positions_for_order = positions_by_order.get(mark.order_id, [])
            if not positions_for_order:
                product_id = order_products.get(mark.order_id)
            elif len(positions_for_order) == 1:
                product_id = positions_for_order[0]
            elif code_product_id is not None and counts[(mark.order_id, code_product_id)] == 1:
                product_id = code_product_id
            else:
                product_id = None
            if product_id is None:
                continue
        if counts[(mark.order_id, product_id)] > 1:
            # Recipe has no order_product_id: this instance cannot be proven.
            continue
        codes[(mark.order_id, product_id)].append(cis_code or mark.value)
    return codes


async def _documents(
    session: AsyncSession, tenant_id: uuid.UUID, movements: list[InventoryMovement]
) -> tuple[dict[uuid.UUID, dict[str, str | None]], dict[uuid.UUID, str | None]]:
    docs: dict[uuid.UUID, dict[str, str | None]] = {}
    return_docs: dict[uuid.UUID, dict[str, str | None]] = {}
    return_marketplaces: dict[uuid.UUID, str | None] = {}
    inbound_ids = {m.inbound_intake_line_id for m in movements if m.inbound_intake_line_id}
    if inbound_ids:
        data = {
            line_id: request
            for line_id, request in (
                await session.execute(
                    select(InboundIntakeLine.id, InboundIntakeRequest)
                    .join(
                        InboundIntakeRequest,
                        InboundIntakeRequest.id == InboundIntakeLine.request_id,
                    )
                    .where(
                        InboundIntakeLine.id.in_(inbound_ids),
                        InboundIntakeRequest.tenant_id == tenant_id,
                    )
                )
            ).all()
        }
        for m in movements:
            if m.inbound_intake_line_id in data:
                doc = data[m.inbound_intake_line_id]
                is_return = m.movement_type == "inbound_intake" and doc.operation_type == "return"
                docs[m.id] = {
                    "id": str(doc.id),
                    "type": "return" if is_return else "inbound",
                    "number": doc.display_number or doc.document_number or "",
                }
                if is_return:
                    return_docs[m.id] = docs[m.id]
                    return_marketplaces[m.id] = (
                        doc.marketplace if doc.marketplace in {"wb", "ozon"} else None
                    )
    unload_ids = {
        m.marketplace_unload_request_id for m in movements if m.marketplace_unload_request_id
    }
    if unload_ids:
        data = {
            d.id: d
            for d in (
                await session.scalars(
                    select(MarketplaceUnloadRequest).where(
                        MarketplaceUnloadRequest.id.in_(unload_ids),
                        MarketplaceUnloadRequest.tenant_id == tenant_id,
                    )
                )
            ).all()
        }
        for m in movements:
            if m.marketplace_unload_request_id in data:
                doc = data[m.marketplace_unload_request_id]
                docs[m.id] = {
                    "id": str(doc.id),
                    "type": "marketplace_unload",
                    "number": doc.display_number or doc.document_number or "",
                    "status": doc.status,
                    "shipped_at": (
                        doc.shipped_at.replace(tzinfo=UTC).isoformat()
                        if doc.shipped_at is not None and doc.shipped_at.tzinfo is None
                        else doc.shipped_at.astimezone(UTC).isoformat()
                        if doc.shipped_at is not None
                        else None
                    ),
                }
    outbound_ids = {m.outbound_shipment_line_id for m in movements if m.outbound_shipment_line_id}
    if outbound_ids:
        data = {
            line_id: request_id
            for line_id, request_id in (
                await session.execute(
                    select(OutboundShipmentLine.id, OutboundShipmentLine.request_id)
                    .join(
                        OutboundShipmentRequest,
                        OutboundShipmentRequest.id == OutboundShipmentLine.request_id,
                    )
                    .where(
                        OutboundShipmentLine.id.in_(outbound_ids),
                        OutboundShipmentRequest.tenant_id == tenant_id,
                    )
                )
            ).all()
        }
        for m in movements:
            if m.outbound_shipment_line_id in data:
                docs[m.id] = {
                    "id": str(data[m.outbound_shipment_line_id]),
                    "type": "outbound_shipment",
                    "number": "",
                }
    count_ids = {m.inventory_count_line_id for m in movements if m.inventory_count_line_id}
    if count_ids:
        data = {
            line_id: count_id
            for line_id, count_id in (
                await session.execute(
                    select(InventoryCountLine.id, InventoryCountLine.count_id)
                    .join(InventoryCount, InventoryCount.id == InventoryCountLine.count_id)
                    .where(
                        InventoryCountLine.id.in_(count_ids), InventoryCount.tenant_id == tenant_id
                    )
                )
            ).all()
        }
        for m in movements:
            if m.inventory_count_line_id in data:
                count_id = data[m.inventory_count_line_id]
                docs[m.id] = {
                    "id": str(count_id),
                    "type": "inventory_count",
                    "number": f"ИНВ-{str(count_id).split('-')[0].upper()}",
                }
    docs.update(return_docs)
    return docs, return_marketplaces


async def list_client_movements(
    session: AsyncSession,
    tenant_id: uuid.UUID,
    *,
    date_from: datetime,
    date_to: datetime,
    warehouse_id: uuid.UUID | None = None,
    sku: str | None = None,
    shk: str | None = None,
    marketplace: str | None = None,
    seller_id: uuid.UUID | None = None,
    cursor: str | None = None,
    limit: int = 200,
) -> tuple[list[dict[str, Any]], str | None]:
    start, end = _period(date_from, date_to)
    if sku is not None and shk is not None:
        raise ValueError("sku and shk cannot be used together")
    if not 1 <= limit <= MAX_LIMIT:
        raise ValueError("limit must be between 1 and 1000")
    if marketplace not in (None, "wb", "ozon"):
        raise ValueError("marketplace must be wb or ozon")
    after = _decode_cursor(cursor)
    if after is not None and not start <= after[0] < end:
        raise ValueError("cursor is outside the requested period")
    result: list[dict[str, Any]] = []
    last_key: tuple[datetime, uuid.UUID, int] | None = None
    scan_at = after[0] if after else start
    scan_id = after[1] if after else None
    resume_unit = after is not None
    while True:
        event_at = case(
            (
                InventoryMovement.movement_type == "marketplace_unload",
                MarketplaceUnloadRequest.shipped_at,
            ),
            else_=InventoryMovement.created_at,
        )
        stmt = (
            select(InventoryMovement, Product, event_at)
            .join(
                Product,
                and_(Product.id == InventoryMovement.product_id, Product.tenant_id == tenant_id),
            )
            .outerjoin(
                MarketplaceUnloadRequest,
                and_(
                    MarketplaceUnloadRequest.id == InventoryMovement.marketplace_unload_request_id,
                    MarketplaceUnloadRequest.tenant_id == tenant_id,
                ),
            )
            .where(
                InventoryMovement.tenant_id == tenant_id,
                InventoryMovement.movement_type.in_(_STOCK_EVENTS),
                or_(
                    and_(
                        InventoryMovement.movement_type == "inbound_intake",
                        InventoryMovement.quantity_delta > 0,
                    ),
                    and_(
                        InventoryMovement.movement_type.in_(("fbs_shipment", "marketplace_unload")),
                        InventoryMovement.quantity_delta < 0,
                    ),
                    and_(
                        InventoryMovement.movement_type == "inventory_count",
                        InventoryMovement.quantity_delta != 0,
                    ),
                ),
                event_at >= start,
                event_at < end,
                or_(
                    InventoryMovement.movement_type != "marketplace_unload",
                    and_(
                        MarketplaceUnloadRequest.status.in_(("shipped", "done")),
                        MarketplaceUnloadRequest.shipped_at.is_not(None),
                    ),
                ),
            )
            .order_by(event_at, InventoryMovement.id)
            .limit(_BATCH)
        )
        if scan_id is not None:
            stmt = stmt.where(
                or_(
                    event_at > scan_at,
                    and_(
                        event_at == scan_at,
                        InventoryMovement.id >= scan_id
                        if resume_unit
                        else InventoryMovement.id > scan_id,
                    ),
                )
            )
        if warehouse_id is not None:
            stmt = stmt.where(InventoryMovement.warehouse_id == warehouse_id)
        if seller_id is not None:
            stmt = stmt.where(Product.seller_id == seller_id)
        if sku is not None:
            stmt = stmt.where(Product.sku_code == sku)
        if shk is not None:
            extra_barcode = (
                select(ProductBarcode.id)
                .where(
                    ProductBarcode.tenant_id == tenant_id,
                    ProductBarcode.seller_id == Product.seller_id,
                    ProductBarcode.product_id == Product.id,
                    ProductBarcode.barcode == shk,
                )
                .exists()
            )
            stmt = stmt.where(or_(Product.wb_barcode == shk, extra_barcode))
        batch = (await session.execute(stmt)).all()
        if not batch:
            break
        movements = [movement for movement, _, _ in batch]
        ledger_index = await _fbs_index(session, tenant_id, seller_id, movements)
        product_ids = {product.id for _, product, _ in batch}
        alternate_barcodes: dict[uuid.UUID, str] = {}
        if product_ids:
            barcode_rows = await session.execute(
                select(ProductBarcode.product_id, ProductBarcode.barcode)
                .where(
                    ProductBarcode.tenant_id == tenant_id,
                    ProductBarcode.product_id.in_(product_ids),
                )
                .order_by(ProductBarcode.barcode)
            )
            for product_id, alternate in barcode_rows:
                alternate_barcodes.setdefault(product_id, alternate)
        docs, return_marketplaces = await _documents(session, tenant_id, movements)
        related_orders = {ledger_index[m.id][0].id for m in movements if m.id in ledger_index}
        marks = await _marking_index(session, tenant_id, related_orders)
        unload_ids = {
            m.marketplace_unload_request_id for m in movements if m.marketplace_unload_request_id
        }
        unload_markets: dict[uuid.UUID, str] = (
            {
                request_id: market
                for request_id, market in (
                    await session.execute(
                        select(
                            MarketplaceUnloadRequest.id, MarketplaceUnloadRequest.marketplace
                        ).where(
                            MarketplaceUnloadRequest.tenant_id == tenant_id,
                            MarketplaceUnloadRequest.id.in_(unload_ids),
                        )
                    )
                ).all()
            }
            if unload_ids
            else {}
        )
        for movement, product, effective_at in batch:
            movement_at = (
                effective_at.replace(tzinfo=UTC)
                if effective_at.tzinfo is None
                else effective_at.astimezone(UTC)
            )
            scan_at, scan_id = effective_at, movement.id
            linked = (
                ledger_index.get(movement.id) if movement.movement_type == "fbs_shipment" else None
            )
            order, offset = linked if linked is not None else (None, None)
            row_marketplace = (
                order.marketplace
                if order is not None
                else return_marketplaces.get(
                    movement.id, unload_markets.get(movement.marketplace_unload_request_id)
                )
            )
            if marketplace is not None and row_marketplace != marketplace:
                continue
            document = docs.get(movement.id)
            if order is not None:
                document = {
                    "id": str(order.id),
                    "type": "fbs_order",
                    "number": order.external_order_id or str(order.wb_order_id),
                }
            count = abs(int(movement.quantity_delta)) if order is not None else 1
            sign = 1 if movement.quantity_delta >= 0 else -1
            values = (
                marks.get((order.id, product.id), [])
                if order is not None and offset is not None and movement.quantity_delta < 0
                else []
            )
            for unit in range(count):
                if after is not None and (movement_at, movement.id, unit) <= after:
                    continue
                if len(result) == limit:
                    assert last_key is not None
                    return result, _encode_cursor(*last_key)
                kiz = (
                    values[offset + unit]
                    if offset is not None and offset + unit < len(values)
                    else None
                )
                result.append(
                    {
                        "id": f"{movement.id}:{unit}" if order is not None else str(movement.id),
                        "movement_id": str(movement.id),
                        "occurred_at": movement_at.isoformat(),
                        "operation": "return"
                        if movement.id in return_marketplaces
                        else movement.movement_type,
                        "warehouse_id": str(movement.warehouse_id),
                        "product_id": str(product.id),
                        "sku": product.sku_code,
                        "product_name": product.name,
                        "shk": shk
                        if shk is not None
                        else (product.wb_barcode or alternate_barcodes.get(product.id)),
                        "size": product.wb_size,
                        "marketplace": row_marketplace,
                        "document": document,
                        "quantity_delta": sign
                        if order is not None
                        else int(movement.quantity_delta),
                        "kiz": kiz,
                    }
                )
                last_key = (movement_at, movement.id, unit)
        if len(batch) < _BATCH:
            break
        # The last movement of this batch was fully consumed in this request.
        resume_unit = False
    return result, None


def _excel_text(value: object) -> str:
    """Excel strings stay strings; control characters use reversible JSON escapes."""
    return json.dumps(str(value), ensure_ascii=False)[1:-1]


async def build_client_movement_workbook(
    session: AsyncSession, tenant_id: uuid.UUID, **filters: Any
) -> bytes:
    _period(filters["date_from"], filters["date_to"])
    if filters.get("sku") is not None and filters.get("shk") is not None:
        raise ValueError("sku and shk cannot be used together")
    workbook = Workbook(write_only=True)
    headers = [
        "Дата и время (UTC)",
        "Движение",
        "Причина",
        "Документ / заказ",
        "Код товара",
        "Товар",
        "Размер",
        "Штрихкод",
        "Приход, шт.",
        "Расход, шт.",
        "Склад",
        "КИЗ",
    ]
    sheets = {name: workbook.create_sheet(name) for name in ("WB", "Ozon", "Общие")}
    for sheet in sheets.values():
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = "A1:L1"
        for index, width in enumerate((23, 19, 27, 28, 22, 42, 12, 22, 14, 14, 24, 45), 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        heading = []
        for name in headers:
            cell = WriteOnlyCell(sheet, value=name)
            cell.font = Font(bold=True, color="000000")
            cell.fill = PatternFill("solid", fgColor="D9D9D9")
            heading.append(cell)
        sheet.append(heading)
    cursor: str | None = None
    while True:
        rows, next_cursor = await list_client_movements(
            session, tenant_id, **filters, cursor=cursor, limit=MAX_LIMIT
        )
        warehouse_ids = {uuid.UUID(row["warehouse_id"]) for row in rows}
        warehouses: dict[uuid.UUID, str] = {}
        if warehouse_ids:
            warehouse_rows = await session.execute(
                select(Warehouse.id, Warehouse.name).where(
                    Warehouse.tenant_id == tenant_id, Warehouse.id.in_(warehouse_ids)
                )
            )
            warehouses = {warehouse_id: name for warehouse_id, name in warehouse_rows}
        for row in rows:
            sheet = sheets[{"wb": "WB", "ozon": "Ozon"}.get(row["marketplace"], "Общие")]
            document = row["document"] or {}
            delta = row["quantity_delta"]
            values = [
                datetime.fromisoformat(row["occurred_at"]).strftime("%d.%m.%Y %H:%M:%S"),
                "Оприходование" if delta > 0 else "Списание",
                _EVENT_NAMES[row["operation"]],
                document.get("number"),
                row["sku"],
                row["product_name"],
                row["size"],
                row["shk"],
                delta if delta > 0 else None,
                -delta if delta < 0 else None,
                warehouses.get(uuid.UUID(row["warehouse_id"])),
                row["kiz"],
            ]
            cells = []
            for value in values:
                if value is None or isinstance(value, int):
                    cells.append(value)
                else:
                    cell = WriteOnlyCell(sheet, value=_excel_text(value))
                    cell.data_type = "s"
                    cells.append(cell)
            sheet.append(cells)
        if next_cursor is None:
            break
        cursor = next_cursor
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
