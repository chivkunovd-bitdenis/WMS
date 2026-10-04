"""Build the operator instruction without shortening its recovered source text."""
from pathlib import Path
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/WMS Print установка и сканирование.md'
OUTPUT = SOURCE.with_suffix('.docx')
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.7)
section.left_margin = section.right_margin = Inches(.75)
for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.1
for name, size in [('Title', 21), ('Heading 1', 17), ('Heading 2', 13)]:
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].paragraph_format.keep_with_next = True
    doc.styles[name].paragraph_format.space_before = Pt(12)
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)


def hyperlink(paragraph, title, url):
    link = OxmlElement('w:hyperlink')
    rel = paragraph.part.relate_to(
        url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
        is_external=True)
    link.set(qn('r:id'), rel)
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '175A91')
    props.append(color)
    run.append(props)
    text = OxmlElement('w:t')
    text.text = title
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def text_runs(paragraph, text):
    pattern = r'\[([^\]]+)\]\((https?://[^)]+)\)|(https?://[^\s]+)'
    cursor = 0
    for match in re.finditer(pattern, text):
        paragraph.add_run(text[cursor:match.start()])
        hyperlink(paragraph, match[1] or match[3], match[2] or match[3])
        cursor = match.end()
    paragraph.add_run(text[cursor:])


lines = SOURCE.read_text().splitlines()
for index, line in enumerate(lines):
    if not line.strip():
        continue
    if index == 0:
        doc.add_paragraph(line, 'Title')
    elif line.startswith('# '):
        p = doc.add_paragraph(line[2:], 'Heading 1')
    elif line.startswith('## '):
        doc.add_paragraph(line[3:], 'Heading 2')
    elif line.startswith('━'):
        doc.add_paragraph(line.strip('━ '), 'Heading 2')
    elif line.startswith('/bin/bash'):
        p = doc.add_paragraph()
        p.paragraph_format.keep_together = True
        p.paragraph_format.line_spacing = 1
        run = p.add_run(line)
        run.font.name = 'Courier New'
        run.font.size = Pt(9)
    else:
        p = doc.add_paragraph()
        if line.startswith('   '):
            p.paragraph_format.left_indent = Inches(.15)
        text_runs(p, line)
doc.core_properties.title = 'WMS Print установка и сканирование'
doc.core_properties.author = 'WMS'
doc.save(OUTPUT)
print(OUTPUT)
